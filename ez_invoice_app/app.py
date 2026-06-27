# app.py (Windows-friendly, Python 3.9+)
# Install:
#   python -m pip install streamlit pandas pymupdf pypdf
# Run:
#   python -m streamlit run app.py
from accounting_routing import DEFAULT_HOME_COMPANIES, apply_accounting_route
from item_classifier import ItemClassifier, classifier_settings_ui
from gst_invoice_parser import looks_like_gst_invoice, parse_gst_invoice
from qb_integration import (
    qb_sidebar,
    send_to_quickbooks,
    is_connected,
    resolve_line_account_for_export,
    complete_qb_auth_if_ready,
)
from tally_integration import build_tally_xml, make_tally_zip, send_to_tally, send_tally_batch, tally_is_connected, tally_sidebar
from zoho_integration import (
    build_zoho_export,
    complete_zoho_auth_if_ready,
    send_to_zoho,
    zoho_is_connected,
    zoho_sidebar,
)
from api_client import (
    EzInvoiceApiClient,
    EzInvoiceApiError,
    invoice_to_session_data,
    parser_mode_for_api,
    session_key_for_invoice,
)
from universal_parser import looks_like_structured_vendor_invoice, parse_generic_invoice
import streamlit as st
import pandas as pd
import json
import io
import zipfile
import re
import base64
import html
import os
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional, Tuple, Dict, Any, List
import openpyxl

APP_DIR = Path(__file__).resolve().parent
CLIENT_PROFILE_FILE = APP_DIR / "data" / "client_profile.json"
DEFAULT_CLIENT_PROFILE = {
    "client_name": "Client Workspace",
    "home_company_names": "",
    "primary_accounting_system": "TallyPrime",
    "tally_url": "http://localhost:9000",
    "tally_company": "",
}
classifier = ItemClassifier(str(APP_DIR / "subcat.csv"), use_seed_mappings=False)

# =========================
# Utilities
# =========================
DATE_PAT = r"\b\d{2}-[A-Z]{3}-\d{4}\b"    # 26-DEC-2023
DATE_PAT_SHORT = r"\b\d{2}-[A-Z]{3}-\d{2}\b"  # 23-OCT-24
INVNO_PAT = r"\b\d{6,}-\d{3,}\b"          # 22885940-21101
THREE_AMTS_PAT = r"([0-9,]+\.\d{2})\s+([0-9,]+\.\d{2})\s+([0-9,]+\.\d{2})"
CURRENCY_CODES = "USD|EUR|GBP|CAD|AUD|JPY|CHF|ZAR|ARS|RON"


def now_ts() -> str:
    return datetime.now().strftime("%Y-%m-%dT%H:%M:%S")


def _split_company_names(raw: Any) -> List[str]:
    if isinstance(raw, list):
        values = raw
    else:
        values = re.split(r"[\n,]+", str(raw or ""))
    return [value.strip() for value in values if value and value.strip()]


def _load_client_profile() -> Dict[str, Any]:
    profile = dict(DEFAULT_CLIENT_PROFILE)
    try:
        if CLIENT_PROFILE_FILE.exists():
            saved = json.loads(CLIENT_PROFILE_FILE.read_text(encoding="utf-8"))
            if isinstance(saved, dict):
                profile.update({k: v for k, v in saved.items() if v is not None})
    except Exception:
        pass
    return profile


def _save_client_profile(profile: Dict[str, Any]) -> bool:
    try:
        CLIENT_PROFILE_FILE.parent.mkdir(exist_ok=True)
        CLIENT_PROFILE_FILE.write_text(json.dumps(profile, indent=2, sort_keys=True), encoding="utf-8")
        return True
    except Exception:
        return False


def current_client_profile() -> Dict[str, Any]:
    if "client_profile" not in st.session_state:
        st.session_state["client_profile"] = _load_client_profile()
    return dict(st.session_state["client_profile"])


def current_home_company_names() -> List[str]:
    return _split_company_names(current_client_profile().get("home_company_names"))


def _api_status_label(status: str, fallback_ready: bool = False) -> str:
    normalized = str(status or "").strip().lower()
    labels = {
        "uploaded": "Uploaded",
        "extracted": "Extracted",
        "needs_review": "Needs Review",
        "validated": "Ready",
        "approved": "Ready",
        "posting": "Posting",
        "posted": "Sent",
        "failed": "Needs Review",
    }
    return labels.get(normalized, "Ready" if fallback_ready else "Needs Review")


def _api_bridge_active() -> bool:
    return bool(st.session_state.get("ez_api_bridge_active"))


def _api_bridge_client() -> Optional[EzInvoiceApiClient]:
    client = st.session_state.get("ez_api_client")
    return client if isinstance(client, EzInvoiceApiClient) else None


def _initialize_api_bridge() -> None:
    requested_mode = os.environ.get("EZ_INVOICE_BACKEND", "auto").strip().lower()
    if requested_mode not in {"auto", "api", "legacy"}:
        requested_mode = "auto"
    st.session_state["ez_backend_mode_requested"] = requested_mode
    st.session_state["ez_api_bridge_active"] = False
    st.session_state.pop("ez_api_bridge_error", None)
    if requested_mode == "legacy":
        return

    client = EzInvoiceApiClient(
        base_url=os.environ.get("EZ_API_BASE_URL", "http://127.0.0.1:8000"),
        timeout=float(os.environ.get("EZ_API_TIMEOUT_SECONDS", "30")),
    )
    try:
        health = client.health()
        if health.get("status") != "ok":
            raise EzInvoiceApiError("EZ-Invoice API health check did not return OK.")
        profile = current_client_profile()
        client_name = (
            os.environ.get("EZ_ORGANIZATION_NAME")
            or profile.get("client_name")
            or "Client Workspace"
        )
        legal_names = current_home_company_names() or [str(client_name)]
        client.authenticate_for_pilot(
            email=os.environ.get("EZ_API_EMAIL", "owner@ezinvoice.local").strip(),
            password=os.environ.get(
                "EZ_API_PASSWORD",
                "local-development-password",
            ),
            full_name=os.environ.get("EZ_API_FULL_NAME", "SiftEntry Owner").strip(),
            organization_name=str(client_name),
            legal_names=legal_names,
            default_currency=str(profile.get("default_currency") or "USD"),
        )
        organization = client.ensure_organization(
            name=str(client_name),
            legal_names=legal_names,
            default_currency=str(profile.get("default_currency") or "USD"),
            organization_id=os.environ.get("EZ_ORGANIZATION_ID", "").strip(),
        )
    except EzInvoiceApiError as exc:
        st.session_state["ez_api_bridge_error"] = str(exc)
        if requested_mode == "api":
            st.error(
                "FastAPI mode is enabled but the backend is unavailable. "
                "Start FastAPI on port 8000 or set EZ_INVOICE_BACKEND=legacy. "
                + str(exc)
            )
        return

    st.session_state["ez_api_client"] = client
    st.session_state["ez_api_organization"] = organization
    st.session_state["ez_api_bridge_active"] = True


def _api_summary_row(fname: str, data: Dict[str, Any]) -> Dict[str, Any]:
    inv = data.get("payload", {}).get("INVOICE", {})
    header = inv.get("INVOICE HEADER", {})
    currency = inv.get("PAYMENT", {}).get("ELECTRONIC", {}).get("CURRENCY", "USD")
    return {
        "File": fname,
        "Invoice #": header.get("INVOICE NO.", ""),
        "Direction": inv.get("ROUTING", {}).get("DIRECTION", "unknown"),
        "Date": header.get("INVOICE DATE", ""),
        "Items": len(inv.get("LINE ITEMS", {}).get("ROWS", []) or []),
        "Total": f"{currency} {money(header.get('INVOICE AMOUNT', 0))}",
        "Status": "✅" if data.get("ok") else "⚠️",
        "Mode": inv.get("DOCUMENT", {}).get("PARSER", ""),
        "Pages": data.get("pages", 0),
        "Speed": "",
    }


def _sync_api_queue() -> Dict[str, str]:
    client = _api_bridge_client()
    organization = st.session_state.get("ez_api_organization") or {}
    if not client or not organization.get("id"):
        return {}

    invoices = client.list_invoices(str(organization["id"]))
    previous = st.session_state.get("payloads", {})
    previous_by_id = {
        str(data.get("api_invoice_id")): data
        for data in previous.values()
        if data.get("api_invoice_id")
    }
    used: set[str] = set()
    payloads: Dict[str, Dict[str, Any]] = {}
    id_to_key: Dict[str, str] = {}
    for invoice in invoices:
        key = session_key_for_invoice(invoice, used)
        data = invoice_to_session_data(invoice)
        data["missing_df"] = pd.DataFrame()
        prior = previous_by_id.get(str(invoice.get("id")))
        if prior:
            for preserved_key in ("tally_status", "text", "bytes"):
                if prior.get(preserved_key):
                    data[preserved_key] = prior[preserved_key]
        payloads[key] = data
        id_to_key[str(invoice.get("id"))] = key

    st.session_state.payloads = payloads
    st.session_state.summary = [
        _api_summary_row(fname, data) for fname, data in payloads.items()
    ]
    st.session_state.ptimes = {
        fname: st.session_state.get("ptimes", {}).get(fname, 0)
        for fname in payloads
    }
    selected = st.session_state.get("selected_invoice_file")
    if selected not in payloads:
        st.session_state["selected_invoice_file"] = next(iter(payloads), None)
    return id_to_key


def _record_api_posting_result(
    payload_data: Dict[str, Any],
    target: str,
    result: Dict[str, Any],
    dry_run: bool = False,
) -> None:
    client = _api_bridge_client()
    invoice_id = str(payload_data.get("api_invoice_id") or "")
    if not client or not invoice_id:
        return
    try:
        client.record_posting_result(
            invoice_id=invoice_id,
            target=target,
            result=result,
            dry_run=dry_run,
        )
        if not dry_run:
            payload_data["api_status"] = "posted" if result.get("success") else "failed"
            payload_data["ok"] = bool(result.get("success"))
    except EzInvoiceApiError as exc:
        st.warning("Posting completed, but FastAPI could not save its audit result: " + str(exc))


def normalize_text(text: str) -> str:
    return (text or "").replace("\u00a0", " ")


def normalize_lines(text: str) -> List[str]:
    return [ln.strip() for ln in normalize_text(text).splitlines() if ln.strip()]


def to_float(x: Any) -> Optional[float]:
    if x is None:
        return None
    s = str(x).strip()
    if not s:
        return None
    s = re.sub(r"[^0-9.\-]", "", s)
    try:
        return float(s)
    except Exception:
        return None


def safe_float0(x: Any) -> float:
    v = to_float(x)
    return float(v) if v is not None else 0.0


def money(v: Any) -> str:
    try:
        return f"{float(v):,.2f}"
    except Exception:
        return str(v)


def amount_close(a: Any, b: Any, tolerance: float = 0.50) -> bool:
    return abs(safe_float0(a) - safe_float0(b)) <= tolerance


def snippet_evidence(text: str, value: str, context: int = 70) -> str:
    if not text or not value:
        return ""
    t = normalize_text(text)
    idx = t.find(value)
    if idx == -1:
        idx = t.lower().find(str(value).lower())
        if idx == -1:
            return ""
    start = max(0, idx - context)
    end = min(len(t), idx + len(value) + context)
    snip = t[start:end].replace("\n", " ")
    return ("..." if start > 0 else "") + snip + ("..." if end < len(t) else "")


def normalize_invoice_date_value(value: str) -> str:
    value = (value or "").strip().upper()
    if re.fullmatch(r"\d{2}-[A-Z]{3}-\d{2}", value):
        value = value[:-2] + "20" + value[-2:]
    return value


def add_days_to_invoice_date(invoice_date: str, days: int) -> str:
    try:
        dt = datetime.strptime(normalize_invoice_date_value(invoice_date), "%d-%b-%Y")
        return (dt + timedelta(days=days)).strftime("%d-%b-%Y").upper()
    except Exception:
        return ""


# =========================
# PDF -> Text + Viewer
# =========================
def pdf_page_count(pdf_bytes: bytes) -> int:
    try:
        import fitz
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        return doc.page_count
    except Exception:
        pass
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(pdf_bytes))
        return len(reader.pages)
    except Exception:
        return 0


def pdf_to_text(pdf_bytes: bytes) -> Tuple[str, str]:
    # PyMuPDF first
    try:
        import fitz
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        text = "\n".join([(p.get_text("text") or "") for p in doc]).strip()
        if text:
            return text, "pymupdf"
    except Exception:
        pass

    # fallback
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(pdf_bytes))
        text = "\n".join([(p.extract_text() or "") for p in reader.pages]).strip()
        if text:
            return text, "pypdf"
    except Exception:
        pass

    return "", "none"


def pdf_embed_html(pdf_bytes: bytes, height_px: int = 720) -> str:
    b64 = base64.b64encode(pdf_bytes).decode("utf-8")
    return f"""
    <iframe
      src="data:application/pdf;base64,{b64}"
      width="100%"
      height="{height_px}"
      style="border:1px solid #e2e8f0;border-radius:14px;">
    </iframe>
    """


# =========================
# Parser (keep structure, improve robustness)
# =========================
def section_between(text: str, start_marker: str, end_markers: List[str]) -> str:
    t = normalize_text(text)
    up = t.upper()
    s = up.find(start_marker.upper())
    if s == -1:
        return ""
    start = s + len(start_marker)
    ends = []
    for m in end_markers:
        p = up.find(m.upper(), start)
        if p != -1:
            ends.append(p)
    end = min(ends) if ends else len(t)
    return t[start:end]


def find_electronic_amount(text: str) -> Tuple[str, float]:
    t = normalize_text(text)
    m = re.search(r"\bELECTRONIC\s+(" + CURRENCY_CODES + r")\s+([0-9,]+\.\d{2})\b", t, re.IGNORECASE)
    if m:
        return m.group(1).upper(), safe_float0(m.group(2))

    lines = normalize_lines(text)
    for i, ln in enumerate(lines):
        if "PLEASE REMIT THIS AMOUNT" in ln.upper():
            for cand in lines[i + 1:i + 16]:
                m_remit = re.search(r"\b(" + CURRENCY_CODES + r")\s+([0-9,]+\.\d{2})\b", cand, re.IGNORECASE)
                if m_remit:
                    return m_remit.group(1).upper(), safe_float0(m_remit.group(2))

    m2 = re.search(r"\b(" + CURRENCY_CODES + r")\s+([0-9,]+\.\d{2})\b", t)
    if m2:
        return m2.group(1).upper(), safe_float0(m2.group(2))
    return "USD", 0.0


def find_exchange_rate(text: str) -> Dict[str, Any]:
    t = normalize_text(text)
    m = re.search(r"\b([0-9,]+\.\d{4,6})\s+(" + CURRENCY_CODES + r")/(" + CURRENCY_CODES + r")\b", t, re.IGNORECASE)
    if not m:
        return {}
    return {
        "RATE": safe_float0(m.group(1)),
        "FROM": m.group(2).upper(),
        "TO": m.group(3).upper(),
        "TEXT": m.group(0),
    }


def find_currencies(text: str) -> List[str]:
    return list(dict.fromkeys(c.upper() for c in re.findall(r"\b(" + CURRENCY_CODES + r")\b", normalize_text(text), re.IGNORECASE)))


_STRUCTURED_INVNO_PAT = r"\d{6,8}-\d{5}(?:-\d+)?"
_STRUCTURED_DATE_ANY_PAT = r"\d{1,2}-[A-Z]{3}-\d{2,4}"
_STRUCTURED_CUSTOMER_PAT = r"\d{3,8}"


def _normalize_structured_date(value: str) -> str:
    value = (value or "").strip().upper()
    if re.fullmatch(r"\d{1,2}-[A-Z]{3}-\d{2}", value):
        value = value[:-2] + "20" + value[-2:]
    return value


def _looks_like_customer_no(value: str) -> bool:
    value = (value or "").strip()
    return bool(re.fullmatch(_STRUCTURED_CUSTOMER_PAT, value)) and not set(value) <= {"0"}


def _invoice_header_windows(text: str) -> List[str]:
    lines = normalize_lines(text)
    windows: List[str] = []
    label_re = re.compile(
        r"(CUSTOMER\s+NO|CLIENTE\s+NR|INVOICE\s+NO|FATTURA\s+NR|INVOICE\s+DATE|DATA\s+FATTURA)",
        re.IGNORECASE,
    )
    for idx, line in enumerate(lines):
        if label_re.search(line):
            windows.append("\n".join(lines[max(0, idx - 10): min(len(lines), idx + 16)]))
    windows.append(normalize_text(text))
    return windows


def _extract_triplet_from_window(window: str) -> Tuple[str, str, str]:
    patterns = [
        re.compile(
            rf"(?<![.\d])(?P<customer>{_STRUCTURED_CUSTOMER_PAT})\s+"
            rf"(?P<invoice>{_STRUCTURED_INVNO_PAT})\s+"
            rf"(?P<date>{_STRUCTURED_DATE_ANY_PAT})(?![A-Z0-9-])",
            re.IGNORECASE,
        ),
        re.compile(
            rf"(?P<date>{_STRUCTURED_DATE_ANY_PAT})\s*"
            rf"(?P<invoice>{_STRUCTURED_INVNO_PAT})\s*"
            rf"(?P<customer>{_STRUCTURED_CUSTOMER_PAT})(?![.\d])",
            re.IGNORECASE,
        ),
        re.compile(
            rf"(?P<invoice>{_STRUCTURED_INVNO_PAT})\s+"
            rf"(?P<date>{_STRUCTURED_DATE_ANY_PAT})\s+"
            rf"(?P<customer>{_STRUCTURED_CUSTOMER_PAT})(?![.\d])",
            re.IGNORECASE,
        ),
        re.compile(
            rf"(?P<invoice>{_STRUCTURED_INVNO_PAT})\s+"
            rf"(?P<customer>{_STRUCTURED_CUSTOMER_PAT})\s+"
            rf"(?P<date>{_STRUCTURED_DATE_ANY_PAT})(?![A-Z0-9-])",
            re.IGNORECASE,
        ),
    ]
    for pattern in patterns:
        for match in pattern.finditer(window):
            customer = match.group("customer")
            if _looks_like_customer_no(customer):
                return customer, match.group("invoice"), _normalize_structured_date(match.group("date"))
    return "", "", ""


def _is_colt_invoice(text: str) -> bool:
    up = normalize_text(text).upper()
    return "TRIP NO." in up and "SERVICE DATE" in up


def _parse_colt_flight_location(text: str) -> Dict[str, str]:
    out = {
        "FUEL TICKET": "",
        "AIRCRAFT TYPE": "",
        "TAIL NO.": "",
        "TERRITORY": "",
        "DATE UPLIFTED": "",
        "TERMS": "",
        "LOCATION": "",
        "CONTACT": "",
        "DESTINATION": "",
        "TRIP NO.": "",
        "FLIGHT NO.": "",
    }
    lines = normalize_lines(text)
    service_idx = next((i for i, ln in enumerate(lines) if ln.upper() == "SERVICE DATE"), None)
    trip_idx = next((i for i, ln in enumerate(lines) if ln.upper() == "TRIP NO."), None)

    if trip_idx is not None and service_idx is not None:
        values = [
            ln for ln in lines[trip_idx + 1:service_idx]
            if ln.upper() not in {
                "FLIGHT NO.", "TAIL NO.", "AIRCRAFT TYPE",
                "SERVICE TYPE", "TERMS", "CONTACT"
            }
        ]
        if values:
            out["TRIP NO."] = values[0]
        tail_re = re.compile(r"^(?:N\d{2,}[A-Z]*|[A-Z]-[A-Z]{3,5}|[A-Z]{1,2}-[A-Z0-9]{2,5}|T7-[A-Z]+|C-[A-Z]{3,5})$")
        tail_pos = None
        for idx, value in enumerate(values[1:], start=1):
            if tail_re.match(value):
                out["TAIL NO."] = value
                tail_pos = idx
                break
        if tail_pos is not None:
            for value in values[tail_pos + 1:]:
                up = value.upper()
                if "INVOICE" in up or "CONTRACT" in up or re.search(r"\d+\s+(?:D FROM INV|NET)", up):
                    break
                if "," in value:
                    break
                out["AIRCRAFT TYPE"] = (out["AIRCRAFT TYPE"] + " " + value).strip()
        for value in values:
            m_terms = re.search(r"\d+\s+(?:D FROM INV|NET[^\s]*)", value, re.IGNORECASE)
            if m_terms:
                out["TERMS"] = m_terms.group(0).strip()
            if "," in value and not out["CONTACT"]:
                out["CONTACT"] = value.strip()

    date_idx = next((i for i, ln in enumerate(lines) if re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", ln)), None)
    if date_idx is not None:
        out["DATE UPLIFTED"] = lines[date_idx]
        codes: List[str] = []
        j = date_idx + 1
        while j < len(lines) and re.fullmatch(r"[A-Z0-9]{3,4}", lines[j]) and len(codes) < 2:
            codes.append(lines[j])
            j += 1
        if codes:
            out["LOCATION"] = " / ".join(codes)
            out["DESTINATION"] = codes[-1]

    country_labels = {"COUNTRY", "COUNTRY SEQUENCE NO.", "REFERENCE NO. / PURCHASE ORDER", "PLEASE REMIT THIS AMOUNT"}
    for i, line in enumerate(lines):
        if line.upper() == "COUNTRY":
            for candidate in lines[i + 1:i + 8]:
                up = candidate.upper()
                if up in country_labels or re.search(r"\d", up):
                    continue
                if re.fullmatch(r"[A-Z][A-Z\s,.'-]+", up):
                    out["TERRITORY"] = candidate.strip()
                    break
            break

    return out


def find_invoice_triplet(text: str) -> Tuple[str, str, str]:
    for window in _invoice_header_windows(text):
        customer, invoice, date = _extract_triplet_from_window(window)
        if customer and invoice and date:
            return customer, invoice, date
    return "", "", ""


def find_due_date(text: str) -> str:
    t = normalize_text(text)
    m = re.search(r"\b(?:PAYMENT\s+)?DUE DATE\b[^\n]*\n(?:[^\n]*\n){0,2}?\s*(" + DATE_PAT + r")\b", t, re.IGNORECASE)
    if m:
        return m.group(1)

    lines = normalize_lines(text)
    for i, ln in enumerate(lines):
        if "DUE DATE" in ln.upper():
            for j in range(i + 1, min(i + 6, len(lines))):
                m2 = re.search(DATE_PAT, lines[j])
                if m2:
                    return m2.group(0)
    _, _, inv_date = find_invoice_triplet(text)
    m_terms = re.search(r"\b(\d{1,3})\s*(?:D\s+FROM\s+INV|NET)\b", t, re.IGNORECASE)
    if inv_date and m_terms:
        return add_days_to_invoice_date(inv_date, int(m_terms.group(1)))
    return ""


def find_sales_order_no(text: str) -> str:
    lines = normalize_lines(text)
    for i, ln in enumerate(lines):
        if ln.upper().startswith("SALES ORDER NO"):
            for j in range(i + 1, min(i + 10, len(lines))):
                v = lines[j].strip()
                if re.fullmatch(r"\d+", v):
                    return v
            return ""
    m = re.search(r"\bSALES ORDER NO\.?[^\n]*\n(?:[^\n]*\n){0,3}?\s*([0-9]+)\b", normalize_text(text), re.IGNORECASE)
    return m.group(1) if m else ""


def find_po_no(text: str) -> str:
    t = normalize_text(text)
    m = re.search(r"\bPO NO\./CONTRACT NO\.\b\s*\n\s*([A-Z0-9/.\-]+)\b", t, re.IGNORECASE)
    return m.group(1) if m else ""


def find_page_no(text: str) -> str:
    t = normalize_text(text)
    for m in re.finditer(r"\b(\d{1,3})\s*-\s*(\d{1,3})\b", t):
        return f"{m.group(1)} - {m.group(2)}"
    return ""


def parse_bill_to_and_seller(text: str) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    lines = normalize_lines(text)

    seller = {"NAME": "WORLD FUEL SERVICES, INC.", "ADDRESS": [], "Tel": "", "Email": "", "Internet": ""}
    bill_to = {"NAME": "", "ADDRESS": []}

    try:
        idx = lines.index("WORLD FUEL SERVICES, INC.")
    except ValueError:
        idx = -1
    if idx >= 2:
        seller["ADDRESS"] = [lines[idx - 2], lines[idx - 1]]

    # Tel / Email / Internet
    m_tel = re.search(r"\bTel:\s*([^\n]+)", text)
    if m_tel:
        tel_line = m_tel.group(1).strip()
        if "Email:" in tel_line:
            tel_part, email_part = tel_line.split("Email:", 1)
            seller["Tel"] = tel_part.strip()
            seller["Email"] = email_part.strip()
        else:
            seller["Tel"] = tel_line

    m_email = re.search(r"\bEmail:\s*([^\n]+)", text)
    if m_email:
        seller["Email"] = m_email.group(1).strip()

    m_inet = re.search(r"\bInternet:\s*([^\n]+)", text)
    if m_inet:
        seller["Internet"] = m_inet.group(1).strip()

    # Bill-to begins after Internet: line
    start = 0
    for i, ln in enumerate(lines):
        if ln.lower().startswith("internet:"):
            start = i + 1
            break
    end = max(start, idx - 2) if idx != -1 else len(lines)
    bill_block = lines[start:end]
    if bill_block:
        bill_to["NAME"] = bill_block[0]
        bill_to["ADDRESS"] = bill_block[1:]

    return bill_to, seller


def parse_flight_location(text: str) -> Dict[str, str]:
    """
    Context parser for structured supplier invoice variants.

    Variant A (older): marker 'FLIGHT NO. TERMS' exists.
    Variant B (newer): headers are on separate lines:
        FLIGHT NO.
        TERMS
        TAIL NO.
        LOCATION
        TERRITORY
        CONTACT
        ...
        <values...>
        TAX AMOUNT
        DATE UPLIFTED
        FUEL TICKET
        AIRCRAFT TYPE

    Keeps your JSON schema the same; only improves extraction.
    """
    if _is_colt_invoice(text):
        return _parse_colt_flight_location(text)

    out = {
        "FUEL TICKET": "",
        "AIRCRAFT TYPE": "",
        "TAIL NO.": "",
        "TERRITORY": "",
        "DATE UPLIFTED": "",
        "TERMS": "",
        "LOCATION": "",
        "CONTACT": "",
        "DESTINATION": ""
    }
    # --- COLT FORMAT: detect and handle separately ---
    t_check = normalize_text(text)
    if "TRIP NO." in t_check and ("SERVICE DATE" in t_check or "SERVICE TYPE" in t_check):
        colt_lines = normalize_lines(text)
        for i, l in enumerate(colt_lines):
            if l.startswith("TRIP NO.") and "TAIL NO." in l and i + 1 < len(colt_lines):
                val = colt_lines[i + 1]
                parts = val.split()
                if parts:
                    out["FUEL TICKET"] = parts[0]
                tail_re = re.compile(r'(N\d{2,}[A-Z]*|[A-Z]{1,2}-[A-Z]{2,5}|\d[A-Z]-[A-Z]{2,4}|T7-[A-Z]+|C-[A-Z]{3,5})')
                m = tail_re.search(val)
                if m:
                    out["TAIL NO."] = m.group(1)
                for kw in ["Dassault", "Challenger", "Gulfstream", "Bombardier", "Cessna", "Boeing", "Pilatus",
                           "Embraer"]:
                    if kw in val:
                        idx = val.index(kw)
                        rest = val[idx:]
                        for stop in ["Service Invoice", "3rd Party", "Contract", "Retail"]:
                            if stop in rest:
                                rest = rest[:rest.index(stop)].strip()
                        out["AIRCRAFT TYPE"] = rest.strip()
                        break
                m_t = re.search(r'(\d+\s+D FROM INV|\d+\s+NET\S*)', val)
                if m_t:
                    out["TERMS"] = m_t.group(1)
                m_c = re.search(r'([A-Z][a-z]+,\s*(?:Ms\.?\s*|Mr\.?\s*)?[A-Z][a-z]+)', val)
                if m_c:
                    out["CONTACT"] = m_c.group(1)
                break
        for l in colt_lines:
            m = re.match(r'^(\d{2}/\d{2}/\d{4})\s+(.*)', l)
            if m:
                out["DATE UPLIFTED"] = m.group(1)
                rest = m.group(2).strip().split()
                found_codes = []
                for token in rest:
                    if re.match(r'^[A-Z]{3,4}$', token) and token not in ("FUEL", "ADMIN", "FEES"):
                        found_codes.append(token)
                    else:
                        break
                if len(found_codes) >= 2:
                    out["LOCATION"] = found_codes[0]
                    out["DESTINATION"] = found_codes[1]
                elif len(found_codes) == 1:
                    out["DESTINATION"] = found_codes[0]
                break
        m_country = re.search(r'\n([A-Z][A-Z\s]+?)\s+[A-Z]{2}-\d{5}', t_check)
        if m_country:
            out["TERRITORY"] = m_country.group(1).strip()
        return out
    # --- END COLT FORMAT ---

    t = normalize_text(text)

    # Destination is usually outside the flight block
    m_dest = re.search(r"\bDESTINATION\b\s*\n\s*([^\n\r]+)", t, re.IGNORECASE)
    if m_dest:
        out["DESTINATION"] = m_dest.group(1).strip()

    # 1) Get a flight block.
    flight_block = section_between(
        text,
        start_marker="FLIGHT NO. TERMS",
        end_markers=["TAX AMOUNT", "DATE UPLIFTED", "FUEL TICKET", "AIRCRAFT TYPE", "DESCRIPTION", "MAIL INSTRUCTIONS", "COMMENTS"]
    )
    if not flight_block.strip():
        # Variant B
        flight_block = section_between(
            text,
            start_marker="FLIGHT NO.",
            end_markers=["TAX AMOUNT", "DATE UPLIFTED", "FUEL TICKET", "AIRCRAFT TYPE", "DESCRIPTION", "MAIL INSTRUCTIONS", "COMMENTS"]
        )
    if not flight_block.strip():
        idx = t.upper().find("FLIGHT NO.")
        if idx != -1:
            flight_block = t[idx: idx + 2500]

    lines = normalize_lines(flight_block)

    # Remove header labels that show up as individual lines
    header_labels = {
        "FLIGHT NO.", "TERMS", "TAIL NO.", "LOCATION", "TERRITORY", "CONTACT",
        "QUANTITY", "UNIT PRICE", "DATE UPLIFTED", "FUEL TICKET", "AIRCRAFT TYPE"
    }
    header_prefixes = (
        "FLIGHT NO", "VOLO NR", "TERMS", "TERMINI", "TAIL NO", "NUMERO DI CODA",
        "LOCATION", "LOCALITA", "TERRITORY", "TERRITORIO", "CONTACT", "CONTATTO",
        "QUANTITY", "QUANTITA", "UNIT PRICE", "PREZZO UNITARIO", "TAX AMOUNT",
        "IMPORTO TASSE", "DATE UPLIFTED", "DATA RIFORNIMENTO", "FUEL TICKET",
        "BOLLA NUM", "AIRCRAFT TYPE", "TIPO AEREOMOBILE", "DESCRIPTION",
        "DESCRIZIONE", "EXTENDED AMOUNT", "IMPORTO", "MAIL INSTRUCTIONS",
        "ISTRUZIONI DI POSTA", "EXCHANGE RATE", "TASSO DI CAMBIO",
        "COUNTRY SEQUENCE", "N. SEQUENZA", "DESTINATION", "DESTINAZIONE",
    )

    def is_header_noise(line: str) -> bool:
        up = line.upper().strip().rstrip("/")
        return up in header_labels or any(up.startswith(prefix) for prefix in header_prefixes)

    value_lines = [ln for ln in lines if not is_header_noise(ln)]

    term_idx = next((i for i, ln in enumerate(value_lines) if re.fullmatch(r"\d+\s+(?:NET|D\s+FROM\s+INV)", ln.strip(), re.IGNORECASE)), None)
    if term_idx is not None:
        if term_idx - 1 >= 0:
            out["TAIL NO."] = value_lines[term_idx - 1].strip()
        if term_idx - 2 >= 0:
            out["TERRITORY"] = value_lines[term_idx - 2].strip()
        if term_idx - 3 >= 0:
            out["AIRCRAFT TYPE"] = value_lines[term_idx - 3].strip()
        if term_idx - 4 >= 0:
            out["FUEL TICKET"] = value_lines[term_idx - 4].strip()
        out["TERMS"] = value_lines[term_idx].strip()
        if term_idx + 1 < len(value_lines) and re.fullmatch(DATE_PAT, value_lines[term_idx + 1].strip()):
            out["DATE UPLIFTED"] = value_lines[term_idx + 1].strip()
        if term_idx + 2 < len(value_lines) and "/" in value_lines[term_idx + 2]:
            out["LOCATION"] = value_lines[term_idx + 2].strip()
        if term_idx + 3 < len(value_lines):
            out["CONTACT"] = value_lines[term_idx + 3].strip()

    # 2) Terms + date uplifted
    m_terms = re.search(r"\b(\d+)\s+(?:NET|D\s+FROM\s+INV)[^\n\r]*", flight_block, re.IGNORECASE)
    if m_terms and not out["TERMS"]:
        out["TERMS"] = m_terms.group(0).strip()

    dates = list(re.finditer(DATE_PAT, flight_block))
    if dates and not out["DATE UPLIFTED"]:
        if m_terms:
            after = [d for d in dates if d.start() > m_terms.start()]
            pick = after[0] if after else dates[0]
        else:
            pick = dates[0]
        out["DATE UPLIFTED"] = pick.group(0)

    # 3) Location + Contact (inside flight block)
    m_loc = re.search(r"\b([A-Z]{3}\s*/\s*[A-Z]{4})\b", flight_block)
    if m_loc and not out["LOCATION"]:
        out["LOCATION"] = m_loc.group(1).strip()

    m_contact = re.search(r"\b([A-Za-z]+,\s*[A-Za-z]+)\b", flight_block)
    if m_contact and not out["CONTACT"]:
        out["CONTACT"] = m_contact.group(1).strip()

    # 4) Tail + Territory
    tail_idx = None
    for i, ln in enumerate(value_lines):
        if out["TAIL NO."]:
            break
        up = ln.upper()
        if "ANY A/C/" in up or re.search(r"\bN[0-9][0-9A-Z-]{2,}\b", ln):
            m_any = re.search(r"(ANY A/C/)?(N[0-9][0-9A-Z-]{2,})", ln)
            if m_any:
                out["TAIL NO."] = (("ANY A/C/" if m_any.group(1) else "") + m_any.group(2)).strip()
                tail_idx = i
                break
            out["TAIL NO."] = ln.strip()
            tail_idx = i
            break

    if tail_idx is not None and tail_idx - 1 >= 0 and not out["TERRITORY"]:
        cand = value_lines[tail_idx - 1].strip()
        if cand.isupper() and "/" not in cand and not re.search(r"\d", cand):
            out["TERRITORY"] = cand

    # 5) Fuel ticket + aircraft type
    ticket_pat = re.compile(r"^(?:INV)?[A-Z]*\d{4,}$", re.IGNORECASE)
    for i, ln in enumerate(value_lines):
        if out["FUEL TICKET"]:
            break
        s = ln.strip()
        if not s or re.fullmatch(DATE_PAT, s):
            continue
        if "/" in s or "," in s or "NET" in s.upper() or "ANY A/C/N" in s.upper():
            continue

        if re.fullmatch(r"\d{4,6}", s):  # numeric ticket
            out["FUEL TICKET"] = s
            if i + 1 < len(value_lines):
                nxt = value_lines[i + 1].strip()
                if nxt and nxt != out["TERRITORY"] and nxt != out["TAIL NO."] and "NET" not in nxt.upper():
                    out["AIRCRAFT TYPE"] = nxt
            break

        if ticket_pat.match(s):  # e.g., INV00037984
            out["FUEL TICKET"] = s
            if i + 1 < len(value_lines):
                nxt = value_lines[i + 1].strip()
                if nxt and (("A/C" in nxt.upper()) or (len(nxt) > 6 and "NET" not in nxt.upper() and "/" not in nxt and "," not in nxt)):
                    out["AIRCRAFT TYPE"] = nxt
            break

    if out["TAIL NO."].upper().endswith("NET"):
        out["TAIL NO."] = ""

    if not out["AIRCRAFT TYPE"]:
        for ln in value_lines:
            up = ln.upper()
            if "A/C" in up and "ANY A/C/N" not in up:
                out["AIRCRAFT TYPE"] = ln.strip()
                break

    return out



def parse_summary_totals(text: str) -> Dict[str, Any]:
    t = normalize_text(text)
    label = "CUSTOMER NO. INVOICE NO. INVOICE DATE"
    idx = t.upper().rfind(label)
    window = t[max(0, idx - 2600):idx] if idx != -1 else t

    ms = list(re.finditer(THREE_AMTS_PAT, window))
    if ms:
        m = ms[0]
        return {"USD USD USD": {
            "TOTAL EXCLUSIVE OF TAX": safe_float0(m.group(1)),
            "TOTAL TAX": safe_float0(m.group(2)),
            "INVOICE TOTAL": safe_float0(m.group(3)),
        }}
    return {"USD USD USD": {}}


def parse_invoice_tax_summary(text: str) -> Dict[str, Any]:
    t = normalize_text(text)
    out: Dict[str, Any] = {}

    def grab(label: str) -> Optional[float]:
        m = re.search(r"\b" + re.escape(label) + r"\b\s+([0-9,]+\.\d{2})", t, re.IGNORECASE)
        return to_float(m.group(1)) if m else None

    for k in ["TOTAL EXCLUSIVE OF TAX", "AMOUNT SUBJECT TO TAX", "AMOUNT NOT SUBJECT TO TAX", "TOTAL TAX", "INVOICE TOTAL"]:
        v = grab(k)
        if v is not None:
            out[k] = float(v)

    m_sales = re.search(r"\bSALES TAX\s*@\s*([0-9.]+%)\s+([0-9,]+\.\d{2})", t, re.IGNORECASE)
    if m_sales:
        out["SALES TAX @ " + m_sales.group(1)] = safe_float0(m_sales.group(2))

    return out


def parse_comments(text: str) -> Dict[str, str]:
    lines = normalize_lines(text)
    comments = {
        "FBO Ticket #": "",
        "Alliance Card No.": "",
        "Customer Card No.": "",
        "FBO Name": "",
        "SUPPLIER NAME": "",
    }

    start = -1
    for i, ln in enumerate(lines):
        if ln.upper() == "COMMENTS":
            start = i + 1
            break
    if start == -1:
        return comments

    block = []
    for j in range(start, len(lines)):
        if lines[j].upper().startswith("OR WIRE TO"):
            break
        block.append(lines[j])

    def set_if(prefix: str, key: str):
        for ln in block:
            if ln.upper().startswith(prefix.upper()):
                parts = ln.split(":", 1)
                comments[key] = (parts[1].strip() if len(parts) > 1 else "").strip()
                return

    set_if("FBO Ticket #", "FBO Ticket #")
    set_if("Alliance Card No.", "Alliance Card No.")
    set_if("Customer Card No.", "Customer Card No.")
    set_if("FBO Name", "FBO Name")
    set_if("Supplier Name", "SUPPLIER NAME")
    set_if("Supplier", "SUPPLIER NAME")
    return comments


def parse_remit_and_wire(text: str) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    lines = normalize_lines(text)
    remit = {"WORLD FUEL SERVICES, INC.": []}
    wire = {"Bank": "", "SWIFT": "", "FedWire ABA": "", "ACH ABA": "", "ACCT": "", "ACCT#": ""}

    # REMIT TO:
    s = -1
    for i, ln in enumerate(lines):
        if ln.upper().startswith("REMIT TO"):
            s = i + 1
            break
    if s != -1:
        block = []
        for j in range(s, len(lines)):
            up = lines[j].upper()
            if up == "COMMENTS" or up.startswith("OR WIRE TO"):
                break
            if up.startswith("SALES TAX @"):
                continue
            block.append(lines[j])
        remit["WORLD FUEL SERVICES, INC."] = block

    # OR WIRE TO:
    s2 = -1
    for i, ln in enumerate(lines):
        if ln.upper().startswith("OR WIRE TO"):
            s2 = i + 1
            break
    if s2 != -1:
        for j in range(s2, len(lines)):
            ln = lines[j]
            up = ln.upper()
            if re.match(r"^\d+\s*-\s*\d+$", ln):
                break

            if up.startswith("BANK OF AMERICA"):
                wire["Bank"] = ln.strip()
            elif up.startswith("SWIFT:"):
                wire["SWIFT"] = ln.split(":", 1)[1].strip()
            elif up.startswith("FEDWIRE ABA:"):
                wire["FedWire ABA"] = ln.split(":", 1)[1].strip()
            elif up.startswith("ACH ABA:"):
                wire["ACH ABA"] = ln.split(":", 1)[1].strip()
            elif up.startswith("ACCT:"):
                wire["ACCT"] = ln.split(":", 1)[1].strip()
            elif up.startswith("ACCT#"):
                wire["ACCT#"] = ln.replace("ACCT#", "").replace(":", "").strip()

    return remit, wire


def parse_line_items_inline_rows(text: str) -> List[Dict[str, Any]]:
    """
    One-line row format:
      FACILITY USE FEE 1 EA 250.00000 USD/EA 250.00 0.00 250.00
    """
    lines = normalize_lines(text)
    row_pat = re.compile(
        r"^(?P<desc>.+?)\s+"
        r"(?P<qty>[0-9,]+(?:\.\d+)?)\s+"
        r"(?P<uom>[A-Z]{1,4})\s+"
        r"(?P<unit_price>[0-9,]+(?:\.\d+)?)\s+"
        r"(?P<ccy>[A-Z]{3})/(?P<uom2>[A-Z]{1,4})\s+"
        r"(?P<ext>[0-9,]+\.\d{2})\s+"
        r"(?P<tax>[0-9,]+\.\d{2})\s+"
        r"(?P<total>[0-9,]+\.\d{2})$"
    )
    rows: List[Dict[str, Any]] = []
    for ln in lines:
        s = " ".join(ln.split())
        m = row_pat.match(s)
        if not m:
            continue
        rows.append({
            "DESCRIPTION": m.group("desc").strip(),
            "QUANTITY": safe_float0(m.group("qty")),
            "UOM": m.group("uom"),
            "UNIT PRICE": safe_float0(m.group("unit_price")),
            "PRICE BASIS": f"{m.group('ccy')}/{m.group('uom2')}",
            "EXTENDED AMOUNT": safe_float0(m.group("ext")),
            "TAX AMOUNT": safe_float0(m.group("tax")),
            "AMOUNT": safe_float0(m.group("total")),
        })
    return rows


def _table_noise_filter(lines: List[str]) -> List[str]:
    noise_exact = {
        "DESCRIPTION", "QUANTITY", "UNIT PRICE", "EXTENDED AMOUNT", "TAX AMOUNT", "DESTINATION",
        "MAIL INSTRUCTIONS", "DATE UPLIFTED FUEL TICKET AIRCRAFT TYPE", "FLIGHT NO. TERMS",
        "TAIL NO. LOCATION TERRITORY CONTACT", "INVOICE", "PAGE NO.", "USD USD USD",
        "CUSTOMER NO.", "INVOICE NO.", "INVOICE DATE", "IMPORTO TOTALE", "N. ORDINE DI VENDITA",
        "FECHA DE VENCIMIENTO", "PAYMENT DUE DATE", "INVOICE AMOUNT",
    }
    out = []
    for ln in lines:
        up = ln.upper().strip()
        if up in noise_exact:
            continue
        if up.startswith(("DESCRIPTION/", "QUANTITY/", "UNIT PRICE/", "EXTENDED AMOUNT/", "TAX AMOUNT/",
                          "SALES ORDER NO", "INVOICE AMOUNT/", "PAYMENT DUE DATE/")):
            continue
        if re.match(r"^\d+\s*-\s*\d+$", ln):
            continue
        # drop seller header if it repeats mid-table
        if "WORLD FUEL SERVICES" in up and "INC" in up:
            continue
        out.append(ln)
    return out


def parse_line_items_column_blocks(text: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    lines = normalize_lines(text)

    # Start after SALES ORDER NO value; fallback after DESCRIPTION header
    start_idx = None
    for i, ln in enumerate(lines):
        if ln.upper().startswith("SALES ORDER NO"):
            for j in range(i + 1, min(i + 12, len(lines))):
                if re.fullmatch(r"\d+", lines[j]):
                    start_idx = j + 1
                    break
            if start_idx is None and i + 1 < len(lines):
                start_idx = i + 1
            break
    if start_idx is None:
        for i, ln in enumerate(lines):
            if ln.upper().startswith("DESCRIPTION"):
                start_idx = i + 1
                break
    if start_idx is None:
        start_idx = 0

    # Stop at 3-number totals line; prefer one that is followed by triplet label
    label = "CUSTOMER NO. INVOICE NO. INVOICE DATE"
    three_num_line = re.compile(r"^[0-9,]+\.\d{2}\s+[0-9,]+\.\d{2}\s+[0-9,]+\.\d{2}$")
    last_three_idx = None
    stop_idx = len(lines)

    for i in range(start_idx, len(lines)):
        if three_num_line.match(lines[i]):
            last_three_idx = i
            lookahead = " ".join(lines[i:i + 60]).upper()
            if label in lookahead:
                stop_idx = i
                break

    if stop_idx == len(lines) and last_three_idx is not None:
        stop_idx = last_three_idx

    tbl = _table_noise_filter(lines[start_idx:stop_idx])

    qty_pat = re.compile(r"^[0-9,]+(?:\.\d+)?(?:\s+[A-Z]{1,4})(?:\s+[0-9,]+(?:\.\d+)?\s+[A-Z]{1,4})?$")
    unit_price_pat = re.compile(r"^[0-9,]+(?:\.\d+)?\s+[A-Z]{3}/[A-Z]{1,4}$") # 25.00000 USD/EA
    amt_pat = re.compile(r"^[0-9,]+\.\d{2}$")

    # Descriptions until first qty
    desc: List[str] = []
    i = 0
    while i < len(tbl) and not qty_pat.match(tbl[i]):
        desc.append(tbl[i])
        i += 1

    # Qty until first unit price
    qty: List[Tuple[float, str]] = []
    while i < len(tbl) and qty_pat.match(tbl[i]):
        q, uom = tbl[i].split()[0:2]
        qty.append((safe_float0(q), uom))
        i += 1

    # Unit prices
    unit_prices: List[Tuple[float, str]] = []
    while i < len(tbl) and unit_price_pat.match(tbl[i]):
        up, basis = tbl[i].split()
        unit_prices.append((safe_float0(up), basis))
        i += 1

    # Amounts
    amounts: List[float] = []
    while i < len(tbl):
        if amt_pat.match(tbl[i]):
            amounts.append(safe_float0(tbl[i]))
        i += 1

    n = min(len(desc), len(qty), len(unit_prices)) if (desc and qty and unit_prices) else 0
    diag = {
        "mode": "column_blocks",
        "desc_count": len(desc),
        "qty_count": len(qty),
        "unit_price_count": len(unit_prices),
        "amount_count": len(amounts),
        "rows_expected": n
    }

    if n == 0:
        return [], diag

    ext = amounts[0:n] if len(amounts) >= n else [0.0] * n
    tax = amounts[n:2 * n] if len(amounts) >= 2 * n else [0.0] * n
    tot = amounts[2 * n:3 * n] if len(amounts) >= 3 * n else ext

    rows: List[Dict[str, Any]] = []
    for idx in range(n):
        rows.append({
            "DESCRIPTION": desc[idx],
            "QUANTITY": qty[idx][0],
            "UOM": qty[idx][1],
            "UNIT PRICE": unit_prices[idx][0],
            "PRICE BASIS": unit_prices[idx][1],
            "EXTENDED AMOUNT": ext[idx] if idx < len(ext) else 0.0,
            "TAX AMOUNT": tax[idx] if idx < len(tax) else 0.0,
            "AMOUNT": tot[idx] if idx < len(tot) else (ext[idx] if idx < len(ext) else 0.0),
        })

    diag["rows"] = len(rows)
    return rows, diag


def parse_colt_line_items(text: str) -> List[Dict[str, Any]]:
    if not _is_colt_invoice(text):
        return []

    lines = normalize_lines(text)
    currency, _ = find_electronic_amount(text)
    date_pat = re.compile(r"^\d{1,2}/\d{1,2}/\d{4}$")
    qty_pat = re.compile(r"^(?P<qty>[0-9,]+(?:\.\d+)?)\s+(?P<uom>[A-Z]{1,4})$")
    num_pat = re.compile(r"^[0-9,]+(?:\.\d+)?$")
    remit_pat = re.compile(r"^(?:" + CURRENCY_CODES + r")\s+[0-9,]+\.\d{2}$", re.IGNORECASE)

    start_idx = 0
    for i, line in enumerate(lines):
        if "PLEASE REMIT THIS AMOUNT" in line.upper():
            for j in range(i + 1, min(i + 8, len(lines))):
                if remit_pat.match(lines[j]):
                    start_idx = j + 1
                    break
            break

    rows: List[Dict[str, Any]] = []
    seen = set()
    i = start_idx
    while i < len(lines):
        up = lines[i].upper()
        if up.startswith("COMMENTS") or up.startswith("VENDOR CODE"):
            break
        if not date_pat.match(lines[i]):
            i += 1
            continue

        i += 1
        while i < len(lines) and re.fullmatch(r"[A-Z0-9]{3,4}", lines[i]) and len(lines[i]) <= 4:
            i += 1

        desc_parts: List[str] = []
        while i < len(lines) and not qty_pat.match(lines[i]):
            up = lines[i].upper()
            if date_pat.match(lines[i]) or up.startswith("COMMENTS") or up.startswith("VENDOR CODE"):
                break
            desc_parts.append(lines[i])
            i += 1

        if i >= len(lines):
            break
        qty_match = qty_pat.match(lines[i])
        if not qty_match or not desc_parts:
            i += 1
            continue

        qty = safe_float0(qty_match.group("qty"))
        uom = qty_match.group("uom")
        i += 1

        nums: List[float] = []
        while i < len(lines) and len(nums) < 5 and num_pat.match(lines[i]):
            nums.append(safe_float0(lines[i]))
            i += 1
        if len(nums) < 2:
            continue

        row = {
            "DESCRIPTION": " ".join(desc_parts).strip(),
            "QUANTITY": qty,
            "UOM": uom,
            "UNIT PRICE": nums[0],
            "PRICE BASIS": f"{currency}/{uom}",
            "EXTENDED AMOUNT": nums[1],
            "TAX AMOUNT": 0.0,
            "AMOUNT": nums[4] if len(nums) >= 5 else nums[1],
        }
        key = (row["DESCRIPTION"], round(row["AMOUNT"], 2))
        if key not in seen:
            seen.add(key)
            rows.append(row)

    return rows


def parse_line_items_dynamic(text: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    colt = parse_colt_line_items(text)
    if colt:
        return colt, {"mode": "colt_service_rows", "rows": len(colt)}
    inline = parse_line_items_inline_rows(text)
    if inline:
        return inline, {"mode": "inline_rows", "rows": len(inline)}
    return parse_line_items_column_blocks(text)


def build_currency_audit(payload: Dict[str, Any], text: str) -> Dict[str, Any]:
    inv = payload["INVOICE"]
    rows = inv.get("LINE ITEMS", {}).get("ROWS", []) or []
    payment = inv.get("PAYMENT", {}).get("ELECTRONIC", {})
    summary_total = inv.get("INVOICE TAX SUMMARY", {}).get("INVOICE TOTAL")
    if summary_total is None:
        summary_total = inv.get("USD USD USD", {}).get("INVOICE TOTAL")

    return {
        "PAYMENT CURRENCY": payment.get("CURRENCY", ""),
        "PAYMENT AMOUNT": safe_float0(payment.get("AMOUNT", 0)),
        "LINE ITEM SUM": round(sum(safe_float0(r.get("AMOUNT", 0)) for r in rows), 2),
        "SUMMARY INVOICE TOTAL": safe_float0(summary_total),
        "CURRENCIES FOUND": find_currencies(text),
        "EXCHANGE RATE": find_exchange_rate(text),
        "ROWS": len(rows),
        "MATCHED TOTAL SOURCE": "",
    }


def reconcile_currency_totals(payload: Dict[str, Any], text: str) -> None:
    """
    Some foreign-currency invoices can show remit and local converted totals.
    Exports should use the same currency as the captured line-item amounts.
    """
    inv = payload["INVOICE"]
    rows = inv.get("LINE ITEMS", {}).get("ROWS", []) or []
    audit = build_currency_audit(payload, text)

    line_sum = safe_float0(audit["LINE ITEM SUM"])
    payment_amt = safe_float0(audit["PAYMENT AMOUNT"])
    summary_amt = safe_float0(audit["SUMMARY INVOICE TOTAL"])

    chosen = None
    source = ""
    if rows and payment_amt and amount_close(line_sum, payment_amt):
        chosen = line_sum
        source = "line_items_match_payment_remit_currency"
    elif rows and summary_amt and amount_close(line_sum, summary_amt):
        chosen = line_sum
        source = "line_items_match_summary"
    elif payment_amt and len(audit["CURRENCIES FOUND"]) > 1:
        chosen = payment_amt
        source = "payment_remit_currency"
    elif summary_amt:
        chosen = summary_amt
        source = "summary"
    elif payment_amt:
        chosen = payment_amt
        source = "payment"

    if chosen is not None:
        inv["INVOICE HEADER"]["INVOICE AMOUNT"] = float(chosen)
        inv["INVOICE HEADER"]["AMOUNT TO BE EFT DRAFTED"] = float(chosen)
        if rows and amount_close(line_sum, chosen):
            inv["USD USD USD"] = {
                "TOTAL EXCLUSIVE OF TAX": round(sum(safe_float0(r.get("EXTENDED AMOUNT", 0)) for r in rows), 2),
                "TOTAL TAX": round(sum(safe_float0(r.get("TAX AMOUNT", 0)) for r in rows), 2),
                "INVOICE TOTAL": round(line_sum, 2),
            }

    audit["MATCHED TOTAL SOURCE"] = source
    audit["CHOSEN INVOICE TOTAL"] = float(chosen or 0.0)
    audit["PAYMENT_MATCHES_LINE_ITEMS"] = bool(rows and payment_amt and amount_close(line_sum, payment_amt))
    audit["SUMMARY_MATCHES_LINE_ITEMS"] = bool(rows and summary_amt and amount_close(line_sum, summary_amt))
    audit["SUMMARY_TOTAL_IS_CONVERTED_LOCAL_CURRENCY"] = bool(
        rows and summary_amt and not amount_close(line_sum, summary_amt) and payment_amt and amount_close(line_sum, payment_amt)
    )
    if audit["SUMMARY_TOTAL_IS_CONVERTED_LOCAL_CURRENCY"]:
        audit["SUMMARY_MATCH_NOTE"] = (
            "False is expected: the printed summary total appears to be the converted local-currency total, "
            "while payment and line items are in the remit currency."
        )
    inv["PARSER_DIAGNOSTICS"]["currency_audit"] = audit


def parse_structured_invoice(filename: str, pdf_bytes: bytes) -> Tuple[Dict[str, Any], str, str, int]:
    pages = pdf_page_count(pdf_bytes)
    text, engine = pdf_to_text(pdf_bytes)

    currency, electronic_amt = find_electronic_amount(text)
    bill_to, seller = parse_bill_to_and_seller(text)
    cust_no, inv_no, inv_date = find_invoice_triplet(text)

    payload: Dict[str, Any] = {
        "INVOICE": {
            "DOCUMENT": {
                "DOCUMENT TYPE": "INVOICE",
                "PAGE NO.": find_page_no(text),
                "PAGES": pages,
                "EXTRACTED AT": now_ts(),
                "SOURCE FILE": filename,
                "EXTRACTION ENGINE": engine,
                "SCHEMA VERSION": "4.0"
            },
            "BILL TO": bill_to,
            "SELLER": seller,
            "INVOICE HEADER": {
                "CUSTOMER NO.": cust_no,
                "INVOICE NO.": inv_no,
                "INVOICE DATE": inv_date,
                "DUE DATE": find_due_date(text),
                "SALES ORDER NO.": find_sales_order_no(text),
                "PO NO./CONTRACT NO.": find_po_no(text) or "N/A",
                "AMOUNT TO BE EFT DRAFTED": 0.0,
                "INVOICE AMOUNT": 0.0
            },
            "PAYMENT": {"ELECTRONIC": {"CURRENCY": currency, "AMOUNT": float(electronic_amt)}},
            "CONTEXT": parse_flight_location(text),
            "LINE ITEMS": {
                "COLUMNS": ["DESCRIPTION", "QUANTITY", "UNIT PRICE", "EXTENDED AMOUNT", "TAX AMOUNT", "AMOUNT"],
                "ROWS": []
            },
            "USD USD USD": parse_summary_totals(text).get("USD USD USD", {}),
            "INVOICE TAX SUMMARY": parse_invoice_tax_summary(text),
            "REMIT TO": parse_remit_and_wire(text)[0],
            "OR WIRE TO": parse_remit_and_wire(text)[1],
            "COMMENTS": parse_comments(text),
            "PARSER_DIAGNOSTICS": {}
        }
    }

    rows, diag = parse_line_items_dynamic(text)
    payload["INVOICE"]["LINE ITEMS"]["ROWS"] = rows
    payload["INVOICE"]["PARSER_DIAGNOSTICS"] = diag
    if not payload["INVOICE"]["LINE ITEMS"]["ROWS"]:
        payload["INVOICE"]["PARSER_DIAGNOSTICS"]["structured_adapter_note"] = "No structured rows detected; universal parser fallback can be used."

    reconcile_currency_totals(payload, text)

    # Safety: if customer number still looks like a decimal fragment,
    # re-run the robust extractor and overwrite.
    try:
        hdr = payload.get("INVOICE", {}).get("INVOICE HEADER", {}) \
              or payload.get("header", {})
        cust_val = str(hdr.get("CUSTOMER NO.", hdr.get("customer_no", "")))
        if not cust_val.isdigit() or len(cust_val) > 8:
            c, i, d = find_invoice_triplet(text)
            if c:
                if "CUSTOMER NO." in hdr:
                    hdr["CUSTOMER NO."] = c
                if "customer_no" in hdr:
                    hdr["customer_no"] = c
    except Exception:
        pass

    return payload, text, engine, pages


def parse_ez_invoice(filename: str, pdf_bytes: bytes, parser_mode: str = "Auto") -> Tuple[Dict[str, Any], str, str, int]:
    """Route invoice PDFs through the best available extraction adapter."""
    pages = pdf_page_count(pdf_bytes)
    text, engine = pdf_to_text(pdf_bytes)
    use_gst_adapter = parser_mode == "GST/e-Invoice adapter" or (
        parser_mode == "Auto" and looks_like_gst_invoice(text)
    )
    if use_gst_adapter:
        payload = parse_gst_invoice(
            filename=filename,
            pdf_bytes=pdf_bytes,
            text=text,
            engine=engine,
            pages=pages,
            extracted_at=now_ts(),
        )
        if payload:
            apply_accounting_route(payload, homes=current_home_company_names())
            return payload, text, engine, pages

    payload, text, engine, pages = parse_structured_invoice(filename, pdf_bytes)
    use_structured_adapter = parser_mode == "Structured adapter" or (
        parser_mode == "Auto" and looks_like_structured_vendor_invoice(text)
    )
    if use_structured_adapter:
        payload["INVOICE"]["DOCUMENT"]["PARSER"] = "Structured Adapter"
        payload["INVOICE"].setdefault("PARSER_DIAGNOSTICS", {})["parser"] = "structured_adapter"
        apply_accounting_route(payload, homes=current_home_company_names())
        return payload, text, engine, pages

    generic_payload = parse_generic_invoice(
        filename=filename,
        text=text,
        engine=engine,
        pages=pages,
        extracted_at=now_ts(),
    )
    apply_accounting_route(generic_payload, homes=current_home_company_names())
    return generic_payload, text, engine, pages


# =========================
# Capture check (what you wanted instead of "risk/anomalies")
# =========================
REQUIRED_LINE_FIELDS = ["DESCRIPTION", "QUANTITY", "UOM", "UNIT PRICE", "AMOUNT"]


def capture_check(payload: Dict[str, Any]) -> Tuple[bool, List[str], pd.DataFrame]:
    """
    Returns:
      ok: bool
      issues: list[str]
      df_missing: dataframe with per-line missing info
    """
    inv = payload["INVOICE"]
    h = inv["INVOICE HEADER"]
    issues: List[str] = []

    # header must-haves
    if not str(h.get("INVOICE NO.", "")).strip():
        issues.append("Missing INVOICE NO.")
    if not str(h.get("INVOICE DATE", "")).strip():
        issues.append("Missing INVOICE DATE.")
    if not str(h.get("DUE DATE", "")).strip():
        issues.append("Missing DUE DATE.")

    rows = inv.get("LINE ITEMS", {}).get("ROWS", []) or []
    if not rows:
        issues.append("No line items captured.")
        return False, issues, pd.DataFrame()

    missing_rows = []
    for i, r in enumerate(rows, start=1):
        missing = []
        for f in REQUIRED_LINE_FIELDS:
            val = r.get(f, None)
            if f in ("QUANTITY", "UNIT PRICE", "AMOUNT"):
                # numeric must exist (can be 0 but should be present)
                if val is None or str(val).strip() == "":
                    missing.append(f)
            else:
                if val is None or str(val).strip() == "":
                    missing.append(f)
        if missing:
            missing_rows.append({
                "Line#": i,
                "DESCRIPTION": r.get("DESCRIPTION", ""),
                "Missing Fields": ", ".join(missing)
            })

    df_missing = pd.DataFrame(missing_rows)

    # sum reconciliation (fixed tolerance, no UI slider)
    inv_total = safe_float0(h.get("INVOICE AMOUNT", 0))
    line_sum = sum(safe_float0(r.get("AMOUNT", 0)) for r in rows)
    if inv_total and abs(line_sum - inv_total) > 0.50:
        issues.append(f"Possible missing/extra line(s): line sum {line_sum:.2f} != invoice total {inv_total:.2f} (Δ {(line_sum-inv_total):.2f})")

    if not df_missing.empty:
        issues.append("Some line items have missing required fields.")

    ok = len(issues) == 0
    return ok, issues, df_missing


# =========================
# Converters (selector + ZIPs + per-invoice tabs)
# =========================
def to_quickbooks(payload: Dict[str, Any]) -> Dict[str, Any]:
    inv = payload["INVOICE"]
    h = inv["INVOICE HEADER"]
    cur = inv.get("PAYMENT", {}).get("ELECTRONIC", {}).get("CURRENCY", "USD")

    bill = {
        "Bill": {
            "VendorRef": {"name": inv["SELLER"]["NAME"], "value": "WORLD_FUEL_SERVICES"},
            "CustomerRef": {"name": inv["BILL TO"]["NAME"]},
            "TxnDate": h.get("INVOICE DATE", ""),
            "DueDate": h.get("DUE DATE", ""),
            "DocNumber": h.get("INVOICE NO.", ""),
            "PrivateNote": f"Sales Order: {h.get('SALES ORDER NO.', '')} | PO: {h.get('PO NO./CONTRACT NO.', '')}",
            "TotalAmt": float(h.get("INVOICE AMOUNT", 0) or 0),
            "CurrencyRef": {"value": cur},
            "Line": []
        }
    }

    for i, r in enumerate(inv.get("LINE ITEMS", {}).get("ROWS", []) or []):
        bill["Bill"]["Line"].append({
            "Id": str(i + 1),
            "LineNum": i + 1,
            "Description": r.get("DESCRIPTION", ""),
            "Amount": float(r.get("AMOUNT", 0) or 0),
            "DetailType": "AccountBasedExpenseLineDetail",
            "AccountBasedExpenseLineDetail": {
                "AccountRef": {"name": "General Expenses", "value": "1"},
                "Qty": float(r.get("QUANTITY", 0) or 0),
                "UnitPrice": float(r.get("UNIT PRICE", 0) or 0),
                "TaxCodeRef": {"value": "TAX" if float(r.get("TAX AMOUNT", 0) or 0) > 0 else "NON"}
            }
        })

    return bill


def to_coupa(payload: Dict[str, Any]) -> Dict[str, Any]:
    inv = payload["INVOICE"]
    h = inv["INVOICE HEADER"]
    cur = inv.get("PAYMENT", {}).get("ELECTRONIC", {}).get("CURRENCY", "USD")
    return {
        "invoice-header": {
            "invoice-number": h.get("INVOICE NO.", ""),
            "invoice-date": h.get("INVOICE DATE", ""),
            "due-date": h.get("DUE DATE", ""),
            "supplier": {"name": inv["SELLER"]["NAME"], "number": h.get("CUSTOMER NO.", "")},
            "bill-to": {"name": inv["BILL TO"]["NAME"], "address": ", ".join(inv["BILL TO"].get("ADDRESS", []))},
            "currency": {"code": cur},
            "total": float(h.get("INVOICE AMOUNT", 0) or 0),
            "tax-amount": float(inv.get("USD USD USD", {}).get("TOTAL TAX", 0) or 0),
            "payment-term": inv.get("CONTEXT", {}).get("TERMS", "")
        },
        "invoice-lines": [
            {
                "line-num": i + 1,
                "description": r.get("DESCRIPTION", ""),
                "quantity": r.get("QUANTITY", 0),
                "uom": {"code": r.get("UOM", "")},
                "price": r.get("UNIT PRICE", 0),
                "tax-amount": r.get("TAX AMOUNT", 0),
                "total": r.get("AMOUNT", 0),
            }
            for i, r in enumerate(inv.get("LINE ITEMS", {}).get("ROWS", []) or [])
        ]
    }


def to_netsuite(payload: Dict[str, Any]) -> Dict[str, Any]:
    inv = payload["INVOICE"]
    h = inv["INVOICE HEADER"]
    cur = inv.get("PAYMENT", {}).get("ELECTRONIC", {}).get("CURRENCY", "USD")
    return {
        "recordType": "vendorBill",
        "entity": {"name": inv["SELLER"]["NAME"]},
        "tranId": h.get("INVOICE NO.", ""),
        "tranDate": h.get("INVOICE DATE", ""),
        "dueDate": h.get("DUE DATE", ""),
        "currency": {"name": cur},
        "total": float(h.get("INVOICE AMOUNT", 0) or 0),
        "taxTotal": float(inv.get("USD USD USD", {}).get("TOTAL TAX", 0) or 0),
        "memo": f"Sales Order {h.get('SALES ORDER NO.', '')}",
        "item": [
            {
                "item": {"name": r.get("DESCRIPTION", "")},
                "quantity": r.get("QUANTITY", 0),
                "rate": r.get("UNIT PRICE", 0),
                "amount": r.get("AMOUNT", 0),
                "taxAmount": r.get("TAX AMOUNT", 0)
            }
            for r in inv.get("LINE ITEMS", {}).get("ROWS", []) or []
        ]
    }


def to_sap(payload: Dict[str, Any]) -> Dict[str, Any]:
    inv = payload["INVOICE"]
    h = inv["INVOICE HEADER"]
    cur = inv.get("PAYMENT", {}).get("ELECTRONIC", {}).get("CURRENCY", "USD")
    return {
        "DocumentType": "VENDOR_INVOICE",
        "CardCode": h.get("CUSTOMER NO.", ""),
        "CardName": inv["SELLER"]["NAME"],
        "DocNum": h.get("INVOICE NO.", ""),
        "DocDate": h.get("INVOICE DATE", ""),
        "DocDueDate": h.get("DUE DATE", ""),
        "DocCurrency": cur,
        "DocTotal": float(h.get("INVOICE AMOUNT", 0) or 0),
        "VatSum": float(inv.get("USD USD USD", {}).get("TOTAL TAX", 0) or 0),
        "Comments": f"Sales Order: {h.get('SALES ORDER NO.', '')}",
        "DocumentLines": [
            {
                "LineNum": i,
                "ItemDescription": r.get("DESCRIPTION", ""),
                "Quantity": r.get("QUANTITY", 0),
                "UnitPrice": r.get("UNIT PRICE", 0),
                "LineTotal": r.get("AMOUNT", 0),
                "TaxTotal": r.get("TAX AMOUNT", 0)
            }
            for i, r in enumerate(inv.get("LINE ITEMS", {}).get("ROWS", []) or [])
        ]
    }


def to_tally(payload: Dict[str, Any]) -> Dict[str, Any]:
    return {"format": "TallyPrime XML", "voucher_xml": build_tally_xml(payload, classifier=classifier)}


CONVERTERS = {
    "QuickBooks": to_quickbooks,
    "Tally": to_tally,
    "Zoho Books": build_zoho_export,
    "Coupa": to_coupa,
    "NetSuite": to_netsuite,
    "SAP": to_sap,
}


def tally_post_blocker(payload: Dict[str, Any]) -> str:
    route = payload.get("INVOICE", {}).get("ROUTING", {})
    direction = route.get("DIRECTION", "unknown")
    if direction in ("unknown", "", None):
        return ""
    if direction != "inbound":
        return "Tally posting is enabled for inbound purchase invoices in this prototype. Detected: " + direction + "."
    return ""


def tally_readiness_message(data: Dict[str, Any]) -> str:
    if not data.get("ok"):
        return "Needs review before posting."
    blocker = tally_post_blocker(data.get("payload", {}))
    return blocker or "Ready for Tally."


def _store_tally_status(fname: str, result: Dict[str, Any], dry_run: bool = False) -> None:
    if "tally_post_status" not in st.session_state:
        st.session_state["tally_post_status"] = {}
    status = {
        "success": bool(result.get("success")),
        "message": result.get("message", ""),
        "response": result.get("response", {}),
        "mode": "Dry Run" if dry_run else "Live Post",
        "at": now_ts(),
    }
    st.session_state["tally_post_status"][fname] = status
    if fname in st.session_state.get("payloads", {}):
        st.session_state["payloads"][fname]["tally_status"] = status


def make_zip(payloads: Dict[str, Dict[str, Any]], converter, label: str) -> io.BytesIO:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for fname, data in payloads.items():
            base = fname.rsplit(".", 1)[0]
            zf.writestr(f"{base}_parsed.json", json.dumps(data["payload"], indent=2))
            converted = converter(data["payload"])
            zf.writestr(f"{base}_{label}.json", json.dumps(converted, indent=2))
    buf.seek(0)
    return buf

def make_einvoice_excel(payloads, classifier=None):
    """Generate a normalized EZ-Invoice accounting workbook."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EZ-Invoice"

    hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    hdr_fill = PatternFill(start_color="0e7490", end_color="0e7490", fill_type="solid")
    data_font = Font(name="Arial", size=10)
    money_fmt = '#,##0.00'
    bdr = Border(
        left=Side(style="thin", color="D0D0D0"), right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"), bottom=Side(style="thin", color="D0D0D0"))

    headers = [
        "INVOICE NUMBER", "INVOICE DATE", "CUSTOMER NUMBER", "CUSTOMER NAME",
        "PAYMENT / CARD NUMBER", "LOCATION / SITE", "SUPPLIER NAME",
        "ASSET / REFERENCE ID", "ASSET / SERVICE TYPE", "REFERENCE TICKET",
        "PROJECT / JOB", "SERVICE / ACTIVITY",
        "SERVICE DATE", "LOCATION CODE", "LOCATION ID", "DESTINATION / DELIVERY LOCATION",
        "TRANSACTION COUNTRY", "STATE / REGION",
        "SALES ORDER NO", "PO CONTRACT NO",
        "ITEM", "ITEM DESC", "PLATFORM CATEGORY",
        "CUSTOMER CATEGORY", "GL CODE",
        "QB CATEGORY", "QB ACCOUNT", "QB ACCOUNT ID", "QB RULE", "QB MAPPING SOURCE",
        "SPEND BILLING", "INVOICE CURRENCY",
        "EXTENDED AMOUNT", "LOCAL CURRENCY",
        "TAX AMOUNT", "QTY INV", "VOLUME", "UOM", "UNIT COST"
    ]

    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = bdr

    row_num = 2
    customer_name = ""

    def extract_customer_name(inv, raw_text):
        name = inv.get("BILL TO", {}).get("NAME", "")
        bad = ["REMIT TO:", "REMIT TO", "REMIT TO/RIMETTERE A:", "REMIT TO/REMETTRE A:", ""]
        if name in bad or "REMIT" in name.upper():
            lines = [l.strip() for l in (raw_text or "").splitlines() if l.strip()]
            skip = {"WORLD FUEL SERVICES", "HSBC", "BANK OF AMERICA", "REMIT TO", "OR WIRE TO"}
            for i, l in enumerate(lines):
                up = l.upper().strip()
                if any(s in up for s in skip):
                    continue
                if up.startswith("TEL:") or up.startswith("EMAIL:") or up.startswith("INTERNET:"):
                    continue
                if up.startswith("9800 N.W.") or up.startswith("MIAMI"):
                    continue
                if up.startswith("INVOICE") or up.startswith("CUSTOMER NO"):
                    continue
                if up == up.upper() and len(up.split()) >= 2 and not up[0].isdigit():
                    if any(kw in up for kw in ["INC", "LTD", "LLC", "AVIATION", "JET", "SERVICE", "MANAGEMENT", "CHARTER", "FLIGHT", "S.A.", "S.R.L"]):
                        if "WORLD FUEL" not in up and "BANK" not in up:
                            return l.strip()
            return name
        return name

    def extract_colt_fields(raw_text):
        """Extract TRIP, FLIGHT, TAIL, AIRCRAFT, SERVICE DATE, DEST from COLT format."""
        result = {"trip": "", "flight": "", "tail": "", "aircraft": "",
                  "service_dates": [], "destinations": [], "country": "", "terms": ""}
        if not raw_text or "TRIP NO." not in raw_text:
            return result
        lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
        for i, l in enumerate(lines):
            if l.startswith("TRIP NO.") and "FLIGHT NO." in l and "TAIL NO." in l:
                if i + 1 < len(lines):
                    val_line = lines[i + 1]
                    parts = val_line.split()
                    if len(parts) >= 1:
                        result["trip"] = parts[0]
                    # Find tail: look for pattern like N993XZ or C-GLVK
                    for p in parts:
                        if re.match(r'^[A-Z]-[A-Z]{3,4}$', p) or re.match(r'^N\d{2,}[A-Z]*$', p) or re.match(r'^[A-Z]{2}-[A-Z]{3}$', p):
                            result["tail"] = p
                            break
                    # Aircraft type: multi-word like "Dassault Falcon 7X" or "Challenger 650"
                    aircraft_patterns = ["Dassault", "Challenger", "Gulfstream", "Bombardier", "Cessna", "Boeing", "Embraer", "Pilatus", "Hawker", "Learjet"]
                    for ap in aircraft_patterns:
                        if ap in val_line:
                            idx = val_line.index(ap)
                            rest = val_line[idx:]
                            # Take until next known keyword
                            for stop in ["Service Invoice", "3rd Party", "Contract", "Retail"]:
                                if stop in rest:
                                    rest = rest[:rest.index(stop)].strip()
                                    break
                            result["aircraft"] = rest.strip()
                            break
                    # Terms
                    if "D FROM INV" in val_line:
                        m = re.search(r'(\d+\s+D FROM INV)', val_line)
                        if m:
                            result["terms"] = m.group(1)
                    elif "NET" in val_line:
                        m = re.search(r'(\d+\s+NET[^\s]*)', val_line)
                        if m:
                            result["terms"] = m.group(1)
                break
        # Service dates and destinations from line items
        date_pat = re.compile(r'^(\d{2}/\d{2}/\d{4})\s+(\w*)\s+(\w+)\s+')
        for l in lines:
            m = date_pat.match(l)
            if m:
                result["service_dates"].append(m.group(1))
                dest = m.group(3) if m.group(2) == "" else m.group(2)
                # Check: if group(2) is ORIG and group(3) is DEST
                if m.group(2) and m.group(3):
                    result["destinations"].append(m.group(3))
                elif m.group(2):
                    result["destinations"].append(m.group(2))
        # Country from COUNTRY line
        m_c = re.search(r'\n([A-Z][A-Z\s]+?)\s+[A-Z]{2}-\d{5}', raw_text)
        if m_c:
            result["country"] = m_c.group(1).strip()
        return result

    for fname, data in payloads.items():
        inv = data["payload"]["INVOICE"]
        h = inv["INVOICE HEADER"]
        fl = inv.get("CONTEXT", {})
        com = inv.get("COMMENTS", {})
        cur = inv.get("PAYMENT", {}).get("ELECTRONIC", {}).get("CURRENCY", "USD")
        rows = inv.get("LINE ITEMS", {}).get("ROWS", []) or []
        raw_text = data.get("text", "")

        inv_no = h.get("INVOICE NO.", "")
        inv_date = h.get("INVOICE DATE", "")
        cust_no = h.get("CUSTOMER NO.", "")
        cust_name = extract_customer_name(inv, raw_text)
        if not customer_name and cust_name:
            customer_name = cust_name

        # Check if COLT format
        is_colt = "TRIP NO." in (raw_text or "") and "SERVICE DATE" in (raw_text or "")
        colt = extract_colt_fields(raw_text) if is_colt else {}

        # Card number
        card = ""
        for k, v in com.items():
            if v and "card" in k.lower():
                card = v
                break

        fbo = com.get("FBO Name", "")

        # Supplier
        supplier = ""
        for k, v in com.items():
            kl = k.lower()
            if v and ("supplier" in kl):
                supplier = v
                break
        if not supplier and raw_text:
            m = re.search(r"VENDOR:\s*([^\n]+)", raw_text)
            if m and "CODE" not in m.group(1).upper() and "INV" not in m.group(1).upper():
                supplier = m.group(1).strip()

        # Tail, aircraft, fuel ticket - use JSON first, fallback to COLT extraction
        tail = fl.get("TAIL NO.", "").replace("ANY A/C/", "").strip()
        if "/" in tail:
            tail = tail.split("/")[0].strip()
        if not tail and colt:
            tail = colt.get("tail", "")

        aircraft = fl.get("AIRCRAFT TYPE", "")
        if not aircraft and colt:
            aircraft = colt.get("aircraft", "")

        fuel_ticket = fl.get("FUEL TICKET", "")

        trip_no = (colt.get("trip", "") if colt else "") or fl.get("TRIP NO.", "")
        flight_no = fl.get("FLIGHT NO.", "")

        # Location -> IATA + ICAO
        loc = fl.get("LOCATION", "")
        iata = ""
        icao = ""
        if "/" in loc:
            parts = [p.strip() for p in loc.split("/")]
            for p in parts:
                if len(p) == 3 and p.isalpha():
                    iata = p
                elif len(p) == 4 and p.isalpha():
                    icao = p
        else:
            p = loc.strip()
            if len(p) == 3 and p.isalpha():
                iata = p
            elif len(p) == 4 and p.isalpha():
                icao = p

        # Destination
        dest = fl.get("DESTINATION", "")
        if dest in ("N/A", "NA00", ""):
            dest = ""
        # For COLT: get destination from line items
        if not dest and colt and colt.get("destinations"):
            dest = colt["destinations"][0]

        # Date uplifted - use JSON first, fallback to COLT service dates
        date_uplifted = fl.get("DATE UPLIFTED", "")
        if not date_uplifted and colt and colt.get("service_dates"):
            date_uplifted = colt["service_dates"][0]

        # Territory = STATE, TRX_COUNTRY
        territory = fl.get("TERRITORY", "")
        state = territory
        trx_country = ""
        country_map = {
            "FLORIDA": ("US", "FL"), "PENNSYLVANIA": ("US", "PA"),
            "SOUTH CAROLINA": ("US", "SC"), "NEW JERSEY": ("US", "NJ"),
            "VIRGINIA": ("US", "VA"), "OHIO": ("US", "OH"),
            "TEXAS": ("US", "TX"), "CALIFORNIA": ("US", "CA"),
            "VIRGIN ISLANDS, U.S.": ("US", "VI"),
            "ONTARIO": ("CA", "ON"), "QUEBEC": ("CA", "QC"),
            "ALBERTA": ("CA", "AB"), "BRITISH COLUMBIA": ("CA", "BC"),
            "IRELAND": ("IE", "IE"), "GREECE": ("GR", "GR"),
            "ITALY": ("IT", "IT"), "FRANCE": ("FR", "FR"),
            "UNITED KINGDOM": ("GB", "GB"), "ROMANIA": ("RO", "RO"),
            "JAPAN": ("JP", "JP"), "SOUTH AFRICA": ("ZA", "ZA"),
            "ARGENTINA": ("AR", "AR"), "DOMINICAN REPUBLIC": ("DO", "DO"),
            "UNITED STATES": ("US", ""),
        }
        if territory.upper() in country_map:
            trx_country, state = country_map[territory.upper()]
        # COLT: get country from extracted text
        if not trx_country and colt and colt.get("country"):
            c_name = colt["country"].upper()
            for k, v in country_map.items():
                if k in c_name:
                    trx_country = v[0]
                    if not state or state == territory:
                        state = v[1]
                    break

        sales_order = h.get("SALES ORDER NO.", "")
        po_contract = h.get("PO NO./CONTRACT NO.", "")
        if po_contract == "N/A":
            po_contract = ""

        # Classify line items
        if classifier:
            enriched_rows = classifier.classify_invoice_rows(rows)
        else:
            enriched_rows = rows

        for r in enriched_rows:
            desc = r.get("DESCRIPTION", "")
            desc_clean = re.sub(r"\s*\[.*?\]\s*$", "", desc).strip()
            qty = r.get("QUANTITY", 0)
            uom = r.get("UOM", "")
            unit_price = r.get("UNIT PRICE", 0)
            ext_amt = r.get("EXTENDED AMOUNT", 0)
            tax_amt = r.get("TAX AMOUNT", 0)
            amount = r.get("AMOUNT", 0)
            item_code = r.get("ITEM", "")
            volume = qty if uom not in ("EA",) else 0

            has_worksheet = bool(classifier and classifier.client_map)
            subcategory = r.get("SUBCATEGORY", "")
            platform_cat = (
                r.get("PLATFORM_CATEGORY", "")
                or "Expense Item"
            )
            client_cat_name = r.get("CLIENT_CATEGORY", "") if has_worksheet else ""
            gl_code = r.get("GL_CODE", "") if has_worksheet else ""
            qb_info = resolve_line_account_for_export(data["payload"], r, classifier) if is_connected() else {}

            price_basis = r.get("PRICE BASIS", "")
            local_cur = price_basis.split("/")[0] if "/" in price_basis else cur

            vals = [
                inv_no, inv_date, cust_no, cust_name,
                card, fbo, supplier,
                tail, aircraft, fuel_ticket,
                trip_no, flight_no,
                date_uplifted, iata, icao, dest, trx_country, state,
                sales_order, po_contract,
                item_code, desc_clean, platform_cat,
                client_cat_name, gl_code,
                qb_info.get("QB_CATEGORY", ""),
                qb_info.get("QB_ACCOUNT", ""),
                qb_info.get("QB_ACCOUNT_ID", ""),
                qb_info.get("QB_RULE", ""),
                qb_info.get("QB_MAPPING_SOURCE", ""),
                float(amount or 0), cur,
                float(ext_amt or 0), local_cur,
                float(tax_amt or 0),
                float(qty or 0), float(volume or 0), uom,
                float(unit_price or 0),
            ]

            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=row_num, column=c, value=v)
                cell.font = data_font
                cell.border = bdr
                if c in (31, 33, 35, 39):
                    cell.number_format = money_fmt
            row_num += 1

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.column_dimensions["D"].width = 35
    ws.column_dimensions["V"].width = 35
    ws.column_dimensions["X"].width = 24
    ws.column_dimensions["Y"].width = 16
    ws.column_dimensions["Z"].width = 24
    ws.column_dimensions["AA"].width = 28
    ws.column_dimensions["AB"].width = 18
    ws.column_dimensions["AC"].width = 28
    ws.column_dimensions["AD"].width = 24
    ws.column_dimensions["AE"].width = 18
    ws.column_dimensions["F"].width = 30
    ws.column_dimensions["G"].width = 30
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    clean_name = re.sub(r'[^A-Za-z0-9\s]', '', customer_name).strip().replace(' ', '_').upper()
    if not clean_name:
        clean_name = "UNKNOWN_CLIENT"
    date_str = datetime.now().strftime("%Y-%m-%d")
    filename = f"EZ_Invoice_to_{clean_name}__{date_str}.xlsx"

    return buf, filename


# =========================
# Persistent History + Analytics
# =========================
HISTORY_DIR = APP_DIR / "data"
HISTORY_FILE = HISTORY_DIR / "invoice_history.json"


def _history_load() -> List[Dict[str, Any]]:
    try:
        if HISTORY_FILE.exists():
            with HISTORY_FILE.open("r", encoding="utf-8") as f:
                data = json.load(f)
            return data if isinstance(data, list) else []
    except Exception:
        pass
    return []


def _history_save(records: List[Dict[str, Any]]) -> None:
    try:
        HISTORY_DIR.mkdir(exist_ok=True)
        with HISTORY_FILE.open("w", encoding="utf-8") as f:
            json.dump(records[-1000:], f, indent=2)
    except Exception:
        pass


def _history_record(fname: str, data: Dict[str, Any], ms: float) -> Dict[str, Any]:
    payload = data["payload"]
    inv = payload.get("INVOICE", {})
    h = inv.get("INVOICE HEADER", {})
    rows = inv.get("LINE ITEMS", {}).get("ROWS", []) or []
    cur = inv.get("PAYMENT", {}).get("ELECTRONIC", {}).get("CURRENCY", "USD")
    supplier = inv.get("SELLER", {}).get("NAME", "")
    customer = inv.get("BILL TO", {}).get("NAME", "")
    route = inv.get("ROUTING", {})
    total = safe_float0(h.get("INVOICE AMOUNT", 0))
    tax = safe_float0(inv.get("USD USD USD", {}).get("TOTAL TAX", 0))
    status_label = _api_status_label(data.get("api_status", ""), fallback_ready=bool(data.get("ok")))
    return {
        "Processed At": now_ts(),
        "File": fname,
        "Invoice #": h.get("INVOICE NO.", ""),
        "Supplier": supplier,
        "Customer": customer,
        "Direction": route.get("DIRECTION", "unknown"),
        "Transaction Type": route.get("TRANSACTION TYPE", "review"),
        "Accounting Party": route.get("PARTY NAME", supplier or customer),
        "Invoice Date": h.get("INVOICE DATE", ""),
        "Due Date": h.get("DUE DATE", ""),
        "Currency": cur,
        "Amount": total,
        "Tax": tax,
        "Items": len(rows),
        "Status": status_label,
        "Tally Status": (data.get("tally_status") or {}).get("message", ""),
        "Parser": inv.get("DOCUMENT", {}).get("PARSER", inv.get("PARSER_DIAGNOSTICS", {}).get("parser", "Auto")),
        "Engine": data.get("engine", ""),
        "Pages": data.get("pages", 0),
        "Speed ms": round(float(ms or 0), 1),
        "Issues": " | ".join(data.get("issues", []) or []),
    }


def _history_append(record: Dict[str, Any]) -> None:
    records = _history_load()
    records.append(record)
    _history_save(records)


def history_df() -> pd.DataFrame:
    records = _history_load()
    return pd.DataFrame(records) if records else pd.DataFrame()


def session_records_df() -> pd.DataFrame:
    rows = []
    for fname, data in st.session_state.get("payloads", {}).items():
        rows.append(_history_record(fname, data, st.session_state.get("ptimes", {}).get(fname, 0)))
    return pd.DataFrame(rows)


def combined_activity_df() -> pd.DataFrame:
    # Demo/client sessions should reflect the current run. Persisted history is
    # still written for future audit work, but it must not make a fresh launch
    # look like invoices were already processed.
    df = session_records_df()
    if "Processed At" in df.columns:
        df = df.sort_values("Processed At", ascending=False)
    return df


def _line_items_history_df() -> pd.DataFrame:
    records = []
    for fname, data in st.session_state.get("payloads", {}).items():
        payload = data.get("payload", {})
        inv = payload.get("INVOICE", {})
        h = inv.get("INVOICE HEADER", {})
        supplier = inv.get("SELLER", {}).get("NAME", "")
        cur = inv.get("PAYMENT", {}).get("ELECTRONIC", {}).get("CURRENCY", "USD")
        for row in inv.get("LINE ITEMS", {}).get("ROWS", []) or []:
            enriched = classifier.get_gl_info(row.get("DESCRIPTION", "")) if classifier else {}
            records.append({
                "Invoice #": h.get("INVOICE NO.", ""),
                "Supplier": supplier,
                "Description": row.get("DESCRIPTION", ""),
                "Category": enriched.get("subcategory", row.get("PLATFORM_CATEGORY", "Unmapped")),
                "Currency": cur,
                "Amount": safe_float0(row.get("AMOUNT", 0)),
            })
    return pd.DataFrame(records)


def _money_label(value: Any) -> str:
    return "$" + money(value)


def _currency_money_label(value: Any, currency: Any = "") -> str:
    code = str(currency or "").upper().strip()
    symbols = {"USD": "$", "CAD": "C$", "AUD": "A$", "EUR": "€", "GBP": "£"}
    if code in symbols:
        return symbols[code] + money(value)
    if code:
        return f"{code} {money(value)}"
    return money(value)


def _parse_display_date(value: Any) -> Optional[datetime]:
    raw = str(value or "").strip()
    if not raw:
        return None
    candidates = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%d-%b-%Y",
        "%d %b %Y",
        "%d %B %Y",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%m/%d/%Y",
    ]
    for fmt in candidates:
        try:
            return datetime.strptime(raw[:19], fmt)
        except Exception:
            pass
    return None


def _display_date(value: Any) -> str:
    dt = _parse_display_date(value)
    if dt:
        return dt.strftime("%d %b %Y")
    return str(value or "")


def _compact_amount(value: Any, currency: Any = "") -> str:
    amount = safe_float0(value)
    code = str(currency or "").upper().strip()
    symbols = {"USD": "$", "CAD": "C$", "AUD": "A$", "EUR": "€", "GBP": "£"}
    prefix = symbols.get(code, f"{code} " if code else "")
    abs_amount = abs(amount)
    if abs_amount >= 1_000_000_000:
        number = f"{amount / 1_000_000_000:.2f}B"
    elif abs_amount >= 1_000_000:
        number = f"{amount / 1_000_000:.2f}M"
    elif abs_amount >= 100_000:
        number = f"{amount / 1_000:.0f}K"
    else:
        label = _currency_money_label(value, currency)
        if "." in label:
            head, tail = label.rsplit(".", 1)
            if tail == "00":
                return head
        return label
    return f"{prefix}{number}"


def _table_status_chip(status: Any) -> str:
    text = str(status or "Ready")
    klass = _status_class(text)
    return f'<span class="status-pill {klass}">{_ui_escape(text)}</span>'


def _page_heading(title: str, subtitle: str) -> None:
    st.markdown(
        f"""<div class="page-card">
        <span class="eyebrow">{title}</span>
        <h2>{subtitle}</h2>
        </div>""",
        unsafe_allow_html=True,
    )


def render_command_center_page() -> None:
    df = combined_activity_df()
    session_df = session_records_df()
    total_docs = len(df)
    ready = int((df.get("Status", pd.Series(dtype=str)) == "Ready").sum()) if not df.empty else 0
    exceptions = max(total_docs - ready, 0)
    amount = float(df.get("Amount", pd.Series(dtype=float)).sum()) if not df.empty else 0.0
    posted_systems = 2 if (is_connected() or tally_is_connected()) else 0
    st.markdown(f"""<div class="kpi-row">
    <div class="kpi"><div class="n">{total_docs}</div><div class="l">Documents in Memory</div></div>
    <div class="kpi"><div class="n g">{ready}</div><div class="l">Ready to Post</div></div>
    <div class="kpi"><div class="n a">{exceptions}</div><div class="l">Needs Review</div></div>
    <div class="kpi"><div class="n p">{_money_label(amount)}</div><div class="l">Captured Spend</div></div>
    <div class="kpi"><div class="n">{posted_systems}/5</div><div class="l">Live Connectors</div></div>
    </div>""", unsafe_allow_html=True)

    st.markdown("""<div class="capability-grid">
    <div class="capability"><strong>AP Copilot Queue</strong><span>Review low-confidence invoices, duplicate documents, and total mismatches before they hit the ledger.</span></div>
    <div class="capability"><strong>Universal Parser Mesh</strong><span>Use generic extraction plus structured adapters as new supplier formats appear.</span></div>
    <div class="capability"><strong>Posting Control Tower</strong><span>Route the same normalized invoice into QuickBooks, Tally, ERP exports, or Excel audit packs.</span></div>
    <div class="capability"><strong>Learning Layer</strong><span>Client GL mapping and rules become reusable automation memory for future invoices.</span></div>
    </div>""", unsafe_allow_html=True)

    left, right = st.columns([1.15, 0.85])
    with left:
        st.markdown("#### Recent Activity")
        if df.empty:
            st.info("No invoice history yet. Go to Process and upload PDFs to start the activity feed.")
        else:
            cols = [c for c in ["Processed At", "Invoice #", "Direction", "Supplier", "Amount", "Currency", "Status", "Parser"] if c in df.columns]
            st.dataframe(df[cols].head(12), width="stretch", hide_index=True)
    with right:
        st.markdown("#### Automation Lanes")
        lane_html = """
        <div class="timeline-card">
          <div><strong>1. Intake</strong><span>PDF capture and extraction</span></div>
          <div><strong>2. Normalize</strong><span>Common AP invoice schema</span></div>
          <div><strong>3. Validate</strong><span>Totals, required fields, duplicate checks</span></div>
          <div><strong>4. Route</strong><span>GL mapping and connector posting</span></div>
        </div>
        """
        st.markdown(lane_html, unsafe_allow_html=True)
        if not session_df.empty:
            st.caption(f"This session has {len(session_df)} processed invoice(s).")


def render_history_page() -> None:
    df = combined_activity_df()
    query = str(st.session_state.get("history_search_topbar", "") or "").lower().strip()
    if df.empty:
        st.markdown(
            """<div class="history-workspace">
            <div class="history-panel">
              <h3>Audit trail</h3>
              <div class="activity-title">No invoice activity found.</div>
              <div class="activity-meta">Upload invoices to build a searchable trail of extraction, validation, mapping, export, and posting evidence.</div>
            </div>
            </div>""",
            unsafe_allow_html=True,
        )
        return

    view = df.copy()
    if query:
        view = view[view.apply(lambda row: query in " ".join(str(v).lower() for v in row.values), axis=1)]
    if "Processed At" in view.columns:
        view = view.sort_values("Processed At", ascending=False)
    if view.empty:
        st.markdown(
            """<div class="history-workspace">
            <div class="history-panel">
              <h3>Search result</h3>
              <div class="activity-title">No matching invoice activity.</div>
              <div class="activity-meta">Try a different invoice number, vendor, amount, status, or file name.</div>
            </div>
            </div>""",
            unsafe_allow_html=True,
        )
        return

    amounts = pd.to_numeric(view.get("Amount", pd.Series(dtype=float)), errors="coerce").fillna(0)
    status_series = view.get("Status", pd.Series(dtype=str)).astype(str)
    ready_n = int(status_series.isin(["Ready", "Sent"]).sum())
    sent_n = int(status_series.str.contains("sent", case=False, na=False).sum())
    review_n = int((~status_series.isin(["Ready", "Sent"])).sum())
    posted_n = int(view.get("Tally Status", pd.Series(dtype=str)).astype(str).str.len().gt(0).sum()) + sent_n
    vendors = int(view.get("Supplier", pd.Series(dtype=str)).replace("", "Unknown").nunique())
    currencies = [c for c in view.get("Currency", pd.Series(dtype=str)).dropna().astype(str).unique() if c]
    currency_label = currencies[0] if len(currencies) == 1 else ""
    mapping_state = "Locked" if classifier and classifier.client_map else "Not configured"

    def history_kpi(label: str, value: Any, note: str = "", klass: str = "") -> str:
        return (
            '<div class="analytics-card">'
            f'<h3>{_ui_escape(label)}</h3>'
            f'<div class="metric-big {_ui_escape(klass)}">{_ui_escape(str(value))}</div>'
            f'<div class="metric-note">{_ui_escape(note)}</div>'
            '</div>'
        )

    kpis = "".join([
        history_kpi("Invoices", len(view), f"{ready_n} validated", "dark"),
        history_kpi("Ready for ERP", max(ready_n - posted_n, 0), f"{posted_n} already posted", ""),
        history_kpi("Total spend", _compact_amount(amounts.sum(), currency_label), f"{vendors} vendor{'s' if vendors != 1 else ''}", "dark"),
        history_kpi("Exceptions", review_n, "Needs review" if review_n else "Clear", "amber"),
    ])

    stages = [
        ("1", "Extracted", len(view), "PDF captured and invoice fields normalized"),
        ("2", "Validated", ready_n, "Totals, dates, and line items checked"),
        ("3", "Mapped", ready_n if classifier and classifier.client_map else 0, f"Client ledger worksheet: {mapping_state}"),
        ("4", "Posted / exported", posted_n, "Evidence ready for audit and close"),
    ]
    stage_html = "".join(
        '<div class="history-stage">'
        f'<span class="dot">{_ui_escape(num)}</span>'
        f'<div><strong>{_ui_escape(name)}</strong><br><span>{_ui_escape(note)}</span></div>'
        f'<strong>{count}</strong>'
        '</div>'
        for num, name, count, note in stages
    )

    activity_rows = []
    for _, rec in view.head(5).iterrows():
        status_text = str(rec.get("Status", "Ready"))
        action = "Posted to ERP" if "sent" in status_text.lower() else ("Ready for ERP" if status_text == "Ready" else "Needs review")
        vendor = str(rec.get("Supplier") or rec.get("Accounting Party") or "Unknown")
        activity_rows.append(
            '<div class="activity-row">'
            f'<div><div class="activity-title">{_ui_escape(str(rec.get("Invoice #", "")))}</div>'
            f'<div class="activity-meta">{_ui_escape(action)} · {_ui_escape(_display_date(rec.get("Invoice Date") or rec.get("Processed At")))} · {_ui_escape(vendor[:46])}</div></div>'
            f'<div>{_table_status_chip(status_text)}</div>'
            '</div>'
        )

    table_rows = []
    for _, rec in view.head(80).iterrows():
        status_text = str(rec.get("Status", "Ready"))
        destination = "Export ready" if status_text == "Ready" else ("Posted evidence" if rec.get("Tally Status") else "Review")
        table_rows.append(
            '<tr>'
            f'<td class="primary mono">{_ui_escape(str(rec.get("Invoice #", "")))}</td>'
            f'<td>{_table_status_chip(status_text)}</td>'
            f'<td>{_ui_escape(_display_date(rec.get("Invoice Date") or rec.get("Processed At")))}</td>'
            f'<td>{_ui_escape(str(rec.get("Supplier") or rec.get("Accounting Party") or "Unknown")[:48])}</td>'
            f'<td class="money">{_ui_escape(_compact_amount(rec.get("Amount", 0), rec.get("Currency", "")))}</td>'
            f'<td>{_ui_escape(destination)}</td>'
            '</tr>'
        )

    st.markdown(
        f"""<div class="history-workspace">
        <div class="history-kpis">{kpis}</div>
        <div class="history-grid">
          <div class="history-panel">
            <h3>Close-ready pipeline</h3>
            {stage_html}
          </div>
          <div class="history-panel">
            <h3>Recent activity</h3>
            {''.join(activity_rows)}
          </div>
        </div>
        <div class="history-table-wrap">
          <table class="history-table-pro">
            <thead><tr><th>Invoice No.</th><th>Status</th><th>Inv. date</th><th>Vendor</th><th>Amount</th><th>Destination</th></tr></thead>
            <tbody>{''.join(table_rows)}</tbody>
          </table>
        </div>
        <div class="history-foot">Showing {len(view)} of {len(df)} records · EZ-Invoice workspace</div>
        </div>""",
        unsafe_allow_html=True,
    )


def render_analytics_page() -> None:
    df = combined_activity_df()
    if df.empty:
        st.markdown(
            """<div class="analytics-page">
            <div class="analytics-card">
              <h3>Analytics</h3>
              <div class="activity-title">No analytics yet.</div>
              <div class="metric-note">Upload invoices first, then spend, vendor, and parser intelligence will appear here.</div>
            </div>
            </div>""",
            unsafe_allow_html=True,
        )
        return

    df = df.copy()
    date_window = st.session_state.get("analytics_range", "30d")
    if date_window != "All":
        days = int(str(date_window).replace("d", ""))
        cutoff = datetime.now() - timedelta(days=days)
        date_basis = df.get("Processed At", df.get("Invoice Date", pd.Series(dtype=str))).map(_parse_display_date)
        filtered = df[date_basis.map(lambda dt: dt is None or dt >= cutoff)]
        if not filtered.empty:
            df = filtered

    df["Amount"] = pd.to_numeric(df.get("Amount", 0), errors="coerce").fillna(0)
    df["Tax"] = pd.to_numeric(df.get("Tax", 0), errors="coerce").fillna(0)
    df["Items"] = pd.to_numeric(df.get("Items", 0), errors="coerce").fillna(0)
    status_series = df.get("Status", pd.Series(dtype=str)).astype(str)
    exceptions = int((status_series != "Ready").sum())
    total = float(df["Amount"].sum())
    sent = int(status_series.str.contains("sent", case=False, na=False).sum())
    posted = float(df.loc[status_series.str.contains("sent", case=False, na=False), "Amount"].sum()) if sent else 0.0
    avg_ms = sum(st.session_state.ptimes.values()) / max(len(st.session_state.ptimes), 1)
    currencies = [c for c in df.get("Currency", pd.Series(dtype=str)).dropna().astype(str).unique() if c]
    currency_label = currencies[0] if len(currencies) == 1 else ""
    currency_count = len(currencies) or 1
    kpi_html = f"""
    <div class="analytics-page">
      <div class="analytics-kpis">
        <div class="analytics-card"><h3>Total spend</h3><div class="metric-big">{_ui_escape(_compact_amount(total, currency_label))}</div><div class="metric-note">{len(df)} invoice{'s' if len(df) != 1 else ''} · {currency_count} currenc{'ies' if currency_count != 1 else 'y'}</div></div>
        <div class="analytics-card"><h3>Sent to ERP</h3><div class="metric-big green">{sent} / {len(df)}</div><div class="metric-note">{_ui_escape(_compact_amount(posted, currency_label))} posted</div></div>
        <div class="analytics-card"><h3>Avg<br>processing</h3><div class="metric-big dark">{avg_ms / 1000:.1f}s</div><div class="metric-note">Extraction pipeline</div></div>
        <div class="analytics-card"><h3>Exceptions</h3><div class="metric-big amber">{exceptions}</div><div class="metric-note">{'Needs review' if exceptions else 'Clear'}</div></div>
      </div>
    """

    supplier_series = df.assign(Supplier=df.get("Supplier", pd.Series(["Unknown"] * len(df))).replace("", "Unknown")).groupby("Supplier")["Amount"].sum().sort_values(ascending=False).head(5)
    vendor_max = max(float(supplier_series.max()), 1.0) if not supplier_series.empty else 1.0
    vendor_html = ""
    for vendor, amount in supplier_series.items():
        width = max(8, min(100, float(amount) / vendor_max * 100))
        vendor_html += (
            f'<div class="row-line">'
            f'<div>{_ui_escape(str(vendor)[:34])}</div>'
            f'<div class="bar-track"><div class="bar-fill" style="width:{width:.0f}%"></div></div>'
            f'<div class="money">{_ui_escape(_compact_amount(amount, currency_label))}</div>'
            f'</div>'
        )

    line_df = _line_items_history_df()
    if not line_df.empty:
        cat_series = line_df.groupby("Category")["Amount"].sum().sort_values(ascending=False)
    else:
        cat_series = pd.Series({"Invoices": total})
    cat_total = max(float(cat_series.sum()), 1.0) if not cat_series.empty else 1.0
    cat_top = cat_series.head(4)
    colors = [theme["export"], theme["warn"], theme["good"], theme["accent"]]
    start = 0.0
    segments = []
    legend_html = ""
    for idx, (cat, amount) in enumerate(cat_top.items()):
        end = start + float(amount) / cat_total * 360
        color = colors[idx % len(colors)]
        segments.append(f"{color} {start:.1f}deg {end:.1f}deg")
        pct = float(amount) / cat_total * 100
        legend_html += f'<div class="legend-row"><span class="legend-dot" style="background:{color};"></span><span>{_ui_escape(str(cat))}</span><strong>{pct:.0f}%</strong></div>'
        start = end
    if not segments:
        segments = [f"{theme['export']} 0deg 360deg"]
        legend_html = '<div class="metric-note">No category data yet.</div>'
    top_cat = str(cat_top.index[0]) if len(cat_top) else "Spend"
    top_pct = (float(cat_top.iloc[0]) / cat_total * 100) if len(cat_top) else 0

    if "Currency" not in df.columns:
        df["Currency"] = currency_label or "USD"
    if "Invoice #" not in df.columns:
        df["Invoice #"] = [f"INV-{idx + 1}" for idx in range(len(df))]

    currency_html = ""
    currency_df = df.groupby("Currency").agg(Invoices=("Invoice #", "count"), Spend=("Amount", "sum")).sort_values("Spend", ascending=False).head(4)
    for currency, row in currency_df.iterrows():
        currency_html += f'<div class="mini-list-row currency-row"><div><span class="currency-pill">{_ui_escape(str(currency))}</span><span class="currency-count">{int(row["Invoices"])} invoice{"s" if int(row["Invoices"]) != 1 else ""}</span></div><div class="money">{_ui_escape(_compact_amount(row["Spend"], currency))}</div></div>'

    vendor_small = ""
    for vendor, amount in supplier_series.head(5).items():
        vendor_small += f'<div class="mini-list-row"><div>{_ui_escape(str(vendor)[:32])}</div><div class="money">{_ui_escape(_compact_amount(amount, currency_label))}</div></div>'

    item_html = ""
    if not line_df.empty:
        item_df = line_df.groupby("Description")["Amount"].sum().sort_values(ascending=False).head(5)
        for desc, amount in item_df.items():
            item_html += f'<div class="mini-list-row"><div>{_ui_escape(str(desc)[:30])}</div><div class="money">{_ui_escape(_compact_amount(amount, currency_label))}</div></div>'

    analytics_html = (
        kpi_html
        + f"""<div class="analytics-grid">
<div class="analytics-card analytics-card-tall">
<h3>Spend by vendor <span>Equivalent</span></h3>
{vendor_html or '<div class="metric-note">No vendor spend yet.</div>'}
</div>
<div class="analytics-card analytics-card-tall">
<h3>Spend by category</h3>
<div class="donut-wrap">
<div class="donut" style="background:conic-gradient({', '.join(segments)});"><div class="donut-center">{top_pct:.0f}%<br><span>{_ui_escape(top_cat)}</span></div></div>
<div>{legend_html}</div>
</div>
</div>
</div>
<div class="analytics-grid three">
<div class="analytics-card analytics-card-small"><h3>By currency</h3>{currency_html or '<div class="metric-note">No currency data yet.</div>'}</div>
<div class="analytics-card analytics-card-small"><h3>Top vendors</h3>{vendor_small or '<div class="metric-note">No vendor data yet.</div>'}</div>
<div class="analytics-card analytics-card-small"><h3>Top line items</h3>{item_html or '<div class="metric-note">No line items yet.</div>'}</div>
</div>
</div>"""
    )
    st.markdown(analytics_html, unsafe_allow_html=True)


def render_exceptions_page() -> None:
    _page_heading("Exceptions", "Invoices that need human review")
    df = combined_activity_df()
    if df.empty:
        st.info("No exceptions yet.")
        return
    exceptions = df[df.get("Status", "") != "Ready"]
    if exceptions.empty:
        st.success("All processed invoices are currently ready.")
        return
    cols = [c for c in ["Processed At", "Invoice #", "Supplier", "Amount", "Currency", "Issues", "File"] if c in exceptions.columns]
    st.dataframe(exceptions[cols], width="stretch", hide_index=True)


def render_tally_queue_panel() -> None:
    st.markdown("#### Tally Posting Queue")
    payloads = st.session_state.get("payloads", {})
    if not payloads:
        st.info("Process invoices first. Approved inbound purchases will appear here for Tally posting.")
        return

    queue_rows = []
    ready_payloads: Dict[str, Dict[str, Any]] = {}
    status_map = st.session_state.get("tally_post_status", {})
    for fname, data in payloads.items():
        payload = data.get("payload", {})
        inv = payload.get("INVOICE", {})
        header = inv.get("INVOICE HEADER", {})
        route = inv.get("ROUTING", {})
        readiness = tally_readiness_message(data)
        ready = readiness == "Ready for Tally."
        if ready:
            ready_payloads[fname] = data
        post_status = status_map.get(fname, data.get("tally_status", {})) or {}
        queue_rows.append(
            {
                "Select": ready,
                "File": fname,
                "Invoice #": header.get("INVOICE NO.", ""),
                "Direction": route.get("DIRECTION", "unknown"),
                "Vendor/Party": inv.get("SELLER", {}).get("NAME", ""),
                "Amount": safe_float0(header.get("INVOICE AMOUNT", 0)),
                "Readiness": readiness,
                "Last Tally Result": post_status.get("message", ""),
            }
        )

    queue_df = pd.DataFrame(queue_rows)
    edited = st.data_editor(
        queue_df,
        width="stretch",
        hide_index=True,
        disabled=[c for c in queue_df.columns if c != "Select"],
        column_config={
            "Select": st.column_config.CheckboxColumn("Send", help="Only ready inbound invoices can be sent."),
            "Amount": st.column_config.NumberColumn("Amount", format="%.2f"),
        },
        key="tally_queue_editor",
    )

    selected_files = [
        row["File"]
        for _, row in edited.iterrows()
        if bool(row.get("Select")) and row.get("File") in ready_payloads
    ]
    selected_payloads = {fname: ready_payloads[fname] for fname in selected_files}

    c1, c2, c3 = st.columns([1, 1, 1])
    with c1:
        st.metric("Ready", len(ready_payloads))
    with c2:
        st.metric("Selected", len(selected_payloads))
    with c3:
        st.metric("Blocked", len(payloads) - len(ready_payloads))

    a1, a2, a3 = st.columns([1, 1, 1])
    with a1:
        st.download_button(
            "Download Selected XML",
            make_tally_zip(selected_payloads or ready_payloads, classifier),
            "selected_tally_xml.zip",
            "application/zip",
            width="stretch",
            disabled=not bool(selected_payloads or ready_payloads),
            key="tally_queue_download_xml",
        )
    with a2:
        if st.button("Dry Run Selected", width="stretch", disabled=not bool(selected_payloads), key="tally_queue_dry_run"):
            result = send_tally_batch(selected_payloads, classifier=classifier, dry_run=True)
            for fname, item in result.get("results", {}).items():
                _store_tally_status(fname, item, dry_run=True)
            if result.get("success"):
                st.success("Tally connector accepted the selected vouchers.")
            else:
                st.warning(result.get("message", "Dry run completed with issues."))
    with a3:
        if st.button("Post Selected to Tally", type="primary", width="stretch", disabled=not bool(selected_payloads), key="tally_queue_post"):
            result = send_tally_batch(selected_payloads, classifier=classifier, dry_run=False)
            for fname, item in result.get("results", {}).items():
                _store_tally_status(fname, item, dry_run=False)
            if result.get("success"):
                st.success("Selected vouchers posted to Tally.")
            else:
                st.warning(result.get("message", "Posting completed with issues."))


INTEGRATION_CATALOG = {
    "QuickBooks": {
        "mark": "QB",
        "type": "Cloud accounting",
        "description": "Create supplier Bills with OAuth, vendor resolution, and client Chart of Accounts mappings.",
    },
    "Tally": {
        "mark": "TA",
        "type": "Desktop accounting",
        "description": "Post approved purchase vouchers through a secure connector running beside TallyPrime.",
    },
    "Zoho Books": {
        "mark": "ZB",
        "type": "Cloud accounting",
        "description": "Resolve vendors, map purchase accounts and taxes, and create Bills in a selected organization.",
    },
    "Coupa": {
        "mark": "CO",
        "type": "Spend management",
        "description": "Prepare normalized supplier-invoice packages for a client Coupa integration.",
    },
    "NetSuite": {
        "mark": "NS",
        "type": "Cloud ERP",
        "description": "Prepare vendor-bill payloads for a client-specific NetSuite REST integration.",
    },
    "SAP": {
        "mark": "SAP",
        "type": "Enterprise ERP",
        "description": "Prepare supplier-invoice data for SAP Business One or S/4HANA connector implementation.",
    },
}


def _integration_connection_state(system: str) -> Tuple[str, str]:
    if system == "QuickBooks":
        return ("Connected", "ready") if is_connected() else ("Setup required", "setup")
    if system == "Tally":
        return ("Configured", "ready") if tally_is_connected() else ("Setup required", "setup")
    if system == "Zoho Books":
        return ("Connected", "ready") if zoho_is_connected() else ("Setup required", "setup")
    return "Export ready", "export"


def _render_integration_shell(system: str) -> None:
    meta = INTEGRATION_CATALOG[system]
    status, status_class = _integration_connection_state(system)
    st.markdown(
        f"""<section class="integration-hero">
        <div class="integration-brand-mark">{_ui_escape(meta["mark"])}</div>
        <div class="integration-hero-copy">
          <span class="integration-kind">{_ui_escape(meta["type"])}</span>
          <div class="integration-title-row">
            <strong class="integration-system-name">{_ui_escape(system)}</strong>
            <div class="integration-status {status_class}"><i></i>{_ui_escape(status)}</div>
          </div>
          <p>{_ui_escape(meta["description"])}</p>
        </div>
        </section>
        <section class="integration-flow">
          <div><b>1</b><span><strong>Connect</strong><small>Authorize the client system</small></span></div>
          <div><b>2</b><span><strong>Map</strong><small>Align accounts and tax rules</small></span></div>
          <div><b>3</b><span><strong>Validate</strong><small>Check vendor and invoice data</small></span></div>
          <div><b>4</b><span><strong>Post</strong><small>Return an external document ID</small></span></div>
        </section>""",
        unsafe_allow_html=True,
    )


def _render_tally_client_workspace(profile: Dict[str, Any]) -> None:
    st.markdown(
        '<div class="integration-section-heading"><div><span>Workspace</span>'
        '<strong>Client identity and routing</strong></div>'
        '<p>Legal names determine whether invoices route into the inbound purchase queue.</p></div>',
        unsafe_allow_html=True,
    )
    c1, c2, c3 = st.columns([1, 1, 1])
    with c1:
        client_name = st.text_input(
            "Client / company name",
            value=profile.get("client_name", ""),
            placeholder="Example: Client legal company",
            key="client_profile_name",
        )
    with c2:
        tally_company = st.text_input(
            "Tally company open now",
            value=profile.get("tally_company", ""),
            placeholder="Exact company name open in TallyPrime",
            key="client_tally_company",
        )
    with c3:
        st.text_input("Posting direction", value="Inbound purchases", disabled=True, key="tally_route_display")
    home_names = st.text_area(
        "Client legal names and aliases",
        value=profile.get("home_company_names", ""),
        placeholder="Client Legal Name Pvt Ltd\nClient Trading Alias",
        height=92,
        key="client_home_names",
    )
    if not _split_company_names(home_names):
        st.warning("Add at least one legal name before posting invoices to Tally.")
    b1, b2 = st.columns([2, 1])
    with b1:
        if st.button("Save workspace", type="primary", key="save_client_workspace", width="stretch"):
            legal_names = home_names.strip()
            if not legal_names and client_name.strip() and client_name.strip() != "Client Workspace":
                legal_names = client_name.strip()
            updated = {
                "client_name": client_name.strip() or "Client Workspace",
                "home_company_names": legal_names,
                "primary_accounting_system": "TallyPrime",
                "tally_url": profile.get("tally_url", "http://localhost:9000"),
                "tally_company": tally_company.strip(),
            }
            st.session_state["client_profile"] = updated
            st.session_state["ez_target_system"] = "Tally"
            if _save_client_profile(updated):
                st.success("Workspace saved. Reprocess invoices when legal names change.")
            else:
                st.warning("Workspace is active for this session, but local persistence failed.")
    with b2:
        if st.button("Reset fields", key="reset_client_workspace", width="stretch"):
            st.session_state["client_profile"] = dict(DEFAULT_CLIENT_PROFILE)
            _save_client_profile(dict(DEFAULT_CLIENT_PROFILE))
            st.rerun()


def _render_export_connector(system: str) -> None:
    requirements = {
        "Coupa": ("Client API tenant and OAuth credentials", "Supplier, account and tax-code mappings", "Invoice API acceptance rules"),
        "NetSuite": ("Sandbox account and integration record", "Subsidiary, vendor and currency mappings", "REST role and token permissions"),
        "SAP": ("SAP Business One Service Layer or S/4HANA sandbox", "Business partner, item, tax and warehouse mappings", "Sample accepted A/P invoice"),
    }
    items = requirements[system]
    st.markdown(
        f"""<section class="integration-readiness">
        <div>
          <span>Available now</span>
          <h3>Validated export packages</h3>
          <p>Process invoices once, review the normalized fields, then download the prepared {_ui_escape(system)} package from the invoice workspace.</p>
        </div>
        <div class="integration-readiness-badge">Export ready</div>
        </section>
        <div class="integration-section-heading"><div><span>Live connector</span>
        <strong>Required from the pilot client</strong></div>
        <p>Production posting is enabled only after the client provides a test tenant and accepted master-data contracts.</p></div>
        <section class="integration-requirements">
          <div><b>01</b><strong>{_ui_escape(items[0])}</strong></div>
          <div><b>02</b><strong>{_ui_escape(items[1])}</strong></div>
          <div><b>03</b><strong>{_ui_escape(items[2])}</strong></div>
        </section>""",
        unsafe_allow_html=True,
    )
    if st.button("Open invoice workspace", type="primary", key="open_invoice_workspace_" + _nav_slug(system)):
        st.session_state["ez_nav_page"] = "Invoices"
        st.session_state["ez_target_system"] = system
        st.session_state["ez_fmt_choice"] = system
        st.query_params["page"] = "Invoices"
        st.query_params["target"] = system
        st.rerun()


def render_integrations_page() -> None:
    active_page = st.session_state.get("ez_nav_page", "QuickBooks")
    if active_page not in INTEGRATION_CATALOG:
        active_page = "QuickBooks"
    _render_integration_shell(active_page)
    profile = current_client_profile()

    if active_page == "Tally":
        with st.container(key="integration_page_body"):
            workspace_tab, connector_tab, queue_tab = st.tabs(["Workspace", "Connector & vouchers", "Posting queue"])
            with workspace_tab:
                _render_tally_client_workspace(profile)
            with connector_tab:
                st.markdown(
                    '<div class="integration-section-heading"><div><span>Configuration</span>'
                    '<strong>Local connector and voucher rules</strong></div>'
                    '<p>Keep TallyPrime open with the client company loaded while testing.</p></div>',
                    unsafe_allow_html=True,
                )
                tally_sidebar(show_heading=False)
            with queue_tab:
                render_tally_queue_panel()
        return

    if active_page == "QuickBooks":
        with st.container(key="integration_page_body"):
            st.markdown(
                '<div class="integration-section-heading"><div><span>Configuration</span>'
                '<strong>Connection and account mapping</strong></div>'
                '<p>Authorize the sandbox or client company, then map universal invoice categories.</p></div>',
                unsafe_allow_html=True,
            )
            with st.container(key="quickbooks_connector_workspace"):
                qb_sidebar(show_heading=False)
        return

    if active_page == "Zoho Books":
        with st.container(key="integration_page_body"):
            st.markdown(
                '<div class="integration-section-heading"><div><span>Configuration</span>'
                '<strong>Organization and purchase mapping</strong></div>'
                '<p>Select the client organization, accounts, and purchase-tax behavior before posting.</p></div>',
                unsafe_allow_html=True,
            )
            with st.container(key="zoho_connector_workspace"):
                zoho_sidebar(show_heading=False)
        return

    with st.container(key="integration_page_body"):
        _render_export_connector(active_page)


def render_roadmap_page() -> None:
    _page_heading("Roadmap", "From local prototype to B2B SaaS product")
    roadmap_path = APP_DIR / "ROADMAP.md"
    if roadmap_path.exists():
        st.markdown(roadmap_path.read_text(encoding="utf-8"))
    else:
        st.info("ROADMAP.md is not available.")


def _ui_escape(value: Any) -> str:
    return html.escape(str(value or ""))


def _status_class(status: str) -> str:
    s = (status or "").lower()
    if "sent" in s:
        return "sent"
    if "ready" in s or "validated" in s:
        return "ok"
    if "review" in s or "exception" in s:
        return "bad"
    if "pending" in s:
        return "warn"
    return "neutral"


def _session_invoice_rows() -> List[Dict[str, Any]]:
    rows = []
    for fname, data in st.session_state.get("payloads", {}).items():
        record = _history_record(fname, data, st.session_state.get("ptimes", {}).get(fname, 0))
        record["_source"] = "session"
        record["_file"] = fname
        rows.append(record)
    return rows


def _activity_rows_for_console() -> List[Dict[str, Any]]:
    return _session_invoice_rows()


def select_console_invoice(fname: str) -> None:
    st.session_state["selected_invoice_file"] = fname


def _selected_payload_for_console(rows: List[Dict[str, Any]]) -> Tuple[Optional[str], Optional[Dict[str, Any]]]:
    if not rows:
        return None, None
    row_files = [row.get("_file") or row.get("File") for row in rows]
    selected = st.session_state.get("selected_invoice_file")
    fname = selected if selected in row_files else row_files[0]
    st.session_state["selected_invoice_file"] = fname
    payload_data = st.session_state.get("payloads", {}).get(fname)
    return fname, payload_data


def render_export_package(payloads: Dict[str, Dict[str, Any]], fmt_choice: str) -> None:
    if not payloads:
        return

    invoice_count = len(payloads)
    noun = "invoice" if invoice_count == 1 else "invoices"
    total_amount = sum(
        safe_float0(data.get("payload", {}).get("INVOICE", {}).get("INVOICE HEADER", {}).get("INVOICE AMOUNT", 0))
        for data in payloads.values()
    )
    currencies = []
    for data in payloads.values():
        inv = data.get("payload", {}).get("INVOICE", {})
        cur = inv.get("PAYMENT", {}).get("ELECTRONIC", {}).get("CURRENCY", "")
        if cur and cur not in currencies:
            currencies.append(cur)
    currency_label = currencies[0] if len(currencies) == 1 else ""
    mapping_status = "Client GL mapping applied" if classifier and classifier.client_map else "Generic category mapping"
    target_label = fmt_choice if fmt_choice in CONVERTERS else "Accounting export"

    excel_buf, excel_fname = make_einvoice_excel(payloads, classifier)
    export_options = [
        ("QuickBooks ZIP", make_zip(payloads, to_quickbooks, "quickbooks"), "invoices_quickbooks.zip", "application/zip", "export_qb_zip"),
        ("Tally XML ZIP", make_tally_zip(payloads, classifier), "invoices_tally_xml.zip", "application/zip", "export_tally_zip"),
        ("Zoho Books ZIP", make_zip(payloads, build_zoho_export, "zoho_books"), "invoices_zoho_books.zip", "application/zip", "export_zoho_zip"),
        ("Coupa ZIP", make_zip(payloads, to_coupa, "coupa"), "invoices_coupa.zip", "application/zip", "export_coupa_zip"),
        ("NetSuite ZIP", make_zip(payloads, to_netsuite, "netsuite"), "invoices_netsuite.zip", "application/zip", "export_ns_zip"),
        ("SAP ZIP", make_zip(payloads, to_sap, "sap"), "invoices_sap.zip", "application/zip", "export_sap_zip"),
        ("Excel", excel_buf, excel_fname, "application/vnd.openxmlformats-officedocument.spreadsheetml.document", "export_excel"),
    ]

    with st.container(key="batch_exports"):
        with st.popover("Export Package", width="stretch"):
                st.markdown(
                    f"""<div class="ez-export-menu ez-export-package-menu">
                    <div class="ez-export-menu-head">
                      <div class="ez-export-menu-title">Export package</div>
                      <div class="ez-export-menu-sub">Download one prepared package for the target accounting system.</div>
                    </div>
                    <div class="ez-export-facts">
                      <div><span>Package</span><strong>{invoice_count} {noun}</strong></div>
                      <div><span>Target</span><strong>{_ui_escape(target_label)}</strong></div>
                      <div><span>Mapping</span><strong>{_ui_escape(mapping_status)}</strong></div>
                      <div><span>Total</span><strong>{_ui_escape(_currency_money_label(total_amount, currency_label))}</strong></div>
                    </div>
                    </div>""",
                    unsafe_allow_html=True,
                )
                export_cols = st.columns(len(export_options), gap="small")
                for col, option in zip(export_cols, export_options):
                    with col:
                        st.download_button(*option, width="stretch")


def _line_item_display_value(column: str, value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    if column in {"QUANTITY", "UNIT PRICE", "AMOUNT"}:
        number = safe_float0(value)
        return f"{number:,.0f}" if abs(number - round(number)) < 0.005 else f"{number:,.2f}"
    return str(value)


def _render_line_items_table(line_df: pd.DataFrame, display_cols: List[str]) -> None:
    header_labels = {
        "DESCRIPTION": "Description",
        "QUANTITY": "Quantity",
        "UOM": "UOM",
        "UNIT PRICE": "Unit price",
        "AMOUNT": "Amount",
        "CATEGORY": "Category",
        "GL_CODE": "GL code",
    }
    headers = "".join(f"<th>{_ui_escape(header_labels.get(col, col).upper())}</th>" for col in display_cols)
    body_rows = []
    numeric_cols = {"QUANTITY", "UNIT PRICE", "AMOUNT"}
    for _, row in line_df[display_cols].iterrows():
        cells = []
        for col in display_cols:
            klass = "num" if col in numeric_cols else ""
            cells.append(f'<td class="{klass}">{_ui_escape(_line_item_display_value(col, row.get(col)))}</td>')
        body_rows.append(f"<tr>{''.join(cells)}</tr>")
    st.markdown(
        f"""<div class="ez-line-table-wrap">
        <table class="ez-line-table">
          <thead><tr>{headers}</tr></thead>
          <tbody>{''.join(body_rows)}</tbody>
        </table>
        </div>""",
        unsafe_allow_html=True,
    )


def render_invoice_console(fmt_choice: str) -> None:
    rows = _activity_rows_for_console()
    uploaded_n = len(rows)
    extracted_n = uploaded_n
    validated_n = sum(1 for r in rows if str(r.get("Status", "")).lower() in ("ready", "sent"))
    exceptions_n = sum(1 for r in rows if str(r.get("Status", "")).lower() not in ("ready", "sent"))
    sent_n = sum(1 for r in rows if "sent" in str(r.get("Status", "")).lower())

    st.markdown(f"""<div class="ap-status-strip">
      <div class="ap-status"><strong>{uploaded_n}</strong><span>Uploaded</span></div>
      <div class="ap-status active"><strong>{extracted_n}</strong><span>Extracted</span></div>
      <div class="ap-status"><strong>{validated_n}</strong><span>Validated</span></div>
      <div class="ap-status"><strong>{sent_n}</strong><span>Sent to ERP</span></div>
      <div class="ap-status danger"><strong>{exceptions_n}</strong><span>Exceptions</span></div>
    </div>""", unsafe_allow_html=True)

    render_export_package(st.session_state.get("payloads", {}), fmt_choice)

    list_col, detail_col = st.columns([0.36, 0.64], gap="small")
    with list_col:
        st.markdown(f"""<div class="pane-title"><span>{uploaded_n} invoices</span><span>Date</span></div>""", unsafe_allow_html=True)
        if not rows:
            st.markdown("""<div class="empty-pane"><strong>No invoices yet</strong><span>Upload PDFs above to populate the invoice queue.</span></div>""", unsafe_allow_html=True)
        else:
            selected_file = st.session_state.get("selected_invoice_file")
            row_files = [row.get("_file") or row.get("File") for row in rows]
            if selected_file not in row_files:
                selected_file = row_files[0]
                st.session_state["selected_invoice_file"] = selected_file
            for idx, record in enumerate(rows[:25]):
                fname_for_row = record.get("_file") or record.get("File") or f"row_{idx}"
                invoice_no = _ui_escape(record.get("Invoice #") or "Draft invoice")
                supplier = _ui_escape(record.get("Supplier") or "Unknown supplier")
                date = _ui_escape(record.get("Invoice Date") or record.get("Processed At") or "")
                amount = _currency_money_label(record.get("Amount", 0), record.get("Currency", ""))
                status = _ui_escape(record.get("Status") or "Pending")
                klass = _status_class(status)
                active = " active" if fname_for_row == selected_file else ""
                safe_key = re.sub(r"[^a-zA-Z0-9_]+", "_", str(fname_for_row))[:54]
                with st.container(key=f"invoice_row_{idx}_{safe_key}"):
                    st.markdown(f"""<div class="invoice-list-row{active}">
                    <div>
                      <strong>{invoice_no}</strong>
                      <span>{supplier}</span>
                      <em>{date}</em>
                    </div>
                    <div class="invoice-row-meta">
                      <b>{_ui_escape(amount)}</b>
                      <i class="status-pill {klass}">{status}</i>
                    </div>
                </div>""", unsafe_allow_html=True)
                    st.button(
                        f"Select invoice {record.get('Invoice #') or idx + 1}",
                        key=f"select_invoice_row_{idx}_{safe_key}",
                        on_click=select_console_invoice,
                        args=(fname_for_row,),
                    )

    with detail_col:
        fname, payload_data = _selected_payload_for_console(rows)
        if payload_data:
            payload = payload_data["payload"]
            inv = payload.get("INVOICE", {})
            h = inv.get("INVOICE HEADER", {})
            supplier = inv.get("SELLER", {}).get("NAME", "Unknown supplier")
            invoice_no = h.get("INVOICE NO.", "Draft invoice")
            cur = inv.get("PAYMENT", {}).get("ELECTRONIC", {}).get("CURRENCY", "USD")
            rows_li = inv.get("LINE ITEMS", {}).get("ROWS", []) or []
            inv_total = safe_float0(h.get("INVOICE AMOUNT", 0))
            line_sum = sum(safe_float0(r.get("AMOUNT", 0)) for r in rows_li)
            conf = max(0, 100 - (0 if h.get("INVOICE NO.") else 25) - (0 if rows_li else 30) - (0 if abs(line_sum-inv_total)<0.5 else 20) - (0 if h.get("DUE DATE") else 5))
            status = "Ready" if payload_data.get("ok") else "Needs Review"
            route = inv.get("ROUTING", {})
            tally_status = st.session_state.get("tally_post_status", {}).get(fname, payload_data.get("tally_status", {})) or {}
            dh_left, dh_right = st.columns([0.68, 0.32])
            with dh_left:
                st.markdown(f"""<div class="detail-title-block">
                  <span class="eyebrow">Invoice</span>
                  <h2>{_ui_escape(invoice_no)}</h2>
                  <p>{_ui_escape(supplier)}</p>
                </div>""", unsafe_allow_html=True)
            with dh_right:
                st.markdown(f"""<div class="detail-badges">
                  <span class="status-pill {_status_class(status)}">{_ui_escape(status)}</span>
                  <span class="status-pill neutral">{conf}% confidence</span>
                </div>""", unsafe_allow_html=True)
            st.markdown(f"""<div class="detail-grid">
              <div><span>Invoice Date</span><strong>{_ui_escape(h.get("INVOICE DATE", ""))}</strong></div>
              <div><span>Due Date</span><strong>{_ui_escape(h.get("DUE DATE", ""))}</strong></div>
              <div><span>Currency</span><strong>{_ui_escape(cur)}</strong></div>
              <div><span>Total</span><strong>{_ui_escape(_currency_money_label(inv_total, cur))}</strong></div>
              <div><span>Line Items</span><strong>{len(rows_li)}</strong></div>
              <div><span>Route</span><strong>{_ui_escape(route.get("DIRECTION", "unknown").title())}</strong></div>
            </div>""", unsafe_allow_html=True)
            if tally_status.get("message"):
                if tally_status.get("success"):
                    st.success(f"{tally_status.get('mode', 'Tally')}: {tally_status.get('message')}")
                else:
                    st.warning(f"{tally_status.get('mode', 'Tally')}: {tally_status.get('message')}")

            with st.container(key="invoice_detail_actions"):
                a1, a2, a3 = st.columns(3)
                with a1:
                    st.download_button("Parsed JSON", json.dumps(payload, indent=2), f"{fname.rsplit('.',1)[0]}_parsed.json", "application/json", width="stretch", key=f"console_json_{fname}")
                with a2:
                    if fmt_choice == "Tally":
                        st.download_button("Tally XML", build_tally_xml(payload, classifier=classifier), f"{fname.rsplit('.',1)[0]}_tally.xml", "application/xml", width="stretch", key=f"console_tally_{fname}")
                    else:
                        st.download_button(f"{fmt_choice} JSON", json.dumps(CONVERTERS[fmt_choice](payload), indent=2), f"{fname.rsplit('.',1)[0]}_{fmt_choice.lower()}.json", "application/json", width="stretch", key=f"console_fmt_{fname}")
                with a3:
                    if tally_is_connected() and fmt_choice == "Tally":
                        blocker = tally_post_blocker(payload)
                        if blocker:
                            st.warning(blocker)
                        else:
                            if st.button("Dry Run Tally", width="stretch", key=f"console_dry_tally_{fname}"):
                                r = send_to_tally(payload, classifier, dry_run=True)
                                _store_tally_status(fname, r, dry_run=True)
                                _record_api_posting_result(payload_data, "tally", r, dry_run=True)
                                if r["success"]:
                                    st.success(r["message"])
                                else:
                                    st.warning(r["message"])
                            if st.button("Post to Tally", width="stretch", key=f"console_send_tally_{fname}"):
                                r = send_to_tally(payload, classifier, dry_run=False)
                                _store_tally_status(fname, r, dry_run=False)
                                _record_api_posting_result(payload_data, "tally", r, dry_run=False)
                                if r["success"]:
                                    st.success(r["message"])
                                else:
                                    st.warning(r["message"])
                    elif fmt_choice == "QuickBooks" and is_connected():
                        if st.button("Send to QuickBooks", width="stretch", key=f"console_send_qb_{fname}"):
                            r = send_to_quickbooks(payload, classifier)
                            _record_api_posting_result(payload_data, "quickbooks", r, dry_run=False)
                            if r["success"]:
                                st.success(r["message"])
                            else:
                                st.warning(r["message"])
                    elif fmt_choice == "Zoho Books" and zoho_is_connected():
                        if st.button("Send to Zoho Books", width="stretch", key=f"console_send_zoho_{fname}"):
                            r = send_to_zoho(payload, classifier)
                            _record_api_posting_result(payload_data, "zoho_books", r, dry_run=False)
                            if r["success"]:
                                st.success(r["message"])
                            else:
                                st.warning(r["message"])
                    else:
                        st.button("Send Bill", disabled=True, width="stretch", key=f"console_disabled_{fname}")

            st.markdown('<div class="ez-line-section-title">Line Items</div>', unsafe_allow_html=True)
            if rows_li:
                line_df = pd.DataFrame(classifier.classify_invoice_rows(rows_li) if classifier else rows_li)
                display_cols = [c for c in ["DESCRIPTION", "QUANTITY", "UOM", "UNIT PRICE", "AMOUNT", "CATEGORY", "GL_CODE"] if c in line_df.columns]
                if not display_cols:
                    display_cols = list(line_df.columns)
                _render_line_items_table(line_df, display_cols)
            else:
                st.info("No line items captured for this invoice.")
        elif rows:
            selected_key = st.session_state.get("selected_invoice_file")
            record = next(
                (
                    r for r in rows
                    if (r.get("_file") or r.get("File")) == selected_key
                ),
                rows[0],
            )
            st.markdown(f"""<div class="detail-head">
                <div><span class="eyebrow">Invoice</span><h2>{_ui_escape(record.get("Invoice #"))}</h2><p>{_ui_escape(record.get("Supplier"))}</p></div>
                <div class="detail-actions"><span class="status-pill {_status_class(record.get("Status", ""))}">{_ui_escape(record.get("Status"))}</span></div>
            </div>
            <div class="detail-grid">
              <div><span>Invoice Date</span><strong>{_ui_escape(record.get("Invoice Date"))}</strong></div>
              <div><span>Due Date</span><strong>{_ui_escape(record.get("Due Date"))}</strong></div>
              <div><span>Currency</span><strong>{_ui_escape(record.get("Currency"))}</strong></div>
              <div><span>Total</span><strong>{_ui_escape(_currency_money_label(record.get("Amount", 0), record.get("Currency", "")))}</strong></div>
            </div>""", unsafe_allow_html=True)
            st.info("Detailed review is available for invoices processed in the current session.")
        else:
            st.markdown("""<div class="empty-pane large"><strong>Select an invoice</strong><span>Details, validation, actions, and line items appear here.</span></div>""", unsafe_allow_html=True)
# =========================
# UI — Futuristic SaaS Command Center
# =========================

# Keep a Streamlit theme file for native widgets; app-level light/dark mode is
# driven by CSS tokens below so it can switch at runtime.
import os as _os
_toml_dir = _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), ".streamlit")
_os.makedirs(_toml_dir, exist_ok=True)
_toml_path = _os.path.join(_toml_dir, "config.toml")
if not _os.path.exists(_toml_path):
    with open(_toml_path, "w") as _f:
        _f.write("""[theme]
base = "light"
primaryColor = "#4A6741"
backgroundColor = "#FAF9F6"
secondaryBackgroundColor = "#F0EFEB"
textColor = "#1A1A18"
font = "sans serif"
""")

st.set_page_config(page_title="EZ-Invoice", page_icon="⚡", layout="wide")
complete_qb_auth_if_ready()
complete_zoho_auth_if_ready()

theme_mode = st.session_state.get("ez_theme_mode", "Light")
current_page = st.query_params.get("page", st.session_state.get("ez_nav_page", "Invoices"))
VALID_NAV_PAGES = ["Invoices", "History", "Analytics", "Exceptions", "QuickBooks", "Tally", "Zoho Books", "Coupa", "NetSuite", "SAP", "GL Mapping", "Rules", "Vendors", "Roadmap"]
ERP_SYSTEMS = ["QuickBooks", "Tally", "Zoho Books", "Coupa", "NetSuite", "SAP"]
if current_page not in VALID_NAV_PAGES:
    current_page = "Invoices"
st.session_state["ez_nav_page"] = current_page
target_system = st.query_params.get("target", st.session_state.get("ez_target_system", "QuickBooks"))
if target_system not in ERP_SYSTEMS:
    target_system = "QuickBooks"
st.session_state["ez_target_system"] = target_system
if "ez_fmt_choice" not in st.session_state:
    st.session_state["ez_fmt_choice"] = target_system
if current_page in ERP_SYSTEMS and st.session_state.get("ez_fmt_choice") != current_page:
    st.session_state["ez_fmt_choice"] = current_page
    st.session_state["ez_target_system"] = current_page
    target_system = current_page
is_light_mode = theme_mode == "Light"
theme = {
    "bg": "#FAF9F6" if is_light_mode else "#111110",
    "bg2": "#FAF9F6" if is_light_mode else "#111110",
    "surface": "#F0EFEB" if is_light_mode else "#1A1A18",
    "surface2": "#F7F5EF" if is_light_mode else "#20201E",
    "panel": "#FAF9F6" if is_light_mode else "#111110",
    "panel2": "#F0EFEB" if is_light_mode else "#1A1A18",
    "text": "#1A1A18" if is_light_mode else "#F5F4F0",
    "muted": "#6B6960" if is_light_mode else "#A09D94",
    "faint": "#B0AEA6" if is_light_mode else "#706D65",
    "border": "#E0DED8" if is_light_mode else "#30302D",
    "border2": "#D5D2CA" if is_light_mode else "#3B3934",
    "accent": "#4A6741" if is_light_mode else "#7DA46E",
    "accent2": "#4A6741" if is_light_mode else "#7DA46E",
    "accent3": "#B2762D" if is_light_mode else "#D69A45",
    "export": "#4A6741" if is_light_mode else "#8FB37F",
    "export_hover": "#3F5938" if is_light_mode else "#A4C895",
    "good": "#4A6741" if is_light_mode else "#7DA46E",
    "warn": "#9B681E" if is_light_mode else "#E0A24A",
    "danger": "#A53E38" if is_light_mode else "#F06B64",
    "shadow": "0 1px 2px rgba(26, 26, 24, 0.04)" if is_light_mode else "0 1px 2px rgba(0, 0, 0, 0.2)",
    "hero": "#F0EFEB" if is_light_mode else "#1A1A18",
    "grid": "transparent",
    "input": "#FFFFFF" if is_light_mode else "#181817",
    "active_row": "#EDF3EA" if is_light_mode else "#1A2418",
    "chip_sent_bg": "#D4DFCE" if is_light_mode else "#25351F",
    "chip_validated_bg": "#E8EFE4" if is_light_mode else "#1F2C1B",
    "chip_exception_bg": "#FBDDD8" if is_light_mode else "#351E1C",
    "chip_pending_bg": "#FBE8C8" if is_light_mode else "#332715",
}

st.markdown(f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=JetBrains+Mono:wght@400;500;700&family=Inter:wght@400;500;600;700;800&display=swap');

html, body, .stApp {{
    font-family: 'Inter', sans-serif !important;
    color: {theme["text"]} !important;
    background: {theme["bg"]} !important;
    background-image: none !important;
}}
.block-container {{ padding-top: 2.2rem; max-width: 1500px; }}
section[data-testid="stSidebar"] {{
    background: {theme["surface"]} !important;
    border-right: 1px solid {theme["border"]};
}}
section[data-testid="stSidebar"] * {{ color: {theme["text"]}; }}

.ez-shell {{ display:flex; flex-direction:column; gap:18px; margin-bottom: 18px; }}
.ez-topbar {{
    display:flex; align-items:center; justify-content:space-between; gap:14px;
    padding: 12px 14px; border:1px solid {theme["border"]}; border-radius:18px;
    background:{theme["surface"]};
}}
.ez-brand {{ display:flex; align-items:center; gap:10px; min-width:0; }}
.ez-logo {{
    width:36px; height:36px; border-radius:10px;
    display:grid; place-items:center; color:white; font-weight:800;
    background: {theme["accent"]};
}}
.ez-brand-text strong {{ display:block; font-size:.94rem; letter-spacing:.01em; }}
.ez-brand-text span {{ display:block; color:{theme["muted"]}; font-size:.72rem; margin-top:1px; }}
.ez-top-pills {{ display:flex; gap:8px; align-items:center; flex-wrap:wrap; justify-content:flex-end; }}
.ez-pill {{
    display:inline-flex; align-items:center; gap:7px; min-height:30px; padding:6px 10px;
    border:1px solid {theme["border"]}; border-radius:999px;
    background:{theme["surface2"]}; color:{theme["muted"]}; font-size:.72rem; font-weight:700;
}}
.ez-dot {{ width:7px; height:7px; border-radius:50%; background:{theme["good"]}; box-shadow:0 0 18px {theme["good"]}; }}

.hero {{
    position:relative; overflow:hidden;
    display:grid; grid-template-columns:minmax(0,1.35fr) minmax(310px,.65fr); gap:18px;
    padding:28px; border:1px solid {theme["border"]}; border-radius:12px;
    background:{theme["hero"]};
}}
.hero-main, .hero-side {{ position:relative; z-index:1; }}
.eyebrow {{
    display:inline-flex; align-items:center; gap:8px;
    border:1px solid {theme["border2"]}; border-radius:999px; padding:6px 10px;
    color:{theme["accent"]}; background:{theme["surface2"]};
    font-size:.72rem; font-weight:800; text-transform:uppercase; letter-spacing:.11em;
}}
.hero h1 {{
    margin:16px 0 10px; max-width:820px; color:{theme["text"]};
    font-size:clamp(1.6rem, 2.4vw, 2.4rem); line-height:1.1; font-weight:800; letter-spacing:-0.01em;
}}
.hero p {{
    margin:0; max-width:760px; color:{theme["muted"]}; font-size:1rem; line-height:1.65;
}}
.hero-actions {{ display:flex; gap:10px; flex-wrap:wrap; margin-top:22px; }}
.hero-chip {{
    display:inline-flex; align-items:center; min-height:38px; padding:8px 12px;
    border-radius:12px; border:1px solid {theme["border"]}; background:{theme["surface2"]};
    color:{theme["text"]}; font-size:.78rem; font-weight:800;
}}
.hero-chip.accent {{ color:white; border-color:transparent; background:{theme["accent"]}; }}
.workspace-compact {{
    border:1px solid {theme["border"]}; border-radius:22px; padding:22px 24px;
    background:{theme["surface"]}; box-shadow:{theme["shadow"]}; margin-bottom:18px;
}}
.workspace-compact h1 {{
    margin:12px 0 6px; color:{theme["text"]}; font-size:2.2rem; line-height:1.05; letter-spacing:0; font-weight:850;
}}
.workspace-compact p {{ margin:0; color:{theme["muted"]}; font-size:.96rem; line-height:1.55; max-width:900px; }}
.hero-side {{
    border:1px solid {theme["border"]}; border-radius:18px; background:{theme["surface"]};
    padding:16px; display:flex; flex-direction:column; gap:12px; min-height:250px;
}}
.signal-row {{ display:flex; justify-content:space-between; align-items:center; gap:12px; }}
.signal-label {{ color:{theme["muted"]}; font-size:.74rem; font-weight:800; text-transform:uppercase; letter-spacing:.08em; }}
.signal-value {{ font-family:'JetBrains Mono', monospace; font-size:.86rem; color:{theme["text"]}; }}
.signal-line {{ height:8px; border-radius:999px; background:{theme["panel2"]}; overflow:hidden; border:1px solid {theme["border"]}; }}
.signal-fill {{ height:100%; border-radius:inherit; background:{theme["accent"]}; }}
.side-brand {{
    padding:14px; border:1px solid {theme["border"]}; border-radius:16px; background:{theme["surface2"]}; margin-bottom:10px;
}}
.side-brand strong {{ display:block; font-size:1rem; }}
.side-brand span {{ display:block; color:{theme["muted"]}; font-size:.76rem; margin-top:3px; line-height:1.35; }}

.kpi-row {{ display:grid; grid-template-columns: repeat(5, minmax(0, 1fr)); gap:12px; margin-bottom:18px; }}
.kpi {{
    background:{theme["surface"]}; border:1px solid {theme["border"]}; border-radius:10px;
    padding:16px 14px; min-height:102px; text-align:left; transition:border-color .18s ease;
}}
.kpi:hover {{ border-color:{theme["border2"]}; }}
.kpi .n {{ font-family:'JetBrains Mono', monospace; font-size:1.55rem; font-weight:800; color:{theme["accent2"]}; }}
.kpi .n.g {{ color:{theme["good"]}; }}
.kpi .n.p {{ color:{theme["accent"]}; }}
.kpi .n.a {{ color:{theme["accent3"]}; }}
.kpi .l {{ font-size:.68rem; color:{theme["muted"]}; text-transform:uppercase; letter-spacing:.08em; margin-top:8px; font-weight:800; }}

.pipe {{ display:grid; grid-template-columns: repeat(5, minmax(0, 1fr)); gap:8px; margin:16px 0; }}
.ps {{
    text-align:center; padding:10px 8px; min-height:58px;
    background:{theme["surface"]}; border:1px solid {theme["border"]}; border-radius:14px;
    font-size:.72rem; color:{theme["muted"]}; transition:all .25s ease;
}}
.ps.done {{ border-color:{theme["border2"]}; color:{theme["good"]}; background:{theme["chip_validated_bg"]}; }}
.ps.active {{ border-color:{theme["accent"]}; color:{theme["accent"]}; background:{theme["active_row"]}; }}
.ps .ic {{ font-size:1.05rem; display:block; margin-bottom:2px; }}

.bok, .bwn, .binf {{
    display:inline-flex; align-items:center; min-height:26px; padding:4px 10px; border-radius:999px;
    font-size:.72rem; font-weight:800; border:1px solid {theme["border"]}; background:{theme["surface2"]};
}}
.bok {{ color:{theme["good"]}; border-color:{theme["border2"]}; }}
.bwn {{ color:{theme["warn"]}; border-color:{theme["border2"]}; }}
.binf {{ color:{theme["accent2"]}; }}
.cb {{ height:7px; border-radius:999px; background:{theme["panel2"]}; overflow:hidden; margin:5px 0; border:1px solid {theme["border"]}; }}
.cf {{ height:100%; border-radius:inherit; transition:width .6s ease; }}
.cf.h {{ background:{theme["good"]}; }}
.cf.m {{ background:{theme["warn"]}; }}
.cf.lo {{ background:{theme["danger"]}; }}

.dropzone {{
    border:1.5px dashed {theme["border2"]}; border-radius:10px; padding:48px 28px;
    text-align:center; margin:16px 0 18px; transition:border-color .25s ease;
    background:{theme["surface"]};
}}
.dropzone:hover {{ border-color:{theme["accent"]}; }}
.dropzone .icon {{ font-size:2.4rem; margin-bottom:8px; }}
.dropzone .title {{ color:{theme["text"]}; font-weight:850; font-size:1.15rem; }}
.dropzone .sub {{ color:{theme["muted"]}; font-size:.86rem; margin-top:6px; }}
.capability-grid {{
    display:grid; grid-template-columns:repeat(4, minmax(0,1fr)); gap:12px; margin: 12px 0 20px;
}}
.capability {{
    border:1px solid {theme["border"]}; border-radius:16px; padding:14px; background:{theme["surface"]};
}}
.capability strong {{ display:block; color:{theme["text"]}; font-size:.88rem; margin-bottom:4px; }}
.capability span {{ display:block; color:{theme["muted"]}; font-size:.76rem; line-height:1.4; }}
.setup-banner {{
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 16px;
    margin: 0 0 10px;
    padding: 14px 18px;
    border: 1px solid {theme["border"]};
    border-radius: 8px;
    background: {theme["active_row"]};
}}
.setup-banner.slim {{
    margin-top: 8px;
}}
.setup-banner strong {{
    color: {theme["text"]};
    font-size: .95rem;
    font-weight: 850;
}}
.setup-banner span {{
    color: {theme["muted"]};
    font-size: .82rem;
    line-height: 1.4;
    text-align: right;
}}
.integration-hero {{
    min-height: 112px;
    display: grid;
    grid-template-columns: 48px minmax(0, 1fr);
    align-items: center;
    gap: 18px;
    padding: 20px 28px;
    border-bottom: 1px solid {theme["border"]};
    background: {theme["panel"]};
}}
.integration-brand-mark {{
    width: 48px;
    height: 48px;
    display: grid;
    place-items: center;
    border-radius: 8px;
    background: {theme["accent"]};
    color: white;
    font-size: .82rem;
    font-weight: 900;
}}
.integration-hero-copy .integration-kind,
.integration-section-heading span,
.integration-readiness span {{
    display: block;
    margin-bottom: 6px;
    color: {theme["faint"]};
    font-size: .68rem;
    font-weight: 850;
    text-transform: uppercase;
    letter-spacing: .08em;
}}
.integration-title-row {{
    display: flex;
    align-items: center;
    justify-content: flex-start;
    gap: 12px;
    flex-wrap: wrap;
}}
.integration-system-name {{
    display: block;
    color: {theme["text"]};
    font-size: 1.28rem;
    line-height: 1.2;
    font-weight: 820;
}}
.integration-hero-copy p {{
    max-width: 900px;
    margin: 6px 0 0;
    color: {theme["muted"]};
    font-size: .8rem;
    line-height: 1.45;
}}
.integration-status {{
    display: inline-flex;
    align-items: center;
    gap: 8px;
    min-height: 34px;
    padding: 7px 12px;
    border: 1px solid {theme["border"]};
    border-radius: 999px;
    color: {theme["muted"]};
    font-size: .74rem;
    font-weight: 850;
    white-space: nowrap;
}}
.integration-status i {{
    width: 8px;
    height: 8px;
    border-radius: 50%;
    background: {theme["faint"]};
}}
.integration-status.ready i {{ background: {theme["good"]}; }}
.integration-status.export i {{ background: {theme["warn"]}; }}
.integration-status.setup i {{ background: {theme["danger"]}; }}
.integration-flow {{
    display: grid;
    grid-template-columns: repeat(4, minmax(0, 1fr));
    gap: 0;
    margin: 18px 0 24px;
    border: 1px solid {theme["border"]};
    border-radius: 8px;
    overflow: hidden;
    background: {theme["surface"]};
}}
.integration-flow > div {{
    display: flex;
    align-items: center;
    gap: 12px;
    min-height: 70px;
    padding: 13px 16px;
    border-right: 1px solid {theme["border"]};
}}
.integration-flow > div:last-child {{ border-right: 0; }}
.integration-flow b {{
    flex: 0 0 28px;
    width: 28px;
    height: 28px;
    display: grid;
    place-items: center;
    border-radius: 50%;
    background: {theme["active_row"]};
    color: {theme["accent"]};
    font-size: .72rem;
}}
.integration-flow strong,
.integration-flow small {{ display: block; }}
.integration-flow strong {{ color: {theme["text"]}; font-size: .78rem; }}
.integration-flow small {{ margin-top: 3px; color: {theme["muted"]}; font-size: .67rem; line-height: 1.3; }}
.integration-section-heading {{
    display: flex;
    align-items: end;
    justify-content: space-between;
    gap: 24px;
    padding: 2px 2px 14px;
    border-bottom: 1px solid {theme["border"]};
    margin-bottom: 16px;
}}
.integration-section-heading strong {{
    display: block;
    color: {theme["text"]};
    font-size: 1rem;
}}
.integration-section-heading p {{
    max-width: 560px;
    margin: 0;
    color: {theme["muted"]};
    font-size: .76rem;
    line-height: 1.45;
    text-align: right;
}}
.integration-readiness {{
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 24px;
    padding: 22px 24px;
    margin-bottom: 24px;
    border: 1px solid {theme["border"]};
    border-radius: 8px;
    background: {theme["surface"]};
}}
.integration-readiness h3 {{ margin: 0; color: {theme["text"]}; font-size: 1.08rem; }}
.integration-readiness p {{ margin: 7px 0 0; color: {theme["muted"]}; font-size: .78rem; line-height: 1.45; }}
.integration-readiness-badge {{
    flex: 0 0 auto;
    padding: 8px 12px;
    border-radius: 999px;
    background: {theme["chip_pending_bg"]};
    color: {theme["warn"]};
    font-size: .72rem;
    font-weight: 850;
}}
.integration-requirements {{
    display: grid;
    grid-template-columns: repeat(3, minmax(0, 1fr));
    gap: 12px;
    margin-bottom: 16px;
}}
.integration-requirements > div {{
    min-height: 82px;
    padding: 16px;
    border: 1px solid {theme["border"]};
    border-radius: 8px;
    background: {theme["surface"]};
}}
.integration-requirements b {{
    display: block;
    margin-bottom: 10px;
    color: {theme["accent"]};
    font-size: .68rem;
}}
.integration-requirements strong {{
    color: {theme["text"]};
    font-size: .78rem;
    line-height: 1.4;
}}
div[class*="st-key-quickbooks_connector_workspace"],
div[class*="st-key-zoho_connector_workspace"] {{
    padding: 18px 20px 20px;
    border: 1px solid {theme["border"]};
    border-radius: 8px;
    background: {theme["surface"]};
}}
div[data-testid="stTabs"] [data-baseweb="tab-list"] {{
    gap: 2px;
    padding: 4px;
    border: 1px solid {theme["border"]};
    border-radius: 8px;
    background: {theme["surface"]};
}}
div[data-testid="stTabs"] [data-baseweb="tab"] {{
    min-height: 38px;
    flex: 1 1 0;
    justify-content: center;
    padding: 0 16px !important;
    border-radius: 6px;
}}
div[data-testid="stTabs"] [aria-selected="true"] {{
    background: {theme["active_row"]};
}}
.connector-help {{
    color: {theme["muted"]};
    font-size: .82rem;
    line-height: 1.45;
    margin: 0 0 12px;
}}
.page-card {{
    border:1px solid {theme["border"]}; border-radius:20px; padding:20px 22px;
    margin: 4px 0 18px; background:{theme["surface"]}; box-shadow:{theme["shadow"]};
}}
.page-card h2 {{
    margin:12px 0 0; color:{theme["text"]}; font-size:1.75rem; line-height:1.15; letter-spacing:0;
}}
.timeline-card {{
    border:1px solid {theme["border"]}; border-radius:18px; background:{theme["surface"]};
    padding:16px; display:flex; flex-direction:column; gap:12px;
}}
.timeline-card div {{
    border-left:3px solid {theme["accent"]}; padding:4px 0 4px 12px;
}}
.timeline-card strong {{ display:block; color:{theme["text"]}; font-size:.9rem; }}
.timeline-card span {{ display:block; color:{theme["muted"]}; font-size:.76rem; margin-top:2px; }}

div[data-testid="stExpander"] {{
    border:1px solid {theme["border"]} !important; border-radius:10px !important; margin-bottom:10px;
    background:{theme["surface"]} !important;
}}
div[data-testid="stExpander"] summary {{ font-weight:750 !important; color:{theme["text"]} !important; }}
button[data-baseweb="tab"] {{ font-weight:750 !important; color:{theme["muted"]} !important; }}
button[aria-selected="true"][data-baseweb="tab"] {{ color:{theme["accent"]} !important; }}
div[data-testid="stDataFrame"] {{ border-radius:14px !important; overflow:hidden; border:1px solid {theme["border"]}; }}
div[data-testid="stMetric"] {{
    background:{theme["surface"]}; border:1px solid {theme["border"]}; border-radius:16px; padding:12px 16px;
}}
div[data-testid="stDownloadButton"] > button, div[data-testid="stButton"] > button {{
    border:1px solid {theme["border2"]} !important; border-radius:12px !important;
    font-weight:800 !important; min-height:42px;
}}
div[data-testid="stFileUploader"] section {{
    background:{theme["surface"]} !important; border:1px solid {theme["border"]} !important;
    border-radius:14px !important; color:{theme["text"]} !important;
}}
div[data-testid="stFileUploader"] section * {{ color:{theme["muted"]} !important; }}
div[data-testid="stButton"] > button[kind="primary"], div[data-testid="stFileUploader"] button {{
    border-color:transparent !important; background:{theme["accent"]} !important; color:white !important;
}}
div[data-testid="stFileUploader"] button * {{ color:white !important; }}
div[data-testid="stDownloadButton"] > button:hover, div[data-testid="stButton"] > button:hover {{
    border-color:{theme["accent"]} !important;
}}
.stMarkdown code {{ padding:2px 7px; border-radius:6px; background:{theme["surface2"]}; color:{theme["accent2"]}; }}
div[data-testid="stAlert"] {{ border-radius:14px !important; }}
input, textarea, [data-baseweb="select"] > div {{
    background:{theme["input"]} !important; color:{theme["text"]} !important; border-color:{theme["border"]} !important;
}}
label, p, span, div {{ letter-spacing:0; }}
::-webkit-scrollbar {{ width:8px; height:8px; }}
::-webkit-scrollbar-thumb {{ background:{theme["border2"]}; border-radius:999px; }}
::-webkit-scrollbar-thumb:hover {{ background:{theme["accent"]}; }}
@media (max-width: 980px) {{
    .hero {{ grid-template-columns:1fr; padding:20px; }}
    .kpi-row, .capability-grid {{ grid-template-columns:repeat(2, minmax(0,1fr)); }}
    .ez-topbar {{ align-items:flex-start; flex-direction:column; }}
    .integration-flow {{ grid-template-columns: repeat(2, minmax(0, 1fr)); }}
    .integration-flow > div:nth-child(2) {{ border-right: 0; }}
    .integration-flow > div:nth-child(-n+2) {{ border-bottom: 1px solid {theme["border"]}; }}
    .integration-requirements {{ grid-template-columns: 1fr; }}
}}
@media (max-width: 640px) {{
    .kpi-row, .capability-grid, .pipe {{ grid-template-columns:1fr; }}
    .hero h1 {{ font-size:2rem; }}
    .dropzone {{ padding:34px 18px; }}
    .integration-hero {{
        grid-template-columns: 44px minmax(0, 1fr);
        padding: 20px 16px;
    }}
    .integration-brand-mark {{ width: 44px; height: 44px; }}
    .integration-title-row {{ align-items: flex-start; flex-direction: column; gap: 8px; }}
    .integration-flow {{ grid-template-columns: 1fr; }}
    .integration-flow > div {{
        border-right: 0;
        border-bottom: 1px solid {theme["border"]};
    }}
    .integration-flow > div:last-child {{ border-bottom: 0; }}
    .integration-section-heading,
    .integration-readiness {{ align-items: flex-start; flex-direction: column; }}
    .integration-section-heading p {{ text-align: left; }}
}}

/* Warm olive AP console shell */
:root {{ color-scheme: {"light" if is_light_mode else "dark"}; }}
html, body, .stApp {{
    background: {theme["bg"]} !important;
    color: {theme["text"]} !important;
    background-image: none !important;
    overflow-x: hidden !important;
}}
main, .block-container {{
    overflow-x: hidden !important;
}}
body, .stApp, input, textarea, button {{
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', system-ui, sans-serif !important;
}}
.block-container {{
    max-width: 1440px;
    padding-top: 1rem;
    padding-left: 1.5rem;
    padding-right: 1.5rem;
    padding-bottom: 2rem;
}}
header[data-testid="stHeader"] {{
    height: 0 !important;
    visibility: hidden !important;
}}
#MainMenu, footer {{
    visibility: hidden !important;
}}
section[data-testid="stSidebar"] {{
    width: 260px !important;
    min-width: 260px !important;
    max-width: 260px !important;
    background: {theme["surface"]} !important;
    border-right: 1px solid {theme["border"]};
    box-shadow: none;
}}
section[data-testid="stSidebar"] > div:first-child {{
    width: 260px !important;
    padding: 22px 14px 18px 14px;
}}
section[data-testid="stSidebar"] * {{
    color: {theme["text"]};
}}
.side-brand {{
    display: flex;
    align-items: center;
    gap: 12px;
    padding: 4px 8px 22px;
    border: 0;
    background: transparent;
}}
.side-mark {{
    width: 36px;
    height: 36px;
    border-radius: 10px;
    display: grid;
    place-items: center;
    color: #FFFFFF;
    background: {theme["accent"]};
    font-weight: 800;
    font-size: .75rem;
}}
.side-brand strong {{
    display: block;
    color: {theme["text"]};
    font-size: 1rem;
    line-height: 1.05;
}}
.side-brand span {{
    display: block;
    color: {theme["faint"]};
    font-size: .72rem;
    margin-top: 4px;
}}
section[data-testid="stSidebar"] h3 {{
    margin: 1rem 0 .45rem;
    color: {theme["faint"]} !important;
    font-size: .72rem !important;
    font-weight: 700 !important;
    text-transform: uppercase;
    letter-spacing: .12em !important;
}}
section[data-testid="stSidebar"] div[role="radiogroup"] label {{
    min-height: 44px;
    width: 100%;
    border-radius: 8px;
    padding: 0 10px;
    margin-bottom: 3px;
    border: 1px solid transparent;
    background: transparent;
}}
section[data-testid="stSidebar"] div[data-testid="stRadio"],
section[data-testid="stSidebar"] div[data-testid="stElementContainer"]:has(div[data-testid="stRadio"]),
section[data-testid="stSidebar"] div[role="radiogroup"],
section[data-testid="stSidebar"] div[role="radiogroup"] > div {{
    width: 100% !important;
    max-width: 100% !important;
}}
section[data-testid="stSidebar"] div[role="radiogroup"] label,
section[data-testid="stSidebar"] div[role="radiogroup"] label > div {{
    display: flex;
    align-items: center;
    width: 100% !important;
    max-width: 100% !important;
    box-sizing: border-box;
}}
section[data-testid="stSidebar"] div[role="radiogroup"] label > div:first-child {{
    display: none !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_theme_mode"] div[role="radiogroup"] {{
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 8px;
}}
section[data-testid="stSidebar"] div[role="radiogroup"] label:has(input:checked) {{
    background: {theme["input"] if is_light_mode else theme["active_row"]};
    border-color: {theme["border"]};
    box-shadow: inset 3px 0 0 {theme["accent"]};
}}
section[data-testid="stSidebar"] div[role="radiogroup"] label:has(input:checked) p {{
    color: {theme["text"]} !important;
    font-weight: 800 !important;
}}
section[data-testid="stSidebar"] label p,
section[data-testid="stSidebar"] p,
section[data-testid="stSidebar"] span {{
    font-size: .9rem !important;
}}
.ez-shell {{
    gap: 0;
    margin-bottom: 12px;
}}
.ez-topbar {{
    height: 54px;
    min-height: 54px;
    border-radius: 8px 8px 0 0;
    border: 1px solid {theme["border"]};
    background: {theme["surface"]};
    box-shadow: none;
    padding: 0 28px;
}}
.top-title {{
    color: {theme["text"]};
    font-size: 1.03rem;
    font-weight: 800;
}}
.ez-top-pills {{
    gap: 12px;
}}
.ez-search, .ez-filter, .ez-upload-pill {{
    height: 36px;
    border-radius: 8px;
    display: inline-flex;
    align-items: center;
    gap: 9px;
    border: 1px solid {theme["border"]};
    background: {theme["input"]};
    color: {theme["faint"]};
    font-size: .86rem;
    font-weight: 650;
}}
.ez-search {{
    width: clamp(210px, 24vw, 320px);
    padding: 0 14px;
    justify-content: flex-start;
}}
.search-box {{
    width: 12px;
    height: 12px;
    border: 1.5px solid {theme["faint"]};
    border-radius: 2px;
    opacity: .85;
}}
.ez-filter {{
    padding: 0 18px;
}}
.ez-upload-pill {{
    padding: 0 22px;
    border-color: transparent;
    background: {theme["accent"]};
    color: #FFFFFF !important;
    font-weight: 800;
    text-decoration: none !important;
}}
a.ez-upload-pill,
a.ez-upload-pill:visited,
a.ez-upload-pill:hover,
a.ez-upload-pill:active {{
    color: #FFFFFF !important;
    text-decoration: none !important;
}}
.workspace-compact, .page-card {{
    border-radius: 0 0 8px 8px;
    background: {theme["surface"]};
    border: 1px solid {theme["border"]};
    border-top: 0;
    box-shadow: none;
}}
.workspace-compact {{
    padding: 20px 28px;
    margin-bottom: 18px;
}}
.workspace-compact h1 {{
    font-size: clamp(1.5rem, 2vw, 2rem);
    font-weight: 800;
    letter-spacing: 0;
}}
.eyebrow {{
    background: transparent;
    color: {theme["faint"]};
    border: 0;
    padding: 0;
    letter-spacing: .12em;
}}
.kpi {{
    background: {theme["surface"]};
    border-color: {theme["border"]};
    border-radius: 8px;
    box-shadow: none;
    min-height: 92px;
}}
.kpi:hover {{
    transform: none;
    box-shadow: none;
}}
.capability, .timeline-card, div[data-testid="stExpander"], div[data-testid="stMetric"] {{
    background: {theme["surface"]} !important;
    border-color: {theme["border"]} !important;
    border-radius: 8px !important;
    box-shadow: none !important;
}}
.dropzone {{
    background: {theme["surface"]};
    border-color: {theme["border"]};
    border-radius: 8px;
    box-shadow: none;
    padding: 26px 24px;
}}
div[data-testid="stFileUploader"] section {{
    background: {theme["surface"]} !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 8px !important;
}}
div[data-testid="stFileUploader"] button,
div[data-testid="stButton"] > button[kind="primary"] {{
    background: {theme["accent"]} !important;
    color: #FFFFFF !important;
    border-color: transparent !important;
}}
div[data-testid="stFileUploader"] button *,
div[data-testid="stButton"] > button[kind="primary"] * {{
    color: #FFFFFF !important;
}}
div[data-testid="stButton"] > button,
div[data-testid="stDownloadButton"] > button {{
    border-radius: 8px !important;
    min-height: 40px;
    border-color: {theme["border"]} !important;
    box-shadow: none !important;
    font-weight: 800 !important;
}}
div[data-testid="stDataFrame"] {{
    border-radius: 9px !important;
    max-width: 100% !important;
    overflow-x: auto !important;
    overflow-y: hidden !important;
    border: 1px solid {theme["border"]};
}}
div[data-testid="stDataFrame"] table {{
    border-collapse: collapse !important;
}}
input, textarea, [data-baseweb="select"] > div {{
    background: {theme["input"]} !important;
    color: {theme["text"]} !important;
    border-color: {theme["border"]} !important;
    border-radius: 8px !important;
}}
label, p, span, div {{
    letter-spacing: 0;
}}

/* Three-pane AP console inspired by operator review workspaces */
.ap-status-strip {{
    display: grid;
    grid-template-columns: repeat(5, minmax(0, 1fr));
    height: 74px;
    border: 1px solid {theme["border"]};
    border-radius: 0;
    background: {theme["surface"]};
    overflow: hidden;
    margin-bottom: 0;
}}
.ap-status {{
    min-height: 74px;
    display: flex;
    align-items: center;
    justify-content: center;
    gap: 10px;
    border-right: 1px solid {theme["border"]};
    color: {theme["muted"]};
}}
.ap-status:last-child {{ border-right: 0; }}
.ap-status strong {{
    font-family: 'JetBrains Mono', ui-monospace, SFMono-Regular, Menlo, monospace;
    font-size: 24px;
    line-height: 1;
    font-weight: 800;
    color: {theme["accent"]};
}}
.ap-status:nth-child(4) strong {{
    color: {theme["warn"]};
}}
.ap-status span {{
    font-weight: 700;
    font-size: .92rem;
}}
.ap-status.active {{
    box-shadow: inset 0 -3px 0 {theme["accent"]};
    color: {theme["text"]};
}}
.ap-status.danger strong {{
    color: {theme["danger"]};
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="column"]:first-child {{
    flex: 0 0 clamp(320px, 28vw, 400px) !important;
    max-width: clamp(320px, 28vw, 400px) !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="column"]:nth-child(2) {{
    flex: 1 1 auto !important;
    min-width: 0 !important;
}}
.pane-title {{
    height: 48px;
    display: flex;
    justify-content: space-between;
    align-items: center;
    padding: 0 24px;
    border: 1px solid {theme["border"]};
    border-top: 0;
    background: {theme["bg"]};
    color: {theme["muted"]};
    font-weight: 700;
}}
.invoice-list-row {{
    min-height: 100px;
    display: flex;
    justify-content: space-between;
    gap: 16px;
    padding: 18px 28px;
    border-left: 1px solid {theme["border"]};
    border-right: 1px solid {theme["border"]};
    border-bottom: 1px solid {theme["border"]};
    background: {theme["bg"]};
}}
.invoice-list-row.active {{
    border-left: 4px solid {theme["accent"]};
    background: {theme["active_row"]};
}}
.invoice-list-row strong {{
    display: block;
    color: {theme["text"]};
    font-size: 1.06rem;
    letter-spacing: 0;
}}
.invoice-list-row span {{
    display: block;
    color: {theme["muted"]};
    margin-top: 7px;
    font-size: .9rem;
}}
.invoice-list-row em {{
    display: block;
    color: {theme["faint"]};
    margin-top: 11px;
    font-style: normal;
    font-size: .82rem;
}}
.invoice-row-meta {{
    display: flex;
    flex-direction: column;
    align-items: flex-end;
    gap: 20px;
    white-space: nowrap;
}}
.invoice-row-meta b {{
    color: {theme["text"]};
    font-family: 'JetBrains Mono', ui-monospace, SFMono-Regular, Menlo, monospace;
    font-size: 1rem;
}}
.status-pill {{
    display: inline-flex;
    align-items: center;
    justify-content: center;
    min-height: 26px;
    padding: 4px 10px;
    border-radius: 999px;
    font-size: .78rem;
    font-weight: 800;
    border: 1px solid transparent;
    background: {theme["surface"]};
    color: {theme["muted"]};
}}
.status-pill.sent {{
    color: {theme["accent"]};
    background: {theme["chip_sent_bg"]};
    border-color: {theme["chip_sent_bg"]};
}}
.status-pill.ok {{
    color: {theme["accent"]};
    background: {theme["chip_validated_bg"]};
    border-color: {theme["chip_validated_bg"]};
}}
.status-pill.bad {{
    color: {theme["danger"]};
    background: {theme["chip_exception_bg"]};
    border-color: {theme["chip_exception_bg"]};
}}
.status-pill.warn {{
    color: {theme["warn"]};
    background: {theme["chip_pending_bg"]};
    border-color: {theme["chip_pending_bg"]};
}}
.status-pill.neutral {{
    color: {theme["muted"]};
    background: {theme["surface"]};
    border-color: {theme["border"]};
}}
.detail-head {{
    display: flex;
    align-items: flex-start;
    justify-content: space-between;
    gap: 24px;
    padding: 28px 36px 22px;
    border: 0;
    border-bottom: 1px solid {theme["border"]};
    border-radius: 0;
    background: {theme["bg"]};
}}
.detail-head > div:first-child {{
    min-width: 0;
    flex: 1 1 auto;
}}
.detail-head h2 {{
    margin: 10px 0 8px;
    color: {theme["text"]};
    font-size: clamp(22px, 2.2vw, 30px);
    line-height: 1.05;
    letter-spacing: 0;
}}
.detail-head p {{
    color: {theme["muted"]};
    margin: 0;
    font-size: 1rem;
    line-height: 1.35;
}}
.detail-actions {{
    display: flex;
    flex-direction: row;
    align-items: center;
    flex-wrap: wrap;
    justify-content: flex-end;
    gap: 8px;
    flex: 0 0 auto;
    padding-top: 2px;
}}
.detail-grid {{
    display: grid;
    grid-template-columns: repeat(auto-fit, minmax(140px, 1fr));
    gap: 14px;
    padding: 18px 0;
}}
.detail-grid div {{
    min-height: 72px;
    border: 1px solid {theme["border"]};
    border-radius: 8px;
    padding: 18px;
    background: {theme["surface"]};
}}
.detail-grid span {{
    display: block;
    color: {theme["faint"]};
    font-size: .72rem;
    font-weight: 800;
    text-transform: uppercase;
    letter-spacing: .08em;
}}
.detail-grid strong {{
    display: block;
    color: {theme["text"]};
    margin-top: 12px;
    font-size: 1.15rem;
    line-height: 1.1;
}}
.empty-pane {{
    min-height: 220px;
    border: 1px solid {theme["border"]};
    border-top: 0;
    background: {theme["bg"]};
    display: grid;
    place-content: center;
    text-align: center;
    padding: 24px;
}}
.empty-pane.large {{
    min-height: 420px;
    border-top: 1px solid {theme["border"]};
    border-radius: 0 0 8px 8px;
}}
.empty-pane strong {{
    color: {theme["text"]};
    font-size: 1.15rem;
}}
.empty-pane span {{
    color: {theme["muted"]};
    margin-top: 8px;
}}
@media (max-width: 980px) {{
    .block-container {{ padding-left: 1rem; padding-right: 1rem; }}
    .ez-topbar {{ height: auto; min-height: 54px; align-items: flex-start; flex-direction: column; padding: 12px 16px; }}
    .ez-top-pills {{ width: 100%; justify-content: flex-start; }}
    .ez-search {{ width: 100%; }}
    .ap-status-strip, .detail-grid {{ grid-template-columns: 1fr; height: auto; }}
    .ap-status {{ justify-content: flex-start; padding-left: 18px; border-right: 0; border-bottom: 1px solid {theme["border"]}; }}
    div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="column"]:first-child {{
        flex: 1 1 auto !important;
        max-width: none !important;
    }}
    .detail-head {{ flex-direction: column; }}
    .detail-actions {{ align-items: flex-start; }}
}}

/* Reference shell override from ez_app.py */
:root {{
    --ez-sidebar-w: 260px;
    --ez-topbar-h: 84px;
    --ez-status-h: 74px;
    --ez-sidebar-footer-h: 113px;
    --ez-head-h: calc(var(--ez-topbar-h) + var(--ez-status-h));
    --ez-queue-w: clamp(320px, 28vw, 400px);
}}
.block-container {{
    padding: 0 !important;
    max-width: none !important;
}}
div[data-testid="stMainBlockContainer"], section[data-testid="stMain"] > div {{
    padding-top: 0 !important;
    margin-top: 0 !important;
}}
section[data-testid="stMain"] {{
    padding-top: 0 !important;
}}
section[data-testid="stMain"] div[data-testid="stVerticalBlock"] {{
    gap: 0 !important;
}}
header[data-testid="stHeader"],
div[data-testid="stToolbar"],
div[data-testid="stDecoration"],
div[data-testid="stStatusWidget"],
div[data-testid="stHeaderActionElements"],
div[data-testid="stAppToolbar"],
div[data-testid="collapsedControl"],
div[data-testid="stSidebarCollapseButton"],
#MainMenu,
footer {{
    display: none !important;
    visibility: hidden !important;
    height: 0 !important;
}}
section[data-testid="stSidebar"] {{
    width: var(--ez-sidebar-w) !important;
    min-width: var(--ez-sidebar-w) !important;
    max-width: var(--ez-sidebar-w) !important;
    background: {theme["surface"]} !important;
    border-right: 1px solid {theme["border"]} !important;
    box-shadow: none !important;
}}
section[data-testid="stSidebar"] > div {{
    padding: 0 !important;
}}
section[data-testid="stSidebar"] div[data-testid="stSidebarUserContent"] {{
    position: absolute !important;
    top: 0 !important;
    left: 0 !important;
    margin-top: 0 !important;
    padding: 0 !important;
    height: 100vh !important;
    max-height: 100vh !important;
    width: 100% !important;
}}
section[data-testid="stSidebar"] div[data-testid="stSidebarUserContent"] > div[data-testid="stVerticalBlock"] > div:first-child {{
    margin-top: 0 !important;
    padding-top: 0 !important;
}}
section[data-testid="stSidebar"] div[data-testid="stSidebarUserContent"] > div[data-testid="stVerticalBlock"] {{
    height: 100% !important;
    min-height: 100% !important;
    display: flex !important;
    flex-direction: column !important;
    overflow: hidden !important;
}}
section[data-testid="stSidebar"],
section[data-testid="stSidebar"] > div {{
    height: 100vh !important;
    max-height: 100vh !important;
    overflow: hidden !important;
}}
section[data-testid="stSidebar"] div[data-testid="stVerticalBlock"] {{
    gap: 0 !important;
}}
.ez-brand {{
    display: flex !important;
    align-items: center !important;
    gap: 14px !important;
    height: var(--ez-topbar-h) !important;
    min-height: var(--ez-topbar-h) !important;
    padding: 0 24px !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
section[data-testid="stSidebar"] div[data-testid="stElementContainer"]:has(.ez-brand),
section[data-testid="stSidebar"] div[data-testid="stMarkdown"]:has(.ez-brand),
section[data-testid="stSidebar"] div[data-testid="stMarkdownContainer"]:has(.ez-brand) {{
    height: var(--ez-topbar-h) !important;
    min-height: var(--ez-topbar-h) !important;
    margin: 0 !important;
}}
.ez-brand-mark {{
    width: 36px !important;
    height: 36px !important;
    border-radius: 10px !important;
    background: {theme["accent"]} !important;
    display: grid !important;
    place-items: center !important;
    color: #fff !important;
    font-weight: 800 !important;
    font-size: 17px !important;
}}
.ez-brand-name {{
    font-size: 15px !important;
    font-weight: 800 !important;
    color: {theme["text"]} !important;
    letter-spacing: 0 !important;
}}
.ez-brand-ver {{
    font-size: 10px !important;
    color: {theme["faint"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    margin-top: 2px !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_nav_shell"] {{
    flex: 1 1 auto !important;
    min-height: 0 !important;
    height: calc(100vh - var(--ez-topbar-h) - var(--ez-sidebar-footer-h)) !important;
    max-height: calc(100vh - var(--ez-topbar-h) - var(--ez-sidebar-footer-h)) !important;
    overflow-y: auto !important;
    overflow-x: hidden !important;
    padding: 16px 12px 18px !important;
    scrollbar-width: thin !important;
    scrollbar-color: {theme["border"]} transparent !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_nav_shell"] > div[data-testid="stVerticalBlock"] {{
    gap: 0 !important;
}}
.nav-section {{
    display: block !important;
    height: 18px !important;
    line-height: 18px !important;
    margin: 18px 12px 10px !important;
    color: {theme["faint"]} !important;
    font-size: 10px !important;
    letter-spacing: .12em !important;
    font-weight: 800 !important;
    text-transform: uppercase !important;
}}
.nav-section:first-child {{
    margin-top: 8px !important;
}}
section[data-testid="stSidebar"] div[data-testid="stElementContainer"]:has(.nav-section),
section[data-testid="stSidebar"] div[data-testid="stMarkdown"]:has(.nav-section),
section[data-testid="stSidebar"] div[data-testid="stMarkdownContainer"]:has(.nav-section) {{
    min-height: 46px !important;
    height: 46px !important;
    margin: 0 !important;
    overflow: visible !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-nav_item_"] {{
    position: relative !important;
    height: 44px !important;
    min-height: 44px !important;
    margin: 0 !important;
    padding: 0 !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-nav_item_"] > div[data-testid="stVerticalBlock"] {{
    gap: 0 !important;
    height: 44px !important;
    min-height: 44px !important;
}}
.side-nav-row {{
    display: grid !important;
    grid-template-columns: 24px minmax(0, 1fr) auto !important;
    align-items: center !important;
    height: 40px !important;
    min-height: 40px !important;
    margin: 1px 0 !important;
    padding: 0 14px !important;
    border-radius: 8px !important;
    border-left: 3px solid transparent !important;
    column-gap: 12px !important;
    color: {theme["muted"]} !important;
    font-size: 14px !important;
    font-weight: 550 !important;
    pointer-events: none !important;
}}
.side-nav-row > span:nth-child(2) {{
    min-width: 0 !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
}}
.side-nav-row * {{
    pointer-events: none !important;
}}
.side-nav-row.active {{
    background: {theme["input"] if is_light_mode else theme["active_row"]} !important;
    color: {theme["accent"]} !important;
    border-left-color: {theme["accent"]} !important;
    font-weight: 800 !important;
}}
.side-icon {{
    display: grid !important;
    place-items: center !important;
    width: 24px !important;
    height: 24px !important;
    color: currentColor !important;
    font-size: 18px !important;
    font-weight: 700 !important;
}}
.side-nav-count {{
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 12px !important;
    color: {theme["faint"]} !important;
    font-weight: 700 !important;
    background: {theme["surface"]} !important;
    padding: 2px 8px !important;
    border-radius: 20px !important;
}}
.side-red-dot {{
    width: 9px !important;
    height: 9px !important;
    border-radius: 999px !important;
    background: {theme["danger"]} !important;
    justify-self: end !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-nav_item_"] div[data-testid="stButton"] {{
    position: absolute !important;
    inset: 0 !important;
    z-index: 5 !important;
    width: 100% !important;
    height: 44px !important;
    margin: 0 !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-nav_btn_"] {{
    position: absolute !important;
    inset: 0 !important;
    z-index: 8 !important;
    width: 100% !important;
    height: 44px !important;
    margin: 0 !important;
    padding: 0 !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-nav_btn_"] div[data-testid="stButton"] {{
    position: absolute !important;
    inset: 0 !important;
    width: 100% !important;
    height: 44px !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-nav_item_"] div[data-testid="stButton"] > button {{
    width: 100% !important;
    height: 44px !important;
    min-height: 44px !important;
    margin: 0 !important;
    padding: 0 !important;
    border: 0 !important;
    background: transparent !important;
    color: transparent !important;
    box-shadow: none !important;
    opacity: 0 !important;
    cursor: pointer !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_sidebar_footer"] {{
    position: fixed !important;
    left: 0 !important;
    bottom: 0 !important;
    width: var(--ez-sidebar-w) !important;
    box-sizing: border-box !important;
    padding: 12px 24px 16px !important;
    border-top: 1px solid {theme["border"]} !important;
    background: {theme["surface"]} !important;
    z-index: 1004 !important;
}}
.sidebar-status {{
    position: static !important;
    width: 100% !important;
    box-sizing: border-box !important;
    padding: 6px 0 0 !important;
    border-top: 0 !important;
    background: {theme["surface"]} !important;
    font-size: 13px !important;
    color: {theme["muted"]} !important;
}}
.sidebar-status-line {{
    display: flex !important;
    align-items: center !important;
    gap: 8px !important;
    margin-bottom: 5px !important;
}}
.status-dot {{
    display: inline-block !important;
    width: 7px !important;
    height: 7px !important;
    border-radius: 999px !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_dark_mode_toggle"] {{
    position: static !important;
    width: 100% !important;
    min-height: 34px !important;
    height: 34px !important;
    margin: 0 0 8px !important;
    padding: 0 !important;
    background: transparent !important;
    border: 0 !important;
    border-radius: 0 !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_dark_mode_toggle"] div[data-testid="stToggle"] {{
    width: 100% !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_dark_mode_toggle"] label {{
    width: 100% !important;
    min-height: 34px !important;
    height: 34px !important;
    display: flex !important;
    flex-direction: row-reverse !important;
    justify-content: space-between !important;
    align-items: center !important;
    gap: 14px !important;
    cursor: pointer !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_dark_mode_toggle"] p {{
    color: {theme["muted"]} !important;
    font-size: 13px !important;
    line-height: 1 !important;
    font-weight: 650 !important;
    margin: 0 !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_dark_mode_toggle"] div[role="switch"] {{
    flex: 0 0 auto !important;
    width: 36px !important;
    height: 20px !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-top-title) {{
    height: var(--ez-topbar-h) !important;
    min-height: var(--ez-topbar-h) !important;
    max-height: var(--ez-topbar-h) !important;
    padding: 0 32px !important;
    gap: 14px !important;
    border-bottom: 1px solid {theme["border"]} !important;
    align-items: center !important;
    background: {theme["bg"]} !important;
}}
div[data-testid="stLayoutWrapper"]:has(.ez-top-title) {{
    margin-top: 0 !important;
    margin-bottom: 0 !important;
}}
.ez-top-title .title {{
    font-size: 18px !important;
    font-weight: 800 !important;
    color: {theme["text"]} !important;
    letter-spacing: 0 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-top-title) div[data-testid="stTextInput"] input {{
    height: 32px !important;
    min-height: 32px !important;
    border-radius: 7px !important;
    font-size: 13px !important;
    background: {theme["input"]} !important;
    border: 1px solid {theme["border"]} !important;
    color: {theme["text"]} !important;
    padding: 4px 14px !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-top-title) div[data-testid="stTextInput"] input::placeholder {{
    color: {theme["faint"]} !important;
    opacity: 1 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-top-title) div[data-baseweb="input"] {{
    background: {theme["input"]} !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 7px !important;
    height: 32px !important;
    min-height: 32px !important;
    box-shadow: none !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-top-title) div[data-testid="stPopover"] button {{
    height: 32px !important;
    min-height: 32px !important;
    max-height: 32px !important;
    display: inline-flex !important;
    align-items: center !important;
    justify-content: center !important;
    gap: 6px !important;
    overflow: hidden !important;
    white-space: nowrap !important;
    line-height: 1 !important;
    border-radius: 7px !important;
    border: 1px solid {theme["border"]} !important;
    background: transparent !important;
    color: {theme["muted"]} !important;
    font-size: 13px !important;
    font-weight: 700 !important;
    padding: 4px 16px !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-top-title) div[data-testid="stPopover"] button * {{
    white-space: nowrap !important;
    line-height: 1 !important;
    margin-top: 0 !important;
    margin-bottom: 0 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-top-title) div[data-testid="stColumn"]:last-child div[data-testid="stPopover"] > div > button {{
    background: {theme["accent"]} !important;
    border-color: {theme["accent"]} !important;
    color: #fff !important;
    font-weight: 800 !important;
    min-width: 136px !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) {{
    height: var(--ez-topbar-h) !important;
    min-height: var(--ez-topbar-h) !important;
    max-height: var(--ez-topbar-h) !important;
    padding: 0 32px !important;
    gap: 18px !important;
    align-items: center !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
    overflow: visible !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div {{
    min-width: 0 !important;
    display: flex !important;
    align-items: center !important;
}}
div[data-testid="stLayoutWrapper"]:has(.ez-page-heading),
div[data-testid="stElementContainer"]:has(.ez-page-heading),
div[data-testid="stMarkdown"]:has(.ez-page-heading),
div[data-testid="stMarkdownContainer"]:has(.ez-page-heading) {{
    margin: 0 !important;
}}
.ez-page-heading {{
    display: flex !important;
    align-items: baseline !important;
    height: 27px !important;
    min-height: 27px !important;
    min-width: 0 !important;
    max-width: 100% !important;
    overflow: hidden !important;
    white-space: nowrap !important;
    text-overflow: ellipsis !important;
}}
.ez-page-heading .title {{
    margin: 0 !important;
    color: {theme["text"]} !important;
    font-size: 23px !important;
    line-height: 1 !important;
    font-weight: 850 !important;
    letter-spacing: 0 !important;
}}
.ez-page-heading .crumb {{
    color: {theme["muted"]} !important;
    font-size: 18px !important;
    font-weight: 650 !important;
    line-height: 1 !important;
    margin-left: 4px !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInput"],
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-baseweb="input"],
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInputRootElement"],
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-baseweb="base-input"],
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInput"] input,
div[data-testid="stColumn"]:has(.ez-filter-slot) div[data-testid="stPopover"] button,
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button {{
    height: 38px !important;
    min-height: 38px !important;
    max-height: 38px !important;
    box-sizing: border-box !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInput"] {{
    width: 100% !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-baseweb="input"] {{
    width: 100% !important;
    background: {theme["input"]} !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 8px !important;
    box-shadow: none !important;
    overflow: hidden !important;
    padding: 0 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-baseweb="base-input"] {{
    background: transparent !important;
    border: 0 !important;
    box-shadow: none !important;
    padding: 0 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-baseweb="input"]:focus-within {{
    border-color: {theme["accent"]} !important;
    box-shadow: 0 0 0 1px {theme["accent"]} inset !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInput"] input {{
    border: 0 !important;
    border-radius: 8px !important;
    background: transparent !important;
    color: {theme["text"]} !important;
    font-size: 15px !important;
    line-height: 1 !important;
    padding: 8px 14px !important;
    text-align: left !important;
    box-shadow: none !important;
    outline: none !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInput"] input::placeholder {{
    color: {theme["faint"]} !important;
    opacity: 1 !important;
}}
div[data-testid="stColumn"]:has(.ez-filter-slot) div[data-testid="stPopover"],
div[data-testid="stColumn"]:has(.ez-filter-slot) div[data-testid="stPopover"] > div,
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"],
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div {{
    width: 100% !important;
    max-width: none !important;
}}
div[data-testid="stColumn"]:has(.ez-filter-slot) div[data-testid="stPopover"] button,
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button {{
    width: 100% !important;
    min-width: auto !important;
    max-width: none !important;
    border-radius: 8px !important;
    font-size: 15px !important;
    line-height: 1 !important;
    padding: 8px 16px !important;
    box-shadow: none !important;
    font-weight: 760 !important;
    white-space: nowrap !important;
}}
div[data-testid="stColumn"]:has(.ez-filter-slot) div[data-testid="stPopover"] button {{
    background: {theme["input"]} !important;
    border: 1px solid {theme["border"]} !important;
    color: {theme["muted"]} !important;
}}
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button {{
    background: {theme["accent"]} !important;
    border: 1px solid {theme["accent"]} !important;
    color: #FFFFFF !important;
    font-weight: 850 !important;
}}
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button *,
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button p,
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button svg {{
    color: #FFFFFF !important;
    stroke: #FFFFFF !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInput"] {{
    max-width: 360px !important;
}}
div[data-testid="stColumn"]:has(.ez-filter-slot) div[data-testid="stPopover"],
div[data-testid="stColumn"]:has(.ez-filter-slot) div[data-testid="stPopover"] > div,
div[data-testid="stColumn"]:has(.ez-filter-slot) div[data-testid="stPopover"] button {{
    width: 104px !important;
    max-width: 104px !important;
}}
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"],
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div,
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] button {{
    width: 148px !important;
    max-width: 148px !important;
}}
.ez-filter-slot,
.ez-upload-slot {{
    display: none !important;
}}
div[data-testid="stElementContainer"]:has(.ez-filter-slot),
div[data-testid="stElementContainer"]:has(.ez-upload-slot) {{
    display: none !important;
    height: 0 !important;
    min-height: 0 !important;
    margin: 0 !important;
    padding: 0 !important;
    overflow: hidden !important;
}}
.ap-status-strip {{
    display: grid !important;
    grid-template-columns: repeat(5, minmax(0, 1fr)) !important;
    height: var(--ez-status-h) !important;
    min-height: var(--ez-status-h) !important;
    border: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
    border-radius: 0 !important;
    background: {theme["bg"]} !important;
    overflow: hidden !important;
    margin: 0 !important;
}}
.ap-status {{
    min-height: var(--ez-status-h) !important;
    display: flex !important;
    align-items: center !important;
    justify-content: flex-start !important;
    gap: 10px !important;
    padding: 0 28px !important;
    border-right: 1px solid {theme["border"]} !important;
    color: {theme["muted"]} !important;
    font-size: 14px !important;
    font-weight: 700 !important;
}}
.ap-status:last-child {{
    border-right: 0 !important;
}}
.ap-status strong {{
    font-size: 24px !important;
    font-weight: 800 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title) {{
    display: flex !important;
    flex-wrap: nowrap !important;
    align-items: stretch !important;
    gap: 0 !important;
    width: 100% !important;
    min-height: calc(100vh - var(--ez-head-h)) !important;
    background: {theme["bg"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="column"]:first-child {{
    flex: 0 0 var(--ez-queue-w) !important;
    width: var(--ez-queue-w) !important;
    min-width: var(--ez-queue-w) !important;
    max-width: var(--ez-queue-w) !important;
    border-right: 1px solid {theme["border"]} !important;
    overflow-y: auto !important;
    overflow-x: hidden !important;
    background: {theme["bg"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="column"]:nth-child(2) {{
    flex: 1 1 calc(100% - var(--ez-queue-w)) !important;
    width: auto !important;
    min-width: 0 !important;
    max-width: none !important;
    overflow-y: auto !important;
    overflow-x: hidden !important;
    background: {theme["bg"]} !important;
}}
.pane-title {{
    height: 48px !important;
    padding: 0 28px !important;
    border: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
}}
.invoice-list-row {{
    min-height: 92px !important;
    border-left: 0 !important;
    border-right: 0 !important;
    padding: 18px 28px !important;
}}
.empty-pane {{
    border-left: 0 !important;
    border-right: 0 !important;
    border-top: 0 !important;
}}
.detail-head {{
    display: flex !important;
    align-items: flex-start !important;
    justify-content: space-between !important;
    gap: 24px !important;
    border: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
    border-radius: 0 !important;
    padding: 28px 36px 22px !important;
}}
.detail-head > div:first-child {{
    min-width: 0 !important;
    flex: 1 1 auto !important;
}}
.detail-head .detail-actions {{
    flex: 0 0 auto !important;
    padding-top: 2px !important;
}}
.detail-grid {{
    padding: 22px 36px 24px !important;
    gap: 10px !important;
    border-bottom: 0 !important;
}}
.detail-grid div {{
    min-height: 72px !important;
    padding: 11px 12px !important;
}}
/* Column-based detail header (title left, badges right) — no overlap */
div[data-testid="stHorizontalBlock"]:has(.detail-title-block) {{
    padding: 28px 36px 18px !important;
    gap: 18px !important;
    align-items: flex-start !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.detail-title-block) > div {{
    display: flex !important;
    align-items: flex-start !important;
}}
.detail-title-block {{
    min-width: 0 !important;
}}
.detail-title-block .eyebrow {{
    display: block !important;
    color: {theme["faint"]} !important;
    font-size: 11px !important;
    font-weight: 800 !important;
    letter-spacing: .12em !important;
    text-transform: uppercase !important;
}}
.detail-title-block h2 {{
    margin: 8px 0 8px !important;
    color: {theme["text"]} !important;
    font-size: clamp(24px, 2vw, 32px) !important;
    line-height: 1.05 !important;
    font-weight: 850 !important;
    overflow-wrap: anywhere !important;
}}
.detail-title-block p {{
    margin: 0 !important;
    color: {theme["muted"]} !important;
    font-size: 15px !important;
    line-height: 1.3 !important;
}}
.detail-badges {{
    display: flex !important;
    flex-direction: row !important;
    flex-wrap: wrap !important;
    align-items: center !important;
    justify-content: flex-end !important;
    gap: 8px !important;
    width: 100% !important;
    padding-top: 4px !important;
}}
.empty-pane.large {{
    border: 0 !important;
    border-radius: 0 !important;
    min-height: calc(100vh - var(--ez-head-h)) !important;
}}
div[data-testid="stPopoverBody"] {{
    background: {theme["surface"]} !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 10px !important;
    box-shadow: 0 4px 16px rgba(0,0,0,.10) !important;
    padding: 18px !important;
    color: {theme["text"]} !important;
}}
div[data-testid="stPopoverBody"] section[data-testid="stFileUploaderDropzone"],
div[data-testid="stPopoverBody"] div[data-testid="stFileUploaderDropzone"] {{
    background: {theme["surface2"]} !important;
    border: 1px solid {theme["accent"]} !important;
    border-radius: 10px !important;
    box-shadow: none !important;
    padding: 20px !important;
}}
div[data-testid="stPopoverBody"] section[data-testid="stFileUploaderDropzone"] > div,
div[data-testid="stPopoverBody"] div[data-testid="stFileUploaderDropzone"] > div,
div[data-testid="stPopoverBody"] [data-testid="stFileUploaderDropzoneInstructions"],
div[data-testid="stPopoverBody"] [data-testid="stFileUploaderDropzoneInstructions"] * {{
    background: transparent !important;
    color: {theme["muted"]} !important;
}}
div[data-testid="stPopoverBody"] section[data-testid="stFileUploaderDropzone"] button,
div[data-testid="stPopoverBody"] div[data-testid="stFileUploaderDropzone"] button {{
    background: {theme["accent"]} !important;
    border: 1px solid {theme["accent"]} !important;
    color: #FFFFFF !important;
    border-radius: 8px !important;
    font-weight: 800 !important;
}}
div[data-testid="stPopoverBody"] section[data-testid="stFileUploaderDropzone"] button *,
div[data-testid="stPopoverBody"] div[data-testid="stFileUploaderDropzone"] button * {{
    color: #FFFFFF !important;
    stroke: #FFFFFF !important;
}}
div[class*="st-key-batch_exports"] {{
    padding: 14px 32px 16px !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
}}
div[class*="st-key-batch_exports"] > div[data-testid="stVerticalBlock"] {{
    gap: 0 !important;
}}
div[class*="st-key-batch_exports"] div[data-testid="stHorizontalBlock"] {{
    align-items: center !important;
    gap: 14px !important;
}}
div[class*="st-key-batch_exports"] div[data-testid="stPopover"],
div[class*="st-key-batch_exports"] div[data-testid="stPopover"] > div {{
    width: 180px !important;
    max-width: 180px !important;
}}
div[class*="st-key-batch_exports"] div[data-testid="stPopover"] > div > button {{
    width: 180px !important;
    height: 40px !important;
    min-height: 40px !important;
    border-radius: 8px !important;
    border: 1px solid {theme["accent"]} !important;
    background: {theme["accent"]} !important;
    color: #FFFFFF !important;
    font-size: 14px !important;
    font-weight: 800 !important;
    line-height: 1 !important;
    box-shadow: none !important;
}}
div[class*="st-key-batch_exports"] div[data-testid="stPopover"] > div > button:hover {{
    background: {theme["accent2"]} !important;
    border-color: {theme["accent2"]} !important;
}}
div[class*="st-key-batch_exports"] div[data-testid="stPopover"] > div > button *,
div[class*="st-key-batch_exports"] div[data-testid="stPopover"] > div > button svg {{
    color: #FFFFFF !important;
    stroke: #FFFFFF !important;
}}
.ez-export-ribbon {{
    min-height: 40px !important;
    display: flex !important;
    align-items: center !important;
    gap: 18px !important;
    padding: 0 18px !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 8px !important;
    background: {theme["surface"]} !important;
    color: {theme["muted"]} !important;
    overflow: hidden !important;
}}
.ez-export-ribbon strong {{
    flex: 0 0 auto !important;
    color: {theme["text"]} !important;
    font-size: 14px !important;
    font-weight: 850 !important;
}}
.ez-export-ribbon span {{
    min-width: 0 !important;
    color: {theme["muted"]} !important;
    font-size: 13px !important;
    font-weight: 650 !important;
    line-height: 1 !important;
    white-space: nowrap !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
}}
.ez-export-menu {{
    min-width: min(540px, calc(100vw - 96px)) !important;
    display: flex !important;
    flex-direction: column !important;
    gap: 0 !important;
    color: {theme["text"]} !important;
    margin-bottom: 16px !important;
    padding-bottom: 6px !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
.ez-export-menu-title {{
    color: {theme["text"]} !important;
    font-size: 17px !important;
    font-weight: 800 !important;
    line-height: 1.2 !important;
    margin-bottom: 4px !important;
}}
.ez-export-menu-sub {{
    color: {theme["muted"]} !important;
    font-size: 13px !important;
    line-height: 1.4 !important;
    margin-bottom: 10px !important;
}}
.ez-export-menu-row {{
    min-height: 38px !important;
    display: flex !important;
    align-items: center !important;
    justify-content: space-between !important;
    gap: 16px !important;
    padding: 9px 0 !important;
    border-top: 1px solid {theme["border"]} !important;
}}
.ez-export-menu-row span {{
    color: {theme["muted"]} !important;
    font-size: 11px !important;
    font-weight: 700 !important;
    letter-spacing: .08em !important;
    text-transform: uppercase !important;
}}
.ez-export-menu-row strong {{
    min-width: 0 !important;
    color: {theme["text"]} !important;
    font-size: 14px !important;
    font-weight: 700 !important;
    text-align: right !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) {{
    width: min(540px, calc(100vw - 48px)) !important;
    max-width: calc(100vw - 48px) !important;
    background: {theme["surface"]} !important;
    padding: 22px !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stHorizontalBlock"] {{
    gap: 10px !important;
    margin-bottom: 10px !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stDownloadButton"] button {{
    width: 100% !important;
    min-height: 44px !important;
    height: 44px !important;
    border-radius: 9px !important;
    background: {theme["input"]} !important;
    color: {theme["text"]} !important;
    border: 1px solid {theme["border"]} !important;
    font-weight: 700 !important;
    font-size: 14px !important;
    padding: 4px 8px !important;
    box-shadow: none !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stDownloadButton"] button:hover {{
    border-color: {theme["accent"]} !important;
    color: {theme["accent"]} !important;
    background: {theme["active_row"]} !important;
}}
/* Filter dropdown internals */
div[data-testid="stPopoverBody"]:has(.ez-filter-menu-title) {{
    width: 280px !important;
    background: {theme["surface"]} !important;
    padding: 18px !important;
}}
.ez-filter-menu-title {{
    color: {theme["faint"]} !important;
    font-size: 11px !important;
    font-weight: 800 !important;
    letter-spacing: .1em !important;
    text-transform: uppercase !important;
    margin-bottom: 10px !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-menu-title) div[data-testid="stCheckbox"] {{
    margin-bottom: 4px !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-menu-title) div[data-testid="stCheckbox"] label {{
    gap: 9px !important;
    align-items: center !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-menu-title) div[data-testid="stCheckbox"] label p {{
    color: {theme["text"]} !important;
    font-size: 14px !important;
    font-weight: 500 !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-menu-title) div[data-testid="stRadio"] {{
    margin-top: 14px !important;
    padding-top: 14px !important;
    border-top: 1px solid {theme["border"]} !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-menu-title) div[data-testid="stRadio"] > label,
div[data-testid="stPopoverBody"]:has(.ez-filter-menu-title) div[data-testid="stRadio"] > div > label:first-child {{
    color: {theme["faint"]} !important;
    font-size: 11px !important;
    font-weight: 800 !important;
    letter-spacing: .1em !important;
    text-transform: uppercase !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-menu-title) div[data-testid="stRadio"] label p {{
    color: {theme["text"]} !important;
    font-size: 14px !important;
    font-weight: 500 !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-menu-title) div[data-testid="stRadio"] div[role="radiogroup"] {{
    gap: 2px !important;
    margin-top: 6px !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-menu-title) div[data-testid="stRadio"] div[role="radiogroup"] label {{
    padding: 5px 0 !important;
}}
div[class*="st-key-invoice_row_"] {{
    position: relative !important;
    min-height: 92px !important;
    margin: 0 !important;
    padding: 0 !important;
}}
div[class*="st-key-invoice_row_"] > div[data-testid="stVerticalBlock"] {{
    gap: 0 !important;
}}
div[class*="st-key-invoice_row_"] div[data-testid="stElementContainer"]:has(div[data-testid="stButton"]),
div[class*="st-key-invoice_row_"] div[data-testid="stLayoutWrapper"]:has(div[data-testid="stButton"]) {{
    position: absolute !important;
    inset: 0 !important;
    z-index: 20 !important;
    width: 100% !important;
    height: 100% !important;
    min-height: 92px !important;
    margin: 0 !important;
    padding: 0 !important;
    overflow: hidden !important;
}}
div[class*="st-key-invoice_row_"] div[data-testid="stButton"] {{
    position: absolute !important;
    inset: 0 !important;
    z-index: 10 !important;
    width: 100% !important;
    height: 100% !important;
}}
div[class*="st-key-invoice_row_"] div[data-testid="stButton"] > button {{
    width: 100% !important;
    height: 100% !important;
    min-height: 92px !important;
    margin: 0 !important;
    border: 0 !important;
    border-radius: 0 !important;
    background: transparent !important;
    color: transparent !important;
    opacity: 0 !important;
    box-shadow: none !important;
    cursor: pointer !important;
}}
div[class*="st-key-invoice_row_"] div[data-testid="stButton"] > button *,
div[class*="st-key-invoice_row_"] div[data-testid="stButton"] > button p {{
    color: transparent !important;
    opacity: 0 !important;
    pointer-events: none !important;
}}
.invoice-list-row {{
    cursor: pointer !important;
    transition: background .14s ease, box-shadow .14s ease !important;
}}
div[class*="st-key-invoice_row_"]:hover .invoice-list-row {{
    background: {theme["surface"]} !important;
}}
.invoice-list-row.active {{
    box-shadow: inset 4px 0 0 {theme["accent"]} !important;
}}
.detail-head,
.detail-grid,
div[data-testid="stHorizontalBlock"]:has(.detail-head) {{
    max-width: 100% !important;
}}
.report-kpi-grid {{
    display: grid !important;
    grid-template-columns: repeat(5, minmax(132px, 1fr)) !important;
    gap: 10px !important;
    padding: 0 32px 18px !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
.report-kpi-grid.analytics {{
    padding-top: 2px !important;
}}
.report-kpi {{
    min-height: 96px !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 9px !important;
    background: {theme["surface"]} !important;
    padding: 14px 16px !important;
    overflow: hidden !important;
}}
.report-kpi span {{
    display: block !important;
    color: {theme["faint"]} !important;
    font-size: 11px !important;
    font-weight: 850 !important;
    text-transform: uppercase !important;
    letter-spacing: .08em !important;
}}
.report-kpi strong {{
    display: block !important;
    margin-top: 9px !important;
    color: {theme["text"]} !important;
    font-size: clamp(19px, 1.6vw, 26px) !important;
    line-height: 1.05 !important;
    font-weight: 850 !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
}}
.report-kpi em {{
    display: block !important;
    margin-top: 8px !important;
    color: {theme["muted"]} !important;
    font-size: 12px !important;
    font-style: normal !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
}}
.report-kpi.danger strong {{
    color: {theme["danger"]} !important;
}}
.report-section-title,
.report-chart-title {{
    height: 42px !important;
    display: flex !important;
    align-items: end !important;
    padding: 0 32px 9px !important;
    color: {theme["text"]} !important;
    font-size: 18px !important;
    font-weight: 850 !important;
}}
.report-chart-title {{
    padding: 18px 0 8px !important;
    height: auto !important;
    font-size: 15px !important;
}}
.empty-report {{
    min-height: calc(100vh - var(--ez-topbar-h)) !important;
    display: grid !important;
    place-content: center !important;
    text-align: center !important;
    padding: 32px !important;
    background: {theme["bg"]} !important;
    color: {theme["muted"]} !important;
}}
.empty-report strong {{
    display: block !important;
    color: {theme["text"]} !important;
    font-size: 22px !important;
    margin-bottom: 8px !important;
}}
.empty-report span {{
    display: block !important;
    color: {theme["muted"]} !important;
    font-size: 14px !important;
}}
div[data-testid="stHorizontalBlock"]:has(.report-chart-title) {{
    padding: 0 32px 10px !important;
    gap: 18px !important;
}}
.analytics-page,
.history-workspace {{
    width: 100% !important;
    box-sizing: border-box !important;
    padding: 34px 48px 86px !important;
    background: {theme["bg"]} !important;
    color: {theme["text"]} !important;
}}
.analytics-kpis,
.history-kpis {{
    display: grid !important;
    grid-template-columns: repeat(4, minmax(0, 1fr)) !important;
    gap: 22px !important;
    margin-bottom: 26px !important;
}}
.analytics-grid,
.history-grid {{
    display: grid !important;
    grid-template-columns: minmax(0, 1.15fr) minmax(0, .85fr) !important;
    gap: 24px !important;
    margin-bottom: 26px !important;
}}
.analytics-grid.three {{
    grid-template-columns: repeat(3, minmax(0, 1fr)) !important;
}}
.analytics-card,
.history-panel,
.history-table-wrap {{
    min-width: 0 !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 12px !important;
    background: {theme["surface2"]} !important;
    box-shadow: none !important;
    text-align: left !important;
}}
.analytics-card,
.history-panel {{
    padding: 24px 26px !important;
}}
.analytics-kpis .analytics-card,
.history-kpis .analytics-card {{
    min-height: 132px !important;
    padding: 24px 26px !important;
}}
.analytics-card-tall,
.history-panel {{
    min-height: 246px !important;
}}
.analytics-card-small {{
    min-height: 188px !important;
}}
.analytics-card h3,
.history-panel h3 {{
    width: 100% !important;
    display: flex !important;
    align-items: center !important;
    justify-content: space-between !important;
    gap: 12px !important;
    text-align: left !important;
}}
.analytics-card h3,
.history-panel h3,
.history-table-pro th {{
    margin: 0 0 18px !important;
    color: {theme["muted"]} !important;
    font-size: 12px !important;
    line-height: 1.1 !important;
    font-weight: 850 !important;
    letter-spacing: .12em !important;
    text-transform: uppercase !important;
}}
.analytics-card h3 span {{
    float: none !important;
    margin-left: auto !important;
    color: {theme["muted"]} !important;
    letter-spacing: 0 !important;
    text-transform: none !important;
    font-size: 12px !important;
    font-weight: 750 !important;
}}
.metric-big {{
    margin: 14px 0 8px !important;
    color: {theme["text"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, SFMono-Regular, Menlo, monospace !important;
    font-size: clamp(24px, 2.4vw, 32px) !important;
    line-height: 1 !important;
    font-weight: 800 !important;
    letter-spacing: 0 !important;
    white-space: nowrap !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
}}
.metric-big.green {{
    color: {theme["export"]} !important;
}}
.metric-big.amber {{
    color: {theme["warn"]} !important;
}}
.metric-big.dark {{
    color: {theme["text"]} !important;
}}
.metric-note,
.activity-meta {{
    color: {theme["muted"]} !important;
    font-size: 13px !important;
    line-height: 1.35 !important;
    font-weight: 550 !important;
}}
.activity-title {{
    color: {theme["text"]} !important;
    font-size: 16px !important;
    line-height: 1.25 !important;
    font-weight: 850 !important;
}}
.row-line {{
    display: grid !important;
    grid-template-columns: minmax(0, 1fr) minmax(120px, 40%) auto !important;
    gap: 14px !important;
    align-items: center !important;
    min-height: 38px !important;
    color: {theme["muted"]} !important;
    font-size: 14px !important;
    font-weight: 650 !important;
}}
.row-line > div:first-child,
.mini-list-row > div:first-child,
.legend-row > span:nth-child(2) {{
    min-width: 0 !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
}}
.bar-track {{
    height: 9px !important;
    background: {theme["panel2"]} !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 999px !important;
    overflow: hidden !important;
}}
.bar-fill {{
    height: 100% !important;
    background: {theme["export"]} !important;
    border-radius: 999px !important;
}}
.donut-wrap {{
    display: grid !important;
    grid-template-columns: 168px minmax(0, 1fr) !important;
    gap: 24px !important;
    align-items: center !important;
}}
.donut {{
    width: 150px !important;
    height: 150px !important;
    border-radius: 50% !important;
    position: relative !important;
    display: grid !important;
    place-items: center !important;
}}
.donut::after {{
    content: "" !important;
    position: absolute !important;
    inset: 39px !important;
    border-radius: 50% !important;
    background: {theme["surface2"]} !important;
}}
.donut-center {{
    position: relative !important;
    z-index: 1 !important;
    color: {theme["text"]} !important;
    text-align: center !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 18px !important;
    line-height: 1.05 !important;
    font-weight: 800 !important;
}}
.donut-center span {{
    display: block !important;
    color: {theme["muted"]} !important;
    font-family: 'Inter', sans-serif !important;
    font-size: 11px !important;
    font-weight: 800 !important;
}}
.legend-row,
.mini-list-row {{
    display: grid !important;
    grid-template-columns: auto minmax(0, 1fr) auto !important;
    gap: 10px !important;
    align-items: center !important;
    min-height: 36px !important;
    color: {theme["muted"]} !important;
    font-size: 14px !important;
    font-weight: 650 !important;
}}
.mini-list-row {{
    grid-template-columns: minmax(0, 1fr) auto !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
.mini-list-row:last-child {{
    border-bottom: 0 !important;
}}
.legend-dot {{
    width: 10px !important;
    height: 10px !important;
    border-radius: 999px !important;
}}
.money {{
    color: {theme["text"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-weight: 800 !important;
    letter-spacing: 0 !important;
    white-space: nowrap !important;
}}
.currency-pill {{
    display: inline-grid !important;
    place-items: center !important;
    min-width: 44px !important;
    height: 24px !important;
    border-radius: 999px !important;
    background: {theme["chip_validated_bg"]} !important;
    color: {theme["export"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 12px !important;
    font-weight: 850 !important;
    margin-right: 10px !important;
}}
.currency-count {{
    color: {theme["muted"]} !important;
}}
.history-stage {{
    display: grid !important;
    grid-template-columns: 34px minmax(0, 1fr) auto !important;
    gap: 16px !important;
    align-items: center !important;
    min-height: 58px !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
.history-stage:last-child {{
    border-bottom: 0 !important;
}}
.history-stage .dot {{
    width: 26px !important;
    height: 26px !important;
    display: grid !important;
    place-items: center !important;
    border-radius: 999px !important;
    background: {theme["chip_validated_bg"]} !important;
    color: {theme["export"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 12px !important;
    font-weight: 850 !important;
}}
.history-stage strong,
.activity-row strong {{
    color: {theme["text"]} !important;
    font-weight: 850 !important;
}}
.history-stage span:not(.dot) {{
    color: {theme["muted"]} !important;
    font-size: 13px !important;
    line-height: 1.3 !important;
}}
.activity-row {{
    display: grid !important;
    grid-template-columns: minmax(0, 1fr) auto !important;
    gap: 16px !important;
    align-items: center !important;
    min-height: 58px !important;
    padding: 8px 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
.activity-row:last-child {{
    border-bottom: 0 !important;
}}
.history-table-wrap {{
    overflow-x: auto !important;
}}
.history-table-pro {{
    width: 100% !important;
    border-collapse: collapse !important;
    min-width: 850px !important;
}}
.history-table-pro th,
.history-table-pro td {{
    padding: 14px 18px !important;
    border-bottom: 1px solid {theme["border"]} !important;
    text-align: left !important;
}}
.history-table-pro th {{
    margin: 0 !important;
}}
.history-table-pro td {{
    color: {theme["muted"]} !important;
    font-size: 14px !important;
    font-weight: 600 !important;
}}
.history-table-pro tr:last-child td {{
    border-bottom: 0 !important;
}}
.history-table-pro .primary {{
    color: {theme["text"]} !important;
    font-weight: 850 !important;
}}
.history-table-pro .mono {{
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
}}
.table-chip {{
    display: inline-flex !important;
    align-items: center !important;
    min-height: 24px !important;
    padding: 0 10px !important;
    border-radius: 999px !important;
    background: {theme["chip_validated_bg"]} !important;
    color: {theme["good"]} !important;
    font-size: 12px !important;
    font-weight: 850 !important;
}}
.table-chip.warn {{
    background: {theme["chip_pending_bg"]} !important;
    color: {theme["warn"]} !important;
}}
.table-chip.danger {{
    background: {theme["chip_exception_bg"]} !important;
    color: {theme["danger"]} !important;
}}
.history-foot {{
    margin-top: 10px !important;
    color: {theme["muted"]} !important;
    font-size: 12px !important;
    font-weight: 650 !important;
}}
@media (max-width: 980px) {{
    .block-container {{
        padding: 0 !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.ez-top-title) {{
        height: auto !important;
        max-height: none !important;
        min-height: var(--ez-topbar-h) !important;
        padding: 12px 16px !important;
        flex-wrap: wrap !important;
        gap: 8px !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.ez-top-title) > div[data-testid="column"],
    div[data-testid="stHorizontalBlock"]:has(.ez-top-title) > div[data-testid="stColumn"] {{
        width: auto !important;
        min-width: 0 !important;
        max-width: none !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.ez-top-title) > div[data-testid="column"]:nth-child(1),
    div[data-testid="stHorizontalBlock"]:has(.ez-top-title) > div[data-testid="stColumn"]:nth-child(1) {{
        flex: 0 0 136px !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.ez-top-title) > div[data-testid="column"]:nth-child(2),
    div[data-testid="stHorizontalBlock"]:has(.ez-top-title) > div[data-testid="stColumn"]:nth-child(2) {{
        flex: 1 1 190px !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.ez-top-title) > div[data-testid="column"]:nth-child(3),
    div[data-testid="stHorizontalBlock"]:has(.ez-top-title) > div[data-testid="stColumn"]:nth-child(3) {{
        flex: 0 0 86px !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.ez-top-title) > div[data-testid="column"]:nth-child(4),
    div[data-testid="stHorizontalBlock"]:has(.ez-top-title) > div[data-testid="stColumn"]:nth-child(4) {{
        flex: 0 0 136px !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) {{
        height: auto !important;
        max-height: none !important;
        min-height: var(--ez-topbar-h) !important;
        padding: 12px 16px !important;
        flex-wrap: wrap !important;
        gap: 10px !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div[data-testid="column"],
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div[data-testid="stColumn"] {{
        width: auto !important;
        min-width: 0 !important;
        max-width: none !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div[data-testid="column"]:nth-child(1),
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div[data-testid="stColumn"]:nth-child(1) {{
        flex: 0 0 100% !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div[data-testid="column"]:nth-child(2),
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div[data-testid="stColumn"]:nth-child(2) {{
        flex: 1 1 220px !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div[data-testid="column"]:nth-child(3),
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div[data-testid="stColumn"]:nth-child(3) {{
        flex: 0 0 118px !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div[data-testid="column"]:nth-child(4),
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div[data-testid="stColumn"]:nth-child(4) {{
        flex: 0 0 170px !important;
    }}
    .ap-status-strip {{
        grid-template-columns: 1fr !important;
        height: auto !important;
    }}
    .ap-status {{
        justify-content: flex-start !important;
        border-right: 0 !important;
        border-bottom: 1px solid {theme["border"]} !important;
    }}
    div[class*="st-key-batch_exports"] {{
        padding: 12px 16px 24px !important;
    }}
    div[class*="st-key-batch_exports"] div[data-testid="stHorizontalBlock"] {{
        display: block !important;
    }}
    div[class*="st-key-batch_exports"] div[data-testid="column"] {{
        width: 100% !important;
        min-width: 0 !important;
        margin-bottom: 10px !important;
    }}
    .ez-export-ribbon {{
        align-items: flex-start !important;
        flex-direction: column !important;
        justify-content: center !important;
        gap: 4px !important;
        min-height: 54px !important;
        padding: 8px 14px !important;
    }}
    .ez-export-ribbon span {{
        white-space: normal !important;
        line-height: 1.25 !important;
    }}
    div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) {{
        width: calc(100vw - 32px) !important;
        max-width: calc(100vw - 32px) !important;
    }}
    .ez-export-menu {{
        min-width: 0 !important;
        width: 100% !important;
    }}
    .report-kpi-grid {{
        grid-template-columns: 1fr !important;
        padding: 0 16px 16px !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.report-chart-title) {{
        padding: 0 16px 10px !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.pane-title) {{
        display: block !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="column"]:first-child,
    div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="column"]:nth-child(2) {{
        width: 100% !important;
        min-width: 0 !important;
        max-width: none !important;
        border-right: 0 !important;
    }}
    .analytics-page,
    .history-workspace {{
        padding: 24px 16px 72px !important;
    }}
    .analytics-kpis,
    .history-kpis,
    .analytics-grid,
    .analytics-grid.three,
    .history-grid {{
        grid-template-columns: 1fr !important;
        gap: 14px !important;
    }}
    .analytics-card,
    .history-panel {{
        padding: 18px 18px !important;
    }}
    .donut-wrap {{
        grid-template-columns: 1fr !important;
        justify-items: start !important;
    }}
    .row-line {{
        grid-template-columns: 1fr !important;
        gap: 7px !important;
    }}
}}

/* Canonical EZ console stylesheet. This is the final source of truth for
   the demo workspace layout after the older Streamlit rules above. */
:root {{
    --ez-sidebar-w: 260px;
    --ez-topbar-h: 84px;
    --ez-status-h: 74px;
    --ez-head-h: calc(var(--ez-topbar-h) + var(--ez-status-h));
    --ez-queue-w: clamp(340px, 31vw, 420px);
}}
html, body, .stApp {{
    background: {theme["bg"]} !important;
    color: {theme["text"]} !important;
    overflow-x: hidden !important;
    -webkit-font-smoothing: antialiased !important;
    text-rendering: optimizeLegibility !important;
}}
main, .block-container, section[data-testid="stMain"], section[data-testid="stMain"] > div {{
    overflow-x: hidden !important;
}}
button, input, textarea {{
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', system-ui, sans-serif !important;
}}
div[data-testid="stDownloadButton"] > button,
div[data-testid="stButton"] > button {{
    min-height: 38px !important;
    border-radius: 8px !important;
    border: 1px solid {theme["border"]} !important;
    background: {theme["input"]} !important;
    color: {theme["text"]} !important;
    box-shadow: none !important;
    font-size: 13px !important;
    line-height: 1 !important;
    font-weight: 740 !important;
}}
div[data-testid="stDownloadButton"] > button *,
div[data-testid="stButton"] > button * {{
    color: inherit !important;
    stroke: currentColor !important;
}}
div[data-testid="stButton"] > button[kind="primary"],
div[data-testid="stDownloadButton"] > button[kind="primary"] {{
    background: {theme["accent"]} !important;
    border-color: {theme["accent"]} !important;
    color: #FFFFFF !important;
}}
div[data-testid="stButton"] > button:disabled,
div[data-testid="stDownloadButton"] > button:disabled {{
    background: {theme["surface"]} !important;
    color: {theme["faint"]} !important;
    border-color: {theme["border"]} !important;
    opacity: .78 !important;
}}
section[data-testid="stSidebar"] {{
    width: var(--ez-sidebar-w) !important;
    min-width: var(--ez-sidebar-w) !important;
    max-width: var(--ez-sidebar-w) !important;
    background: {theme["surface"]} !important;
    border-right: 1px solid {theme["border"]} !important;
    overflow: hidden !important;
}}
.ez-brand {{
    height: var(--ez-topbar-h) !important;
    min-height: var(--ez-topbar-h) !important;
    padding: 0 24px !important;
    gap: 14px !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
.ez-brand-mark {{
    width: 36px !important;
    height: 36px !important;
    border-radius: 10px !important;
    background: {theme["accent"]} !important;
    color: #FFFFFF !important;
    font-size: 17px !important;
    line-height: 1 !important;
    font-weight: 850 !important;
}}
.ez-brand-name {{
    color: {theme["text"]} !important;
    font-size: 15px !important;
    line-height: 1 !important;
    font-weight: 850 !important;
}}
.ez-brand-ver {{
    color: {theme["faint"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 10px !important;
    line-height: 1 !important;
    margin-top: 7px !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_nav_shell"] {{
    height: calc(100vh - var(--ez-topbar-h) - var(--ez-sidebar-footer-h)) !important;
    max-height: calc(100vh - var(--ez-topbar-h) - var(--ez-sidebar-footer-h)) !important;
    padding: 16px 12px 20px !important;
    overflow-y: auto !important;
    overflow-x: hidden !important;
    scrollbar-gutter: stable !important;
}}
.nav-section {{
    height: 18px !important;
    margin: 18px 14px 10px !important;
    color: {theme["faint"]} !important;
    font-size: 10px !important;
    line-height: 18px !important;
    letter-spacing: .12em !important;
    font-weight: 850 !important;
    text-transform: uppercase !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-nav_item_"] {{
    height: 44px !important;
    min-height: 44px !important;
    margin: 0 0 4px !important;
    position: relative !important;
}}
.side-nav-row {{
    height: 40px !important;
    min-height: 40px !important;
    display: grid !important;
    grid-template-columns: 24px minmax(0, 1fr) auto !important;
    align-items: center !important;
    column-gap: 12px !important;
    padding: 0 14px !important;
    border-left: 3px solid transparent !important;
    border-radius: 8px !important;
    color: {theme["muted"]} !important;
    font-size: 14px !important;
    line-height: 1 !important;
    font-weight: 650 !important;
    pointer-events: none !important;
}}
.side-nav-row.active {{
    background: {theme["input"] if is_light_mode else theme["active_row"]} !important;
    color: {theme["text"]} !important;
    border-left-color: {theme["accent"]} !important;
}}
.side-icon {{
    width: 24px !important;
    height: 24px !important;
    display: grid !important;
    place-items: center !important;
    color: {theme["accent"] if is_light_mode else theme["accent"]} !important;
    font-size: 15px !important;
    font-weight: 800 !important;
}}
.side-nav-row > span:nth-child(2) {{
    min-width: 0 !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
}}
.side-nav-count {{
    min-width: 24px !important;
    height: 24px !important;
    display: inline-grid !important;
    place-items: center !important;
    padding: 0 7px !important;
    border-radius: 999px !important;
    background: {theme["surface2"]} !important;
    color: {theme["faint"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 12px !important;
    line-height: 1 !important;
    font-weight: 800 !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-nav_item_"] div[data-testid="stButton"],
section[data-testid="stSidebar"] div[class*="st-key-nav_btn_"] {{
    position: absolute !important;
    inset: 0 !important;
    width: 100% !important;
    height: 44px !important;
    margin: 0 !important;
    z-index: 10 !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-nav_item_"] div[data-testid="stButton"] > button {{
    width: 100% !important;
    height: 44px !important;
    min-height: 44px !important;
    padding: 0 !important;
    border: 0 !important;
    background: transparent !important;
    color: transparent !important;
    opacity: 0 !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_sidebar_footer"] {{
    border-top: 1px solid {theme["border"]} !important;
    background: {theme["surface"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) {{
    height: var(--ez-topbar-h) !important;
    min-height: var(--ez-topbar-h) !important;
    max-height: var(--ez-topbar-h) !important;
    padding: 0 32px !important;
    gap: 18px !important;
    align-items: center !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
    overflow: visible !important;
}}
.ez-page-heading {{
    display: flex !important;
    align-items: baseline !important;
    gap: 6px !important;
    min-width: 0 !important;
}}
.ez-page-heading .title {{
    color: {theme["text"]} !important;
    font-size: 23px !important;
    line-height: 1 !important;
    font-weight: 850 !important;
}}
.ez-page-heading .crumb {{
    color: {theme["muted"]} !important;
    font-size: 18px !important;
    line-height: 1 !important;
    font-weight: 650 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-baseweb="input"],
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInput"] input {{
    height: 38px !important;
    min-height: 38px !important;
    border-radius: 8px !important;
    background: {theme["input"]} !important;
    color: {theme["text"]} !important;
    border-color: {theme["border"]} !important;
    box-shadow: none !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInput"] input {{
    padding: 8px 14px !important;
    font-size: 15px !important;
    line-height: 1 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInput"] input::placeholder {{
    color: {theme["faint"]} !important;
    opacity: 1 !important;
}}
div[data-testid="stColumn"]:has(.ez-filter-slot) div[data-testid="stPopover"] button,
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button {{
    height: 38px !important;
    min-height: 38px !important;
    border-radius: 8px !important;
    box-shadow: none !important;
    font-size: 15px !important;
    line-height: 1 !important;
    font-weight: 750 !important;
}}
div[data-testid="stColumn"]:has(.ez-filter-slot) div[data-testid="stPopover"] button {{
    background: {theme["input"]} !important;
    border: 1px solid {theme["border"]} !important;
    color: {theme["muted"]} !important;
}}
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button {{
    background: {theme["accent"]} !important;
    border: 1px solid {theme["accent"]} !important;
    color: #FFFFFF !important;
}}
div[data-testid="stPopoverBody"] {{
    background: {theme["surface"]} !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 10px !important;
    box-shadow: 0 18px 42px rgba(0,0,0,.22) !important;
    color: {theme["text"]} !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) {{
    width: 330px !important;
    max-width: calc(100vw - 32px) !important;
    padding: 18px !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploader"] {{
    margin: 0 !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) section[data-testid="stFileUploaderDropzone"],
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploaderDropzone"] {{
    min-height: 132px !important;
    border: 1px dashed {theme["border2"]} !important;
    border-radius: 10px !important;
    background: {theme["surface2"]} !important;
    box-shadow: none !important;
    padding: 20px 16px !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderDropzoneInstructions"],
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderDropzoneInstructions"] * {{
    background: transparent !important;
    color: {theme["muted"]} !important;
    font-size: 14px !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploader"] button {{
    background: {theme["input"]} !important;
    border: 1px solid {theme["border2"]} !important;
    color: {theme["text"]} !important;
    border-radius: 8px !important;
    min-height: 38px !important;
    box-shadow: none !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stButton"] > button[kind="primary"] {{
    height: 42px !important;
    min-height: 42px !important;
    background: {theme["accent"]} !important;
    border-color: {theme["accent"]} !important;
    color: #FFFFFF !important;
    font-size: 15px !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stButton"] > button:not([kind="primary"]) {{
    height: 36px !important;
    min-height: 36px !important;
    margin-top: 8px !important;
    background: {theme["input"]} !important;
    color: {theme["text"]} !important;
}}
.ap-status-strip {{
    height: var(--ez-status-h) !important;
    min-height: var(--ez-status-h) !important;
    max-height: var(--ez-status-h) !important;
    display: grid !important;
    grid-template-columns: repeat(5, minmax(0, 1fr)) !important;
    background: {theme["bg"]} !important;
    border: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
    overflow: hidden !important;
}}
.ap-status {{
    min-height: var(--ez-status-h) !important;
    display: flex !important;
    align-items: center !important;
    justify-content: center !important;
    gap: 12px !important;
    padding: 0 18px !important;
    border-right: 1px solid {theme["border"]} !important;
    color: {theme["muted"]} !important;
    font-size: 15px !important;
    line-height: 1 !important;
    font-weight: 750 !important;
}}
.ap-status:last-child {{
    border-right: 0 !important;
}}
.ap-status strong {{
    color: {theme["accent"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 24px !important;
    line-height: 1 !important;
    font-weight: 850 !important;
}}
.ap-status:nth-child(4) strong {{
    color: {theme["warn"]} !important;
}}
.ap-status.danger strong {{
    color: {theme["danger"]} !important;
}}
.ap-status.active {{
    box-shadow: inset 0 -3px 0 {theme["accent"]} !important;
}}
div[class*="st-key-batch_exports"] {{
    padding: 20px 32px 18px !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
}}
div[class*="st-key-batch_exports"] > div[data-testid="stVerticalBlock"] {{
    gap: 0 !important;
}}
div[class*="st-key-batch_exports"] > div[data-testid="stVerticalBlock"] > div[data-testid="stHorizontalBlock"] {{
    gap: 16px !important;
    align-items: center !important;
}}
div[class*="st-key-batch_exports"] > div[data-testid="stVerticalBlock"] > div[data-testid="stHorizontalBlock"] > div[data-testid="column"]:first-child {{
    flex: 0 0 220px !important;
    width: 220px !important;
    min-width: 220px !important;
    max-width: 220px !important;
}}
div[class*="st-key-batch_exports"] div[data-testid="stPopover"] > div > button {{
    height: 40px !important;
    min-height: 40px !important;
    border-radius: 8px !important;
    background: {theme["accent"]} !important;
    border-color: {theme["accent"]} !important;
    color: #FFFFFF !important;
    font-size: 14px !important;
    font-weight: 800 !important;
}}
.ez-export-ribbon {{
    min-height: 40px !important;
    display: flex !important;
    align-items: center !important;
    gap: 14px !important;
    padding: 0 16px !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 8px !important;
    background: {theme["surface2"]} !important;
    overflow: hidden !important;
}}
.ez-export-ribbon strong {{
    color: {theme["text"]} !important;
    font-size: 14px !important;
    line-height: 1 !important;
    font-weight: 850 !important;
    white-space: nowrap !important;
}}
.ez-export-ribbon span {{
    min-width: 0 !important;
    color: {theme["muted"]} !important;
    font-size: 13px !important;
    line-height: 1 !important;
    font-weight: 650 !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) {{
    width: min(900px, calc(100vw - 64px)) !important;
    max-width: calc(100vw - 64px) !important;
    padding: 22px 24px !important;
    background: {theme["surface"]} !important;
}}
.ez-export-menu {{
    width: 100% !important;
    min-width: 0 !important;
    display: grid !important;
    gap: 18px !important;
    color: {theme["text"]} !important;
}}
.ez-export-menu-head {{
    padding-bottom: 16px !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
.ez-export-menu-title {{
    color: {theme["text"]} !important;
    font-size: 18px !important;
    line-height: 1.1 !important;
    font-weight: 850 !important;
}}
.ez-export-menu-sub {{
    margin-top: 5px !important;
    color: {theme["muted"]} !important;
    font-size: 14px !important;
    line-height: 1.35 !important;
    font-weight: 550 !important;
}}
.ez-export-facts {{
    display: grid !important;
    grid-template-columns: repeat(4, minmax(0, 1fr)) !important;
    gap: 0 !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 8px !important;
    overflow: hidden !important;
}}
.ez-export-facts div {{
    min-height: 58px !important;
    padding: 10px 14px !important;
    background: {theme["surface2"]} !important;
    border-right: 1px solid {theme["border"]} !important;
}}
.ez-export-facts div:last-child {{
    border-right: 0 !important;
}}
.ez-export-facts span {{
    display: block !important;
    color: {theme["faint"]} !important;
    font-size: 10px !important;
    line-height: 1 !important;
    letter-spacing: .12em !important;
    text-transform: uppercase !important;
    font-weight: 850 !important;
}}
.ez-export-facts strong {{
    display: block !important;
    margin-top: 8px !important;
    color: {theme["text"]} !important;
    font-size: 14px !important;
    line-height: 1.15 !important;
    font-weight: 850 !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stHorizontalBlock"] {{
    gap: 10px !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stDownloadButton"] > button {{
    height: 40px !important;
    min-height: 40px !important;
    border-radius: 8px !important;
    background: {theme["input"]} !important;
    border-color: {theme["border"]} !important;
    color: {theme["text"]} !important;
    font-size: 12px !important;
    font-weight: 750 !important;
    padding: 4px 8px !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stDownloadButton"] > button:hover {{
    border-color: {theme["accent"]} !important;
    color: {theme["accent"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title) {{
    display: grid !important;
    grid-template-columns: var(--ez-queue-w) minmax(0, 1fr) !important;
    flex-wrap: nowrap !important;
    gap: 0 !important;
    width: 100% !important;
    min-height: calc(100vh - var(--ez-head-h)) !important;
    background: {theme["bg"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="column"]:first-child,
div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="stColumn"]:first-child {{
    flex: unset !important;
    width: 100% !important;
    min-width: 0 !important;
    max-width: none !important;
    border-right: 1px solid {theme["border"]} !important;
    overflow-x: hidden !important;
    overflow-y: auto !important;
    background: {theme["bg"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="column"]:nth-child(2),
div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="stColumn"]:nth-child(2) {{
    flex: 1 1 auto !important;
    min-width: 0 !important;
    max-width: none !important;
    width: 100% !important;
    overflow-x: hidden !important;
    overflow-y: auto !important;
    background: {theme["bg"]} !important;
}}
.pane-title {{
    height: 52px !important;
    min-height: 52px !important;
    display: flex !important;
    align-items: center !important;
    justify-content: space-between !important;
    padding: 0 30px !important;
    border: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
    color: {theme["muted"]} !important;
    font-size: 14px !important;
    line-height: 1 !important;
    font-weight: 800 !important;
}}
.invoice-list-row {{
    min-height: 112px !important;
    display: grid !important;
    grid-template-columns: minmax(0, 1fr) auto !important;
    gap: 18px !important;
    align-items: start !important;
    padding: 20px 30px !important;
    border: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
}}
.invoice-list-row.active {{
    background: {theme["active_row"]} !important;
    border-left: 4px solid {theme["accent"]} !important;
    padding-left: 26px !important;
}}
.invoice-list-row strong {{
    color: {theme["text"]} !important;
    font-size: 16px !important;
    line-height: 1.15 !important;
    font-weight: 850 !important;
}}
.invoice-list-row span {{
    color: {theme["muted"]} !important;
    margin-top: 9px !important;
    font-size: 14px !important;
    line-height: 1.25 !important;
    font-weight: 560 !important;
}}
.invoice-list-row em {{
    color: {theme["faint"]} !important;
    margin-top: 12px !important;
    font-size: 12px !important;
    line-height: 1 !important;
    font-style: normal !important;
    font-weight: 650 !important;
}}
.invoice-row-meta {{
    align-items: flex-end !important;
    gap: 18px !important;
    white-space: nowrap !important;
}}
.invoice-row-meta b {{
    color: {theme["text"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 15px !important;
    line-height: 1.1 !important;
    font-weight: 800 !important;
}}
.detail-head {{
    display: flex !important;
    justify-content: space-between !important;
    gap: 24px !important;
    padding: 32px 40px 22px !important;
    border: 0 !important;
    background: {theme["bg"]} !important;
}}
.detail-head h2 {{
    margin: 12px 0 10px !important;
    color: {theme["text"]} !important;
    font-size: clamp(28px, 3vw, 38px) !important;
    line-height: 1 !important;
    font-weight: 850 !important;
    letter-spacing: 0 !important;
}}
.detail-head p {{
    margin: 0 !important;
    color: {theme["muted"]} !important;
    font-size: 15px !important;
    line-height: 1.35 !important;
    font-weight: 560 !important;
}}
.detail-actions {{
    flex: 0 0 auto !important;
    display: flex !important;
    flex-direction: column !important;
    align-items: flex-end !important;
    gap: 10px !important;
}}
.detail-grid {{
    display: grid !important;
    grid-template-columns: repeat(auto-fit, minmax(160px, 1fr)) !important;
    gap: 12px !important;
    padding: 0 40px !important;
    margin: 0 0 24px !important;
    border: 0 !important;
    background: {theme["bg"]} !important;
}}
.detail-grid div {{
    min-height: 82px !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 8px !important;
    padding: 14px 16px !important;
    background: {theme["surface2"]} !important;
}}
.detail-grid span {{
    color: {theme["faint"]} !important;
    font-size: 11px !important;
    line-height: 1 !important;
    letter-spacing: .12em !important;
    text-transform: uppercase !important;
    font-weight: 850 !important;
}}
.detail-grid strong {{
    margin-top: 12px !important;
    color: {theme["text"]} !important;
    font-size: 18px !important;
    line-height: 1.12 !important;
    font-weight: 800 !important;
    overflow-wrap: anywhere !important;
}}
div[class*="st-key-invoice_detail_actions"] {{
    clear: both !important;
    position: relative !important;
    z-index: 1 !important;
    margin: 0 40px 0 !important;
    padding: 0 0 24px !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
}}
div[class*="st-key-invoice_detail_actions"] > div[data-testid="stVerticalBlock"] {{
    gap: 0 !important;
}}
div[class*="st-key-invoice_detail_actions"] div[data-testid="stHorizontalBlock"] {{
    display: grid !important;
    grid-template-columns: repeat(3, minmax(0, 1fr)) !important;
    gap: 12px !important;
}}
div[class*="st-key-invoice_detail_actions"] div[data-testid="stHorizontalBlock"] > div[data-testid="column"],
div[class*="st-key-invoice_detail_actions"] div[data-testid="stHorizontalBlock"] > div[data-testid="stColumn"] {{
    width: 100% !important;
    min-width: 0 !important;
    max-width: none !important;
    flex: unset !important;
}}
div[class*="st-key-invoice_detail_actions"] div[data-testid="stDownloadButton"],
div[class*="st-key-invoice_detail_actions"] div[data-testid="stButton"] {{
    width: 100% !important;
}}
div[class*="st-key-invoice_detail_actions"] div[data-testid="stDownloadButton"] > button,
div[class*="st-key-invoice_detail_actions"] div[data-testid="stButton"] > button {{
    width: 100% !important;
    height: 40px !important;
    min-height: 40px !important;
    font-size: 13px !important;
    font-weight: 740 !important;
}}
div[data-testid="stElementContainer"]:has(.ez-line-section-title),
div[data-testid="stMarkdown"]:has(.ez-line-section-title) {{
    margin: 0 !important;
}}
.ez-line-section-title {{
    margin: 0 !important;
    padding: 24px 40px 0 !important;
    color: {theme["text"]} !important;
    font-size: 20px !important;
    line-height: 1.1 !important;
    font-weight: 850 !important;
    background: {theme["bg"]} !important;
}}
.ez-line-table-wrap {{
    max-width: calc(100% - 80px) !important;
    margin: 14px 40px 44px !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 9px !important;
    overflow-x: auto !important;
    overflow-y: hidden !important;
    background: {theme["surface"]} !important;
}}
.ez-line-table {{
    width: 100% !important;
    min-width: 660px !important;
    border-collapse: collapse !important;
    table-layout: auto !important;
}}
.ez-line-table th {{
    background: {theme["surface2"]} !important;
    color: {theme["muted"]} !important;
    border-right: 1px solid {theme["border"]} !important;
    border-bottom: 1px solid {theme["border"]} !important;
    padding: 12px 14px !important;
    text-align: left !important;
    font-size: 11px !important;
    line-height: 1 !important;
    letter-spacing: .06em !important;
    text-transform: uppercase !important;
    font-weight: 750 !important;
    white-space: nowrap !important;
}}
.ez-line-table td {{
    background: {theme["bg"]} !important;
    color: {theme["text"]} !important;
    border-right: 1px solid {theme["border"]} !important;
    border-bottom: 1px solid {theme["border"]} !important;
    padding: 13px 14px !important;
    font-size: 13px !important;
    line-height: 1.28 !important;
    font-weight: 590 !important;
    vertical-align: middle !important;
}}
.ez-line-table th:last-child,
.ez-line-table td:last-child {{
    border-right: 0 !important;
}}
.ez-line-table tr:last-child td {{
    border-bottom: 0 !important;
}}
.ez-line-table td.num {{
    text-align: right !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-weight: 750 !important;
    white-space: nowrap !important;
}}
.empty-pane,
.empty-pane.large {{
    border-color: {theme["border"]} !important;
    background: {theme["bg"]} !important;
    color: {theme["text"]} !important;
}}
.analytics-card,
.history-panel,
.history-table-wrap {{
    background: {theme["surface2"]} !important;
    border-color: {theme["border"]} !important;
}}
.metric-big.green,
.bar-fill,
.currency-pill,
.history-stage .dot {{
    color: {theme["accent"]} !important;
}}
.bar-fill {{
    background: {theme["accent"]} !important;
}}
.donut-center span,
.legend-row,
.mini-list-row,
.metric-note,
.activity-meta,
.history-stage span {{
    color: {theme["muted"]} !important;
}}
.history-stage {{
    display: grid !important;
    grid-template-columns: 34px minmax(0, 1fr) 42px !important;
    gap: 14px !important;
    align-items: center !important;
    min-height: 60px !important;
    padding: 10px 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
.history-stage .dot {{
    width: 26px !important;
    height: 26px !important;
    border-radius: 999px !important;
    display: grid !important;
    place-items: center !important;
    background: {theme["chip_validated_bg"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 12px !important;
    line-height: 1 !important;
    font-weight: 850 !important;
}}
.history-stage > strong:last-child {{
    justify-self: end !important;
    text-align: right !important;
    color: {theme["text"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 14px !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stRadio"] {{
    width: 280px !important;
    max-width: 100% !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stRadio"] div[role="radiogroup"] {{
    width: 280px !important;
    max-width: 100% !important;
    height: 38px !important;
    display: grid !important;
    grid-template-columns: repeat(4, minmax(0, 1fr)) !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 8px !important;
    overflow: hidden !important;
    background: {theme["surface"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stRadio"] div[role="radiogroup"] label {{
    width: 100% !important;
    min-width: 0 !important;
    height: 36px !important;
    min-height: 36px !important;
    margin: 0 !important;
    padding: 0 !important;
    display: flex !important;
    justify-content: center !important;
    align-items: center !important;
    border-right: 1px solid {theme["border"]} !important;
    border-radius: 0 !important;
    background: transparent !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stRadio"] div[role="radiogroup"] label:last-child {{
    border-right: 0 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stRadio"] div[role="radiogroup"] label > div:first-child {{
    display: none !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stRadio"] div[role="radiogroup"] label p {{
    color: {theme["muted"]} !important;
    font-size: 13px !important;
    line-height: 1 !important;
    font-weight: 800 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stRadio"] div[role="radiogroup"] label:has(input:checked) {{
    background: {theme["accent"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stRadio"] div[role="radiogroup"] label:has(input:checked) p {{
    color: #FFFFFF !important;
}}
@media (max-width: 980px) {{
    div[data-testid="stHorizontalBlock"]:has(.pane-title) {{
        display: block !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="column"]:first-child,
    div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="stColumn"]:first-child,
    div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="column"]:nth-child(2),
    div[data-testid="stHorizontalBlock"]:has(.pane-title) > div[data-testid="stColumn"]:nth-child(2) {{
        width: 100% !important;
        min-width: 0 !important;
        max-width: none !important;
        border-right: 0 !important;
    }}
    .ap-status-strip,
    .ez-export-facts {{
        grid-template-columns: 1fr !important;
        height: auto !important;
        max-height: none !important;
    }}
    .ap-status,
    .ez-export-facts div {{
        border-right: 0 !important;
        border-bottom: 1px solid {theme["border"]} !important;
    }}
    div[class*="st-key-batch_exports"] > div[data-testid="stVerticalBlock"] > div[data-testid="stHorizontalBlock"] {{
        display: block !important;
    }}
    div[class*="st-key-batch_exports"] > div[data-testid="stVerticalBlock"] > div[data-testid="stHorizontalBlock"] > div[data-testid="column"]:first-child {{
        width: 100% !important;
        min-width: 0 !important;
        max-width: none !important;
        margin-bottom: 10px !important;
    }}
    .detail-head,
    .detail-grid,
    div[class*="st-key-invoice_detail_actions"],
    .ez-line-section-title {{
        margin-left: 0 !important;
        margin-right: 0 !important;
        padding-left: 18px !important;
        padding-right: 18px !important;
    }}
    div[class*="st-key-invoice_detail_actions"] div[data-testid="stHorizontalBlock"] {{
        grid-template-columns: 1fr !important;
    }}
    .ez-line-table-wrap {{
        max-width: calc(100% - 36px) !important;
        margin-left: 18px !important;
        margin-right: 18px !important;
    }}
}}

/* Final demo QA layer: processed invoice state, popovers, and dark-mode polish. */
div[data-testid="stElementContainer"]:has(.ap-status-strip),
div[data-testid="stMarkdown"]:has(.ap-status-strip),
div[data-testid="stMarkdownContainer"]:has(.ap-status-strip) {{
    height: var(--ez-status-h) !important;
    min-height: var(--ez-status-h) !important;
    max-height: var(--ez-status-h) !important;
    margin: 0 !important;
    padding: 0 !important;
    overflow: hidden !important;
}}
.ap-status-strip {{
    height: var(--ez-status-h) !important;
    min-height: var(--ez-status-h) !important;
    max-height: var(--ez-status-h) !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title) {{
    min-height: calc(100vh - var(--ez-head-h)) !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title) > div:has(.empty-pane.large) {{
    border-left: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
.empty-pane:not(.large) {{
    height: calc(100vh - var(--ez-head-h) - 52px) !important;
    min-height: 420px !important;
    display: grid !important;
    place-content: center !important;
    gap: 14px !important;
    padding: 0 42px !important;
    border: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
    text-align: center !important;
}}
.empty-pane.large {{
    height: calc(100vh - var(--ez-head-h)) !important;
    min-height: 480px !important;
    display: grid !important;
    place-content: center !important;
    gap: 14px !important;
    padding: 0 48px !important;
    border: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
    text-align: center !important;
}}
.empty-pane strong {{
    color: {theme["text"]} !important;
    font-size: 18px !important;
    line-height: 1.2 !important;
    font-weight: 820 !important;
    letter-spacing: 0 !important;
}}
.empty-pane span {{
    max-width: 430px !important;
    color: {theme["muted"]} !important;
    font-size: 15px !important;
    line-height: 1.45 !important;
    font-weight: 520 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title):has(.detail-title-block) {{
    display: grid !important;
    grid-template-columns: var(--ez-queue-w) minmax(0, 1fr) !important;
    gap: 0 !important;
    align-items: stretch !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title):has(.detail-title-block) > div:has(.pane-title) {{
    grid-column: 1 / 2 !important;
    width: var(--ez-queue-w) !important;
    min-width: var(--ez-queue-w) !important;
    max-width: var(--ez-queue-w) !important;
    overflow: hidden auto !important;
    border-right: 1px solid {theme["border"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title):has(.detail-title-block) > div:has(.detail-title-block) {{
    grid-column: 2 / 3 !important;
    width: 100% !important;
    min-width: 0 !important;
    max-width: none !important;
    overflow: hidden auto !important;
}}
div[class*="st-key-invoice_row_"] {{
    position: relative !important;
    width: 100% !important;
    height: 112px !important;
    min-height: 112px !important;
    max-height: 112px !important;
    overflow: hidden !important;
}}
div[class*="st-key-invoice_row_"] .invoice-list-row {{
    position: relative !important;
    z-index: 1 !important;
    width: 100% !important;
    height: 112px !important;
    min-height: 112px !important;
    max-height: 112px !important;
    box-sizing: border-box !important;
    pointer-events: none !important;
    overflow: hidden !important;
}}
div[class*="st-key-invoice_row_"] .invoice-list-row > div:first-child {{
    min-width: 0 !important;
    overflow: hidden !important;
}}
div[class*="st-key-invoice_row_"] .invoice-list-row strong,
div[class*="st-key-invoice_row_"] .invoice-list-row span,
div[class*="st-key-invoice_row_"] .invoice-list-row em {{
    max-width: 100% !important;
    display: block !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
}}
div[class*="st-key-invoice_row_"] .invoice-list-row strong,
div[class*="st-key-invoice_row_"] .invoice-list-row em {{
    white-space: nowrap !important;
}}
div[class*="st-key-invoice_row_"] .invoice-list-row span {{
    display: -webkit-box !important;
    -webkit-line-clamp: 2 !important;
    -webkit-box-orient: vertical !important;
}}
div[class*="st-key-invoice_row_"] .invoice-row-meta {{
    min-width: 120px !important;
    max-width: 142px !important;
}}
div[class*="st-key-invoice_row_"] .invoice-row-meta b {{
    max-width: 142px !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
}}
div[class*="st-key-select_invoice_row_"] {{
    position: absolute !important;
    inset: 0 !important;
    z-index: 8 !important;
    width: 100% !important;
    height: 112px !important;
    min-height: 112px !important;
    max-height: 112px !important;
    margin: 0 !important;
    padding: 0 !important;
}}
div[class*="st-key-select_invoice_row_"] div[data-testid="stButton"],
div[class*="st-key-select_invoice_row_"] div[data-testid="stButton"] > button {{
    position: absolute !important;
    inset: 0 !important;
    width: 100% !important;
    height: 112px !important;
    min-height: 112px !important;
    max-height: 112px !important;
    margin: 0 !important;
    padding: 0 !important;
    border: 0 !important;
    border-radius: 0 !important;
    background: transparent !important;
    color: transparent !important;
    opacity: 0 !important;
    box-shadow: none !important;
    cursor: pointer !important;
}}
div[class*="st-key-select_invoice_row_"] div[data-testid="stButton"] > button *,
div[class*="st-key-select_invoice_row_"] div[data-testid="stButton"] > button p {{
    color: transparent !important;
    opacity: 0 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.detail-title-block) {{
    display: grid !important;
    grid-template-columns: minmax(0, 1fr) auto !important;
    gap: 24px !important;
    align-items: start !important;
    padding: 34px 40px 20px !important;
    border: 0 !important;
    background: {theme["bg"]} !important;
    overflow: visible !important;
}}
div[data-testid="stHorizontalBlock"]:has(.detail-title-block) > div[data-testid="column"],
div[data-testid="stHorizontalBlock"]:has(.detail-title-block) > div[data-testid="stColumn"] {{
    width: auto !important;
    min-width: 0 !important;
    max-width: none !important;
    flex: unset !important;
}}
.detail-title-block {{
    min-width: 0 !important;
    max-width: 100% !important;
}}
.detail-title-block .eyebrow {{
    display: block !important;
    color: {theme["faint"]} !important;
    font-size: 10px !important;
    line-height: 1 !important;
    letter-spacing: .14em !important;
    text-transform: uppercase !important;
    font-weight: 850 !important;
}}
.detail-title-block h2 {{
    max-width: 100% !important;
    margin: 16px 0 12px !important;
    color: {theme["text"]} !important;
    font-size: clamp(26px, 2.3vw, 34px) !important;
    line-height: 1.02 !important;
    font-weight: 850 !important;
    letter-spacing: 0 !important;
    overflow-wrap: anywhere !important;
}}
.detail-title-block p {{
    max-width: 54ch !important;
    margin: 0 !important;
    color: {theme["muted"]} !important;
    font-size: 15px !important;
    line-height: 1.32 !important;
    font-weight: 560 !important;
}}
.detail-badges {{
    min-width: 136px !important;
    display: flex !important;
    flex-direction: column !important;
    align-items: flex-end !important;
    justify-content: flex-start !important;
    gap: 10px !important;
    padding-top: 0 !important;
}}
.status-pill {{
    min-height: 28px !important;
    display: inline-flex !important;
    align-items: center !important;
    justify-content: center !important;
    padding: 0 14px !important;
    border-radius: 999px !important;
    border: 1px solid transparent !important;
    font-size: 12px !important;
    line-height: 1 !important;
    font-weight: 800 !important;
    white-space: nowrap !important;
}}
.status-pill.neutral {{
    color: {theme["muted"]} !important;
    background: {theme["surface"]} !important;
    border-color: {theme["border"]} !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]),
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) {{
    background: {theme["surface"]} !important;
    border-color: {theme["border"]} !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) > div,
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stVerticalBlock"],
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) > div,
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stVerticalBlock"] {{
    background: transparent !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) section[data-testid="stFileUploaderDropzone"],
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploaderDropzone"] {{
    background: {theme["bg"]} !important;
    border: 1px dashed {theme["border2"]} !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploader"] button {{
    background: {theme["input"]} !important;
    color: {theme["text"]} !important;
    border-color: {theme["border2"]} !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploader"] button *,
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploader"] button p,
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploader"] button span {{
    color: {theme["text"]} !important;
    fill: currentColor !important;
    stroke: currentColor !important;
    opacity: 1 !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) section[data-testid="stFileUploaderDropzone"],
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploaderDropzone"] {{
    min-height: 124px !important;
    padding: 18px 16px !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderDropzoneInstructions"],
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderDropzoneInstructions"] * {{
    font-size: 13px !important;
    line-height: 1.35 !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) section[data-testid="stFileUploaderDropzone"]:has([data-testid="stFileUploaderFile"]),
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploaderDropzone"]:has([data-testid="stFileUploaderFile"]) {{
    min-height: 112px !important;
    padding: 12px !important;
    border-style: solid !important;
    background: {theme["surface2"]} !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) section[data-testid="stFileUploaderDropzone"]:has([data-testid="stFileUploaderFile"]) button,
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploaderDropzone"]:has([data-testid="stFileUploaderFile"]) button {{
    width: 34px !important;
    min-width: 34px !important;
    max-width: 34px !important;
    height: 34px !important;
    min-height: 34px !important;
    max-height: 34px !important;
    margin-top: 10px !important;
    padding: 0 !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 8px !important;
    background: {theme["input"]} !important;
    color: {theme["text"]} !important;
    box-shadow: none !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] {{
    width: 100% !important;
    min-height: 58px !important;
    display: grid !important;
    grid-template-columns: 38px minmax(0, 1fr) 30px !important;
    align-items: center !important;
    column-gap: 12px !important;
    padding: 8px 10px !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 9px !important;
    background: {theme["input"]} !important;
    box-sizing: border-box !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] * {{
    min-width: 0 !important;
    max-width: 100% !important;
    color: {theme["text"]} !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] [data-testid="stFileUploaderFileName"],
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] span,
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] p {{
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
    font-size: 13px !important;
    line-height: 1.2 !important;
    font-weight: 720 !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] small,
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] [data-testid="stFileUploaderFileSize"] {{
    color: {theme["muted"]} !important;
    font-size: 11px !important;
    line-height: 1.2 !important;
    font-weight: 560 !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] svg {{
    color: {theme["muted"]} !important;
    fill: none !important;
    stroke: currentColor !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] button,
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] button:hover {{
    width: 28px !important;
    min-width: 28px !important;
    max-width: 28px !important;
    height: 28px !important;
    min-height: 28px !important;
    max-height: 28px !important;
    padding: 0 !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 7px !important;
    background: transparent !important;
    color: {theme["muted"]} !important;
    box-shadow: none !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) section[data-testid="stFileUploaderDropzone"]:has([data-testid="stFileUploaderFile"]) > button,
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploaderDropzone"]:has([data-testid="stFileUploaderFile"]) > button {{
    width: 34px !important;
    min-width: 34px !important;
    max-width: 34px !important;
    height: 34px !important;
    min-height: 34px !important;
    max-height: 34px !important;
    margin-top: 10px !important;
    padding: 0 !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 8px !important;
    background: {theme["input"]} !important;
    color: {theme["text"]} !important;
    box-shadow: none !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stButton"] > button[kind="primary"],
div[class*="st-key-batch_exports"] div[data-testid="stPopover"] > div > button,
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button {{
    background: {theme["accent"]} !important;
    border-color: {theme["accent"]} !important;
    color: #FFFFFF !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) {{
    width: min(980px, calc(100vw - 96px)) !important;
    padding: 22px !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) .ez-export-menu {{
    background: transparent !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stHorizontalBlock"] {{
    display: grid !important;
    grid-template-columns: repeat(6, minmax(112px, 1fr)) !important;
    gap: 10px !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stHorizontalBlock"] > div[data-testid="column"],
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stHorizontalBlock"] > div[data-testid="stColumn"],
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stDownloadButton"] {{
    width: 100% !important;
    min-width: 0 !important;
    max-width: none !important;
    flex: unset !important;
}}
.ez-export-facts {{
    grid-template-columns: repeat(4, minmax(0, 1fr)) !important;
}}
.ez-export-facts div,
.detail-grid div,
.analytics-card,
.history-panel,
.history-table-wrap {{
    background: {theme["surface2"]} !important;
    border-color: {theme["border"]} !important;
}}
.ez-line-table-wrap {{
    max-width: calc(100% - 80px) !important;
    margin-top: 24px !important;
}}
.ez-line-table,
.ez-line-table thead,
.ez-line-table tbody,
.ez-line-table tr,
.ez-line-table th,
.ez-line-table td {{
    background: {theme["bg"]} !important;
}}
.ez-line-table th {{
    background: {theme["surface2"]} !important;
    color: {theme["muted"]} !important;
}}
.ez-line-table td {{
    color: {theme["text"]} !important;
}}
.bar-fill,
.currency-pill {{
    background-color: {theme["accent"]} !important;
    color: #FFFFFF !important;
}}
.metric-big.green {{
    display: block !important;
    width: auto !important;
    max-width: 100% !important;
    padding: 0 !important;
    border: 0 !important;
    border-radius: 0 !important;
    background: transparent !important;
    color: {theme["accent"]} !important;
    box-shadow: none !important;
}}
.donut {{
    filter: saturate(.92) !important;
}}
.history-stage .dot {{
    color: {theme["accent"]} !important;
    background: {theme["chip_validated_bg"]} !important;
}}
@media (max-width: 1220px) {{
    div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stHorizontalBlock"] {{
        grid-template-columns: repeat(3, minmax(120px, 1fr)) !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.detail-title-block) {{
        padding-left: 28px !important;
        padding-right: 28px !important;
    }}
    .detail-grid,
    div[class*="st-key-invoice_detail_actions"],
    .ez-line-section-title {{
        padding-left: 28px !important;
        padding-right: 28px !important;
    }}
    .ez-line-table-wrap {{
        max-width: calc(100% - 56px) !important;
        margin-left: 28px !important;
        margin-right: 28px !important;
    }}
}}
@media (max-width: 760px) {{
    div[data-testid="stHorizontalBlock"]:has(.detail-title-block) {{
        display: block !important;
    }}
    .detail-badges {{
        align-items: flex-start !important;
        margin-top: 16px !important;
    }}
    div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stHorizontalBlock"] {{
        grid-template-columns: 1fr !important;
    }}
}}

/* EZ-Invoice demo polish system: one final visual pass for shell, rhythm,
   controls, panels, and dark-mode parity. */
:root {{
    --ez-sidebar-w: 260px;
    --ez-topbar-h: 84px;
    --ez-status-h: 74px;
    --ez-head-h: calc(var(--ez-topbar-h) + var(--ez-status-h));
    --ez-queue-w: clamp(352px, 30vw, 424px);
    --ez-radius: 10px;
    --ez-card-radius: 12px;
    --ez-pop-shadow: 0 20px 54px rgba(0,0,0,{".16" if is_light_mode else ".46"});
}}
html, body, .stApp {{
    background: {theme["bg"]} !important;
    color: {theme["text"]} !important;
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', system-ui, sans-serif !important;
    -webkit-font-smoothing: antialiased !important;
    text-rendering: geometricPrecision !important;
}}
.block-container,
div[data-testid="stMainBlockContainer"],
section[data-testid="stMain"] > div {{
    padding: 0 !important;
    margin: 0 !important;
    max-width: none !important;
}}
section[data-testid="stSidebar"] {{
    width: var(--ez-sidebar-w) !important;
    min-width: var(--ez-sidebar-w) !important;
    max-width: var(--ez-sidebar-w) !important;
    background: {theme["surface"]} !important;
    border-right: 1px solid {theme["border"]} !important;
}}
.ez-brand {{
    height: var(--ez-topbar-h) !important;
    min-height: var(--ez-topbar-h) !important;
    padding: 0 24px !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
.ez-brand-mark {{
    width: 38px !important;
    height: 38px !important;
    border-radius: 10px !important;
    background: {theme["accent"]} !important;
    color: #fff !important;
    font-size: 17px !important;
    font-weight: 850 !important;
}}
.ez-brand-name {{
    color: {theme["text"]} !important;
    font-size: 15px !important;
    font-weight: 850 !important;
}}
.ez-brand-ver {{
    color: {theme["faint"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 10px !important;
    margin-top: 6px !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_nav_shell"] {{
    padding: 18px 12px 22px !important;
    scrollbar-gutter: stable both-edges !important;
}}
.nav-section {{
    margin: 18px 14px 10px !important;
    color: {theme["faint"]} !important;
    font-size: 10px !important;
    letter-spacing: .14em !important;
    font-weight: 850 !important;
}}
.side-nav-row {{
    height: 40px !important;
    min-height: 40px !important;
    grid-template-columns: 24px minmax(0, 1fr) auto !important;
    column-gap: 12px !important;
    border-radius: 8px !important;
    color: {theme["muted"]} !important;
    font-size: 14px !important;
    line-height: 1 !important;
    font-weight: 680 !important;
}}
.side-nav-row.active {{
    background: {theme["input"] if is_light_mode else theme["active_row"]} !important;
    color: {theme["text"]} !important;
    box-shadow: inset 3px 0 0 {theme["accent"]} !important;
}}
.side-nav-row.active .side-icon,
.side-nav-row.active > span:nth-child(2) {{
    color: {theme["accent"]} !important;
}}
.side-nav-count {{
    min-width: 25px !important;
    height: 25px !important;
    display: inline-grid !important;
    place-items: center !important;
    padding: 0 7px !important;
    background: {theme["surface2"]} !important;
    color: {theme["faint"]} !important;
    border-radius: 999px !important;
    font-size: 12px !important;
}}
section[data-testid="stSidebar"] div[class*="st-key-ez_sidebar_footer"] {{
    padding: 12px 24px 16px !important;
    border-top: 1px solid {theme["border"]} !important;
    background: {theme["surface"]} !important;
}}

div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) {{
    height: var(--ez-topbar-h) !important;
    min-height: var(--ez-topbar-h) !important;
    max-height: var(--ez-topbar-h) !important;
    display: grid !important;
    grid-template-columns: minmax(260px, 1fr) minmax(320px, 520px) 124px 176px !important;
    gap: 18px !important;
    align-items: center !important;
    padding: 0 34px !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
    overflow: visible !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div[data-testid="column"],
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) > div[data-testid="stColumn"] {{
    width: 100% !important;
    min-width: 0 !important;
    max-width: none !important;
    flex: unset !important;
}}
.ez-page-heading {{
    display: flex !important;
    align-items: baseline !important;
    gap: 0 !important;
    min-width: 0 !important;
    max-width: 100% !important;
    white-space: nowrap !important;
    overflow: hidden !important;
}}
.ez-page-heading .title {{
    color: {theme["text"]} !important;
    font-size: 24px !important;
    line-height: 1 !important;
    font-weight: 850 !important;
    letter-spacing: 0 !important;
}}
.ez-page-heading .slash {{
    margin-left: 0 !important;
    color: {theme["muted"]} !important;
    font-size: 22px !important;
    line-height: 1 !important;
    font-weight: 550 !important;
}}
.ez-page-heading .crumb {{
    margin-left: 6px !important;
    color: {theme["muted"]} !important;
    font-size: 18px !important;
    line-height: 1 !important;
    font-weight: 650 !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-baseweb="input"],
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInput"] input,
div[data-testid="stColumn"]:has(.ez-filter-slot) div[data-testid="stPopover"] button,
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button {{
    height: 38px !important;
    min-height: 38px !important;
    max-height: 38px !important;
    border-radius: 8px !important;
    box-shadow: none !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-baseweb="input"] {{
    background: {theme["input"]} !important;
    border: 1px solid {theme["border"]} !important;
    overflow: hidden !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-baseweb="input"]:focus-within {{
    border-color: {theme["accent"]} !important;
    box-shadow: 0 0 0 2px {"rgba(74,103,65,.12)" if is_light_mode else "rgba(125,164,110,.20)"} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInput"] input {{
    padding: 0 14px !important;
    border: 0 !important;
    background: transparent !important;
    color: {theme["text"]} !important;
    font-size: 15px !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stTextInput"] input::placeholder {{
    color: {theme["faint"]} !important;
    opacity: 1 !important;
}}
div[data-testid="stColumn"]:has(.ez-filter-slot) div[data-testid="stPopover"] button {{
    background: {theme["input"]} !important;
    border: 1px solid {theme["border"]} !important;
    color: {theme["muted"]} !important;
    font-size: 15px !important;
    font-weight: 720 !important;
}}
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button {{
    background: {theme["accent"]} !important;
    border: 1px solid {theme["accent"]} !important;
    color: #fff !important;
    font-size: 15px !important;
    font-weight: 830 !important;
}}
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button *,
div[data-testid="stColumn"]:has(.ez-upload-slot) div[data-testid="stPopover"] > div > button svg {{
    color: #fff !important;
    stroke: #fff !important;
}}

div[data-testid="stElementContainer"]:has(.ap-status-strip),
div[data-testid="stMarkdown"]:has(.ap-status-strip),
div[data-testid="stMarkdownContainer"]:has(.ap-status-strip) {{
    height: var(--ez-status-h) !important;
    min-height: var(--ez-status-h) !important;
    margin: 0 !important;
    overflow: hidden !important;
}}
.ap-status-strip {{
    height: var(--ez-status-h) !important;
    min-height: var(--ez-status-h) !important;
    display: grid !important;
    grid-template-columns: repeat(5, minmax(0, 1fr)) !important;
    background: {theme["bg"]} !important;
    border: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
    overflow: hidden !important;
}}
.ap-status {{
    min-height: var(--ez-status-h) !important;
    display: flex !important;
    align-items: center !important;
    justify-content: flex-start !important;
    gap: 12px !important;
    padding: 0 30px !important;
    border-right: 1px solid {theme["border"]} !important;
    color: {theme["muted"]} !important;
    font-size: 15px !important;
    line-height: 1 !important;
    font-weight: 760 !important;
}}
.ap-status:last-child {{
    border-right: 0 !important;
}}
.ap-status strong {{
    color: {theme["accent"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 25px !important;
    line-height: 1 !important;
    font-weight: 850 !important;
}}
.ap-status:nth-child(4) strong {{
    color: {theme["warn"]} !important;
}}
.ap-status.danger strong {{
    color: {theme["danger"]} !important;
}}
.ap-status.active {{
    background: {"#F4F7F1" if is_light_mode else "#151B13"} !important;
    box-shadow: inset 0 -3px 0 {theme["accent"]} !important;
    color: {theme["text"]} !important;
}}

div[class*="st-key-batch_exports"] {{
    padding: 18px 34px !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
}}
div[class*="st-key-batch_exports"] > div[data-testid="stVerticalBlock"] > div[data-testid="stHorizontalBlock"] {{
    display: grid !important;
    grid-template-columns: 220px minmax(0, 1fr) !important;
    gap: 18px !important;
    align-items: center !important;
}}
div[class*="st-key-batch_exports"] > div[data-testid="stVerticalBlock"] > div[data-testid="stHorizontalBlock"] > div[data-testid="column"],
div[class*="st-key-batch_exports"] > div[data-testid="stVerticalBlock"] > div[data-testid="stHorizontalBlock"] > div[data-testid="stColumn"] {{
    width: 100% !important;
    min-width: 0 !important;
    max-width: none !important;
    flex: unset !important;
}}
div[class*="st-key-batch_exports"] div[data-testid="stPopover"] > div > button {{
    width: 220px !important;
    height: 42px !important;
    min-height: 42px !important;
    border-radius: 9px !important;
    background: {theme["accent"]} !important;
    border-color: {theme["accent"]} !important;
    color: #FFFFFF !important;
    font-size: 15px !important;
    font-weight: 820 !important;
}}
.ez-export-ribbon {{
    min-height: 42px !important;
    display: flex !important;
    align-items: center !important;
    gap: 14px !important;
    padding: 0 18px !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 9px !important;
    background: {theme["surface2"]} !important;
}}
.ez-export-ribbon strong {{
    color: {theme["text"]} !important;
    font-size: 14px !important;
    font-weight: 850 !important;
    white-space: nowrap !important;
}}
.ez-export-ribbon span {{
    color: {theme["muted"]} !important;
    font-size: 13px !important;
    font-weight: 650 !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
}}
div[data-testid="stPopoverBody"] {{
    background: {theme["surface"]} !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 12px !important;
    box-shadow: var(--ez-pop-shadow) !important;
    color: {theme["text"]} !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) {{
    width: min(920px, calc(100vw - 72px)) !important;
    max-width: calc(100vw - 72px) !important;
    padding: 24px !important;
}}
.ez-export-menu {{
    display: grid !important;
    gap: 18px !important;
    min-width: 0 !important;
    width: 100% !important;
}}
.ez-export-menu-head {{
    padding-bottom: 16px !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
.ez-export-menu-title {{
    color: {theme["text"]} !important;
    font-size: 18px !important;
    line-height: 1.1 !important;
    font-weight: 850 !important;
}}
.ez-export-menu-sub {{
    margin-top: 6px !important;
    color: {theme["muted"]} !important;
    font-size: 14px !important;
    line-height: 1.35 !important;
}}
.ez-export-facts {{
    display: grid !important;
    grid-template-columns: repeat(4, minmax(0, 1fr)) !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 9px !important;
    overflow: hidden !important;
}}
.ez-export-facts div {{
    min-height: 62px !important;
    padding: 12px 15px !important;
    border-right: 1px solid {theme["border"]} !important;
    background: {theme["surface2"]} !important;
}}
.ez-export-facts div:last-child {{
    border-right: 0 !important;
}}
.ez-export-facts span {{
    display: block !important;
    color: {theme["faint"]} !important;
    font-size: 10px !important;
    line-height: 1 !important;
    letter-spacing: .12em !important;
    text-transform: uppercase !important;
    font-weight: 850 !important;
}}
.ez-export-facts strong {{
    display: block !important;
    margin-top: 9px !important;
    color: {theme["text"]} !important;
    font-size: 14px !important;
    line-height: 1.15 !important;
    font-weight: 850 !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stHorizontalBlock"] {{
    display: grid !important;
    grid-template-columns: repeat(6, minmax(108px, 1fr)) !important;
    gap: 10px !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-export-package-menu) div[data-testid="stDownloadButton"] > button {{
    height: 42px !important;
    min-height: 42px !important;
    border-radius: 9px !important;
    background: {theme["input"]} !important;
    border-color: {theme["border"]} !important;
    color: {theme["text"]} !important;
    font-size: 13px !important;
    font-weight: 760 !important;
}}

div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) {{
    width: 340px !important;
    max-width: calc(100vw - 32px) !important;
    padding: 18px !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) section[data-testid="stFileUploaderDropzone"],
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploaderDropzone"] {{
    min-height: 126px !important;
    padding: 16px !important;
    border: 1px dashed {theme["border2"]} !important;
    border-radius: 10px !important;
    background: {theme["surface2"]} !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploader"] button {{
    min-height: 36px !important;
    border-radius: 8px !important;
    background: {theme["input"]} !important;
    border: 1px solid {theme["border"]} !important;
    color: {theme["text"]} !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploader"] button *,
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) div[data-testid="stFileUploader"] button p {{
    color: {theme["text"]} !important;
    stroke: currentColor !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] {{
    display: grid !important;
    grid-template-columns: 34px minmax(0, 1fr) 28px !important;
    align-items: center !important;
    gap: 10px !important;
    min-height: 56px !important;
    padding: 8px 10px !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 9px !important;
    background: {theme["input"]} !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] * {{
    min-width: 0 !important;
    max-width: 100% !important;
    color: {theme["text"]} !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] span,
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] p {{
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
    font-size: 13px !important;
    line-height: 1.2 !important;
    font-weight: 720 !important;
}}
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFileSize"],
div[data-testid="stPopoverBody"]:has(div[data-testid="stFileUploader"]) [data-testid="stFileUploaderFile"] small {{
    color: {theme["muted"]} !important;
    font-size: 11px !important;
}}

div[data-testid="stHorizontalBlock"]:has(.pane-title) {{
    display: grid !important;
    grid-template-columns: var(--ez-queue-w) minmax(0, 1fr) !important;
    gap: 0 !important;
    align-items: stretch !important;
    min-height: calc(100vh - var(--ez-head-h)) !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title) > div {{
    min-width: 0 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.pane-title) > div:has(.pane-title) {{
    border-right: 1px solid {theme["border"]} !important;
    overflow: hidden auto !important;
}}
.pane-title {{
    height: 52px !important;
    min-height: 52px !important;
    padding: 0 30px !important;
    border: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
    color: {theme["muted"]} !important;
    font-size: 14px !important;
    font-weight: 820 !important;
}}
div[class*="st-key-invoice_row_"],
div[class*="st-key-invoice_row_"] .invoice-list-row {{
    height: 112px !important;
    min-height: 112px !important;
    max-height: 112px !important;
}}
.invoice-list-row {{
    display: grid !important;
    grid-template-columns: minmax(0, 1fr) auto !important;
    align-items: start !important;
    gap: 18px !important;
    padding: 20px 30px !important;
    border: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
}}
.invoice-list-row.active {{
    padding-left: 26px !important;
    background: {theme["active_row"]} !important;
    border-left: 4px solid {theme["accent"]} !important;
}}
.invoice-list-row strong {{
    color: {theme["text"]} !important;
    font-size: 16px !important;
    line-height: 1.15 !important;
    font-weight: 850 !important;
}}
.invoice-list-row span {{
    color: {theme["muted"]} !important;
    margin-top: 9px !important;
    font-size: 14px !important;
    line-height: 1.25 !important;
}}
.invoice-list-row em {{
    color: {theme["faint"]} !important;
    margin-top: 12px !important;
    font-size: 12px !important;
    line-height: 1 !important;
}}
.invoice-row-meta b {{
    color: {theme["text"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: 15px !important;
    font-weight: 820 !important;
}}
div[data-testid="stHorizontalBlock"]:has(.detail-title-block) {{
    display: grid !important;
    grid-template-columns: minmax(0, 1fr) auto !important;
    gap: 24px !important;
    align-items: start !important;
    padding: 36px 42px 20px !important;
    background: {theme["bg"]} !important;
    border: 0 !important;
}}
.detail-title-block h2 {{
    margin: 16px 0 12px !important;
    color: {theme["text"]} !important;
    font-size: clamp(28px, 2.7vw, 38px) !important;
    line-height: 1.02 !important;
    font-weight: 850 !important;
    overflow-wrap: anywhere !important;
}}
.detail-title-block p {{
    color: {theme["muted"]} !important;
    font-size: 15px !important;
    line-height: 1.34 !important;
}}
.detail-badges {{
    min-width: 138px !important;
    display: flex !important;
    flex-direction: column !important;
    align-items: flex-end !important;
    gap: 10px !important;
}}
.status-pill {{
    min-height: 28px !important;
    padding: 0 14px !important;
    border-radius: 999px !important;
    font-size: 12px !important;
    line-height: 1 !important;
    font-weight: 820 !important;
}}
.detail-grid {{
    grid-template-columns: repeat(auto-fit, minmax(168px, 1fr)) !important;
    gap: 12px !important;
    padding: 0 42px 24px !important;
    margin: 0 !important;
    background: {theme["bg"]} !important;
}}
.detail-grid div {{
    min-height: 84px !important;
    padding: 15px 16px !important;
    border-radius: 9px !important;
    border: 1px solid {theme["border"]} !important;
    background: {theme["surface2"]} !important;
}}
.detail-grid span {{
    color: {theme["faint"]} !important;
    font-size: 11px !important;
    letter-spacing: .12em !important;
    font-weight: 850 !important;
}}
.detail-grid strong {{
    margin-top: 12px !important;
    color: {theme["text"]} !important;
    font-size: 18px !important;
    line-height: 1.12 !important;
    font-weight: 820 !important;
}}
div[class*="st-key-invoice_detail_actions"] {{
    margin: 0 42px !important;
    padding: 0 0 24px !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
div[class*="st-key-invoice_detail_actions"] div[data-testid="stHorizontalBlock"] {{
    display: grid !important;
    grid-template-columns: repeat(3, minmax(0, 1fr)) !important;
    gap: 12px !important;
}}
.ez-line-section-title {{
    padding: 32px 42px 0 !important;
    margin: 0 0 22px !important;
    color: {theme["text"]} !important;
    font-size: 20px !important;
    line-height: 1.2 !important;
    font-weight: 850 !important;
}}
.ez-line-table-wrap {{
    max-width: calc(100% - 84px) !important;
    margin: 0 42px 46px !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 10px !important;
    background: {theme["surface"]} !important;
}}
.ez-line-table th {{
    background: {theme["surface2"]} !important;
    color: {theme["muted"]} !important;
    padding: 12px 14px !important;
    font-size: 11px !important;
    letter-spacing: .06em !important;
}}
.ez-line-table td {{
    background: {theme["bg"]} !important;
    color: {theme["text"]} !important;
    padding: 13px 14px !important;
    font-size: 13px !important;
}}
.empty-pane:not(.large) {{
    height: calc(100vh - var(--ez-head-h) - 52px) !important;
    min-height: 430px !important;
}}
.empty-pane.large {{
    height: calc(100vh - var(--ez-head-h)) !important;
    min-height: 500px !important;
}}
.empty-pane,
.empty-pane.large {{
    display: grid !important;
    place-content: center !important;
    gap: 14px !important;
    padding: 0 42px !important;
    border: 0 !important;
    border-bottom: 1px solid {theme["border"]} !important;
    background: {theme["bg"]} !important;
    text-align: center !important;
}}
.empty-pane strong {{
    color: {theme["text"]} !important;
    font-size: 18px !important;
    line-height: 1.2 !important;
    font-weight: 850 !important;
}}
.empty-pane span {{
    color: {theme["muted"]} !important;
    font-size: 15px !important;
    line-height: 1.45 !important;
}}

.analytics-page,
.history-workspace {{
    padding: 36px 48px 86px !important;
    background: {theme["bg"]} !important;
}}
.analytics-kpis,
.history-kpis {{
    grid-template-columns: repeat(4, minmax(0, 1fr)) !important;
    gap: 20px !important;
    margin-bottom: 26px !important;
}}
.analytics-grid,
.history-grid {{
    gap: 24px !important;
    margin-bottom: 26px !important;
}}
.analytics-card,
.history-panel,
.history-table-wrap {{
    border: 1px solid {theme["border"]} !important;
    border-radius: 12px !important;
    background: {theme["surface2"]} !important;
    box-shadow: none !important;
}}
.analytics-card,
.history-panel {{
    padding: 24px 26px !important;
}}
.analytics-card h3,
.history-panel h3,
.history-table-pro th {{
    color: {theme["muted"]} !important;
    font-size: 12px !important;
    line-height: 1.1 !important;
    letter-spacing: .12em !important;
    text-transform: uppercase !important;
    font-weight: 850 !important;
}}
.metric-big {{
    color: {theme["text"]} !important;
    font-family: 'JetBrains Mono', ui-monospace, monospace !important;
    font-size: clamp(25px, 2.4vw, 34px) !important;
    line-height: 1 !important;
    font-weight: 820 !important;
    background: transparent !important;
}}
.metric-big.green {{
    color: {theme["accent"]} !important;
    background: transparent !important;
    padding: 0 !important;
    border: 0 !important;
    box-shadow: none !important;
}}
.metric-big.amber {{
    color: {theme["warn"]} !important;
}}
.bar-fill,
.currency-pill {{
    background: {theme["accent"]} !important;
    color: #fff !important;
}}
.donut {{
    filter: saturate(.88) contrast(.98) !important;
}}
.history-stage {{
    grid-template-columns: 34px minmax(0, 1fr) 46px !important;
    min-height: 60px !important;
}}
.history-stage .dot {{
    background: {theme["chip_validated_bg"]} !important;
    color: {theme["accent"]} !important;
}}
.history-stage > strong:last-child {{
    justify-self: end !important;
    text-align: right !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stRadio"] div[role="radiogroup"] {{
    height: 38px !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 9px !important;
    background: {theme["surface"]} !important;
    overflow: hidden !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stRadio"] div[role="radiogroup"] label:has(input:checked) {{
    background: {theme["accent"]} !important;
}}
div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) div[data-testid="stRadio"] div[role="radiogroup"] label:has(input:checked) p {{
    color: #FFFFFF !important;
}}

div[data-testid="stPopoverBody"]:has(.ez-filter-panel) {{
    width: min(350px, calc(100vw - 36px)) !important;
    max-width: calc(100vw - 36px) !important;
    padding: 20px !important;
    border: 1px solid {theme["border"]} !important;
    border-radius: 12px !important;
    background: {theme["surface"]} !important;
    color: {theme["text"]} !important;
    box-shadow: var(--ez-pop-shadow) !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) > div,
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stVerticalBlock"],
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stElementContainer"],
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stMarkdown"],
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stCheckbox"],
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stRadio"] {{
    background: transparent !important;
    color: {theme["text"]} !important;
}}
.ez-filter-panel {{
    margin: 0 0 16px !important;
    padding: 0 0 14px !important;
    border-bottom: 1px solid {theme["border"]} !important;
}}
.ez-filter-panel-head {{
    display: grid !important;
    gap: 5px !important;
}}
.ez-filter-panel-head span,
.ez-filter-menu-title {{
    color: {theme["faint"]} !important;
    font-size: 11px !important;
    line-height: 1.1 !important;
    font-weight: 850 !important;
    letter-spacing: .12em !important;
    text-transform: uppercase !important;
}}
.ez-filter-panel-head strong {{
    color: {theme["text"]} !important;
    font-size: 18px !important;
    line-height: 1.15 !important;
    font-weight: 850 !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) .ez-filter-menu-title {{
    display: block !important;
    margin: 0 0 9px !important;
    color: {theme["muted"]} !important;
}}
.ez-filter-section-break {{
    height: 1px !important;
    margin: 16px 0 !important;
    background: {theme["border"]} !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) label,
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) label p,
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stMarkdownContainer"] p {{
    color: {theme["text"]} !important;
    opacity: 1 !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stCheckbox"] {{
    margin: 0 !important;
    padding: 2px 0 !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stCheckbox"] label,
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stRadio"] div[role="radiogroup"] label {{
    min-height: 34px !important;
    padding: 4px 0 !important;
    gap: 10px !important;
    align-items: center !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stCheckbox"] label p,
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stRadio"] label p {{
    color: {theme["text"]} !important;
    font-size: 14px !important;
    line-height: 1.2 !important;
    font-weight: 650 !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stRadio"] {{
    margin: 0 !important;
    padding: 0 !important;
    border-top: 0 !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) div[data-testid="stRadio"] div[role="radiogroup"] {{
    gap: 2px !important;
    margin: 0 !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) label > div:first-child {{
    border-color: {theme["border2"]} !important;
    background: {theme["input"]} !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) label:has(input:checked) > div:first-child {{
    border-color: {theme["accent"]} !important;
    background: {theme["accent"]} !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) svg {{
    color: #FFFFFF !important;
    fill: currentColor !important;
}}
div[data-testid="stPopoverBody"]:has(.ez-filter-panel) input {{
    accent-color: {theme["accent"]} !important;
}}

@media (max-width: 1180px) {{
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) {{
        grid-template-columns: minmax(220px, 1fr) minmax(240px, 1fr) 118px 164px !important;
        padding: 0 24px !important;
        gap: 12px !important;
    }}
    .ap-status {{
        padding: 0 20px !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.detail-title-block),
    .detail-grid,
    .ez-line-section-title {{
        padding-left: 30px !important;
        padding-right: 30px !important;
    }}
    div[class*="st-key-invoice_detail_actions"] {{
        margin-left: 30px !important;
        margin-right: 30px !important;
    }}
    .ez-line-table-wrap {{
        max-width: calc(100% - 60px) !important;
        margin-left: 30px !important;
        margin-right: 30px !important;
    }}
}}
@media (max-width: 900px) {{
    div[data-testid="stHorizontalBlock"]:has(.ez-page-heading) {{
        height: auto !important;
        max-height: none !important;
        grid-template-columns: 1fr !important;
        padding: 16px !important;
        align-items: stretch !important;
    }}
    .ap-status-strip {{
        height: auto !important;
        max-height: none !important;
        grid-template-columns: 1fr !important;
    }}
    .ap-status {{
        min-height: 54px !important;
        border-right: 0 !important;
        border-bottom: 1px solid {theme["border"]} !important;
    }}
    div[class*="st-key-batch_exports"] > div[data-testid="stVerticalBlock"] > div[data-testid="stHorizontalBlock"],
    div[data-testid="stHorizontalBlock"]:has(.pane-title) {{
        display: block !important;
    }}
    div[data-testid="stHorizontalBlock"]:has(.pane-title) > div:has(.pane-title) {{
        border-right: 0 !important;
    }}
    .analytics-kpis,
    .history-kpis,
    .analytics-grid,
    .analytics-grid.three,
    .history-grid {{
        grid-template-columns: 1fr !important;
    }}
}}

/* Integration workspace rhythm and typography */
section[data-testid="stMain"]:has(.integration-hero)
div[data-testid="stMarkdownContainer"]:has(.integration-hero) {{
    margin-bottom: 0 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-hero {{
    min-height: 112px !important;
    grid-template-columns: 48px minmax(0, 1fr) !important;
    gap: 18px !important;
    padding: 20px 32px !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-brand-mark {{
    width: 48px !important;
    height: 48px !important;
    border-radius: 8px !important;
    font-size: .82rem !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-hero-copy > .integration-kind {{
    margin: 0 0 4px !important;
    font-size: .64rem !important;
    line-height: 1.2 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-title-row {{
    display: flex !important;
    align-items: center !important;
    justify-content: flex-start !important;
    gap: 12px !important;
    flex-wrap: wrap !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-system-name {{
    display: block !important;
    margin: 0 !important;
    padding: 0 !important;
    color: {theme["text"]} !important;
    font-size: 1.28rem !important;
    line-height: 1.2 !important;
    font-weight: 800 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-hero-copy p {{
    max-width: 900px !important;
    margin: 6px 0 0 !important;
    padding: 0 !important;
    font-size: .8rem !important;
    line-height: 1.45 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-status {{
    min-height: 32px !important;
    padding: 6px 11px !important;
    font-size: .7rem !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-flow {{
    width: calc(100% - 64px) !important;
    margin: 18px 32px 20px !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-flow > div {{
    min-height: 64px !important;
    gap: 10px !important;
    padding: 11px 14px !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-flow b {{
    flex-basis: 26px !important;
    width: 26px !important;
    height: 26px !important;
    font-size: .68rem !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-flow strong {{
    font-size: .76rem !important;
    line-height: 1.25 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-flow small {{
    margin-top: 2px !important;
    font-size: .64rem !important;
    line-height: 1.25 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) div[class*="st-key-integration_page_body"] {{
    padding: 0 32px 32px !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-section-heading {{
    min-height: 54px !important;
    align-items: center !important;
    gap: 20px !important;
    margin: 0 0 18px !important;
    padding: 5px 0 13px !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-section-heading span {{
    margin-bottom: 4px !important;
    font-size: .63rem !important;
    line-height: 1.2 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-section-heading strong {{
    font-size: .94rem !important;
    line-height: 1.3 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-section-heading p {{
    max-width: 500px !important;
    font-size: .72rem !important;
    line-height: 1.4 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stTabs"] [data-baseweb="tab-list"] {{
    margin: 0 0 16px !important;
    padding: 3px !important;
}}
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stTabs"] [data-baseweb="tab"] {{
    min-height: 38px !important;
    padding: 0 14px !important;
    font-size: .76rem !important;
    line-height: 1.2 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) div[class*="st-key-quickbooks_connector_workspace"],
section[data-testid="stMain"]:has(.integration-hero) div[class*="st-key-zoho_connector_workspace"] {{
    padding: 20px 22px 22px !important;
}}
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stHorizontalBlock"] {{
    gap: 16px !important;
}}
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stTextInput"] label p,
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stTextArea"] label p,
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stSelectbox"] label p,
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stNumberInput"] label p {{
    margin: 0 0 5px !important;
    color: {theme["muted"]} !important;
    font-size: .75rem !important;
    line-height: 1.3 !important;
    font-weight: 700 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stTextInputRootElement"],
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stNumberInputContainer"],
section[data-testid="stMain"]:has(.integration-hero) [data-baseweb="select"] > div {{
    min-height: 40px !important;
    border-radius: 7px !important;
}}
section[data-testid="stMain"]:has(.integration-hero) input,
section[data-testid="stMain"]:has(.integration-hero) textarea,
section[data-testid="stMain"]:has(.integration-hero) [data-baseweb="select"] * {{
    font-size: .84rem !important;
    line-height: 1.4 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) input::placeholder,
section[data-testid="stMain"]:has(.integration-hero) textarea::placeholder {{
    color: {theme["faint"]} !important;
    opacity: 1 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) textarea {{
    min-height: 74px !important;
    padding: 10px 12px !important;
}}
section[data-testid="stMain"]:has(.integration-hero)
div[data-testid="stAlert"]:has([data-testid="stAlertContentWarning"])
div[data-testid="stAlertContainer"] {{
    background: {theme["chip_pending_bg"]} !important;
    color: {theme["warn"]} !important;
}}
section[data-testid="stMain"]:has(.integration-hero)
div[data-testid="stAlert"]:has([data-testid="stAlertContentSuccess"])
div[data-testid="stAlertContainer"] {{
    background: {theme["chip_validated_bg"]} !important;
    color: {theme["good"]} !important;
}}
section[data-testid="stMain"]:has(.integration-hero)
div[data-testid="stAlert"]:has([data-testid="stAlertContentError"])
div[data-testid="stAlertContainer"] {{
    background: {theme["chip_exception_bg"]} !important;
    color: {theme["danger"]} !important;
}}
section[data-testid="stMain"]:has(.integration-hero)
div[data-testid="stAlert"]:has([data-testid="stAlertContentInfo"])
div[data-testid="stAlertContainer"] {{
    background: {theme["active_row"]} !important;
    color: {theme["accent"]} !important;
}}
section[data-testid="stMain"]:has(.integration-hero)
div[data-testid="stAlert"] div[data-testid="stAlertContainer"] * {{
    color: inherit !important;
}}
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stExpander"] {{
    margin-bottom: 12px !important;
    border-radius: 8px !important;
}}
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stExpander"] summary {{
    min-height: 44px !important;
    padding: 0 14px !important;
    font-size: .8rem !important;
}}
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stButton"] > button,
section[data-testid="stMain"]:has(.integration-hero) div[data-testid="stDownloadButton"] > button {{
    min-height: 40px !important;
    border-radius: 7px !important;
    font-size: .8rem !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-readiness {{
    margin: 0 0 20px !important;
    padding: 20px 22px !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-readiness h3 {{
    margin: 0 !important;
    padding: 0 !important;
    font-size: 1rem !important;
    line-height: 1.3 !important;
    font-weight: 800 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-readiness p {{
    margin: 6px 0 0 !important;
}}
section[data-testid="stMain"]:has(.integration-hero) .integration-requirements {{
    gap: 14px !important;
}}
@media (max-width: 900px) {{
    :root {{
        --ez-sidebar-w: 220px;
    }}
    section[data-testid="stSidebar"] {{
        width: var(--ez-sidebar-w) !important;
        min-width: var(--ez-sidebar-w) !important;
        max-width: var(--ez-sidebar-w) !important;
        transform: none !important;
    }}
    section[data-testid="stMain"]:has(.integration-hero) div[class*="st-key-integration_page_body"] {{
        padding: 0 20px 24px !important;
    }}
    section[data-testid="stMain"]:has(.integration-hero) .integration-flow {{
        width: calc(100% - 40px) !important;
        margin-left: 20px !important;
        margin-right: 20px !important;
    }}
}}
@media (max-width: 640px) {{
    :root {{
        --ez-sidebar-w: 190px;
    }}
    section[data-testid="stSidebar"] {{
        transform: translateX(-100%) !important;
    }}
    section[data-testid="stMain"]:has(.integration-hero) .integration-hero {{
        grid-template-columns: 42px minmax(0, 1fr) !important;
        padding: 18px 16px !important;
    }}
    section[data-testid="stMain"]:has(.integration-hero) .integration-brand-mark {{
        width: 42px !important;
        height: 42px !important;
    }}
    section[data-testid="stMain"]:has(.integration-hero) .integration-title-row {{
        align-items: flex-start !important;
        flex-direction: column !important;
        gap: 8px !important;
    }}
    section[data-testid="stMain"]:has(.integration-hero) .integration-section-heading {{
        align-items: flex-start !important;
        flex-direction: column !important;
        gap: 7px !important;
    }}
    section[data-testid="stMain"]:has(.integration-hero) .integration-section-heading p {{
        text-align: left !important;
    }}
}}
</style>
""", unsafe_allow_html=True)

# ── Header ────────────────────────────────────────────
_workspace_subtitles = {
    "Invoices": "Upload, extract, validate, and review supplier invoices in one AP workspace.",
    "History": "Searchable audit trail for processed invoices and posting readiness.",
    "Analytics": "Spend, exception, supplier, currency, and parser performance intelligence.",
    "Exceptions": "Focused review queue for invoices that need human attention.",
    "QuickBooks": "QuickBooks connector status, mappings, and bill posting controls.",
    "Tally": "TallyPrime XML voucher settings and posting controls.",
    "Zoho Books": "Zoho Books organization, account mapping, and bill posting controls.",
    "Coupa": "Coupa-ready export routing.",
    "NetSuite": "NetSuite-ready export routing.",
    "SAP": "SAP-ready export routing.",
    "GL Mapping": "Client category and GL mapping controls.",
    "Rules": "Automation rules for routing and review decisions.",
    "Vendors": "Supplier profiles and learning history.",
    "Roadmap": "Product roadmap from local prototype to B2B SaaS platform.",
}
GLYPH = {
    "Invoices": "▤",
    "History": "↺",
    "Analytics": "▥",
    "Exceptions": "⚠",
    "QuickBooks": "∞",
    "Tally": "⊞",
    "Zoho Books": "Z",
    "Coupa": "▦",
    "NetSuite": "⌘",
    "SAP": "▦",
    "GL Mapping": "◇",
    "Rules": "≡",
    "Vendors": "⌾",
    "Roadmap": "→",
}


def _nav_slug(page: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", page.lower()).strip("_")


def render_shell_sidebar() -> str:
    rows = _activity_rows_for_console()
    invoice_count = len(rows)
    exception_count = sum(1 for r in rows if str(r.get("Status", "")).lower() not in ("ready", "sent"))
    history_count = invoice_count

    def nav_row(page: str, label: str, count: Optional[int] = None, dot: bool = False) -> str:
        active = page == st.session_state.get("ez_nav_page", "Invoices")
        right = f'<span class="side-nav-count">{count}</span>' if count is not None else ""
        if dot:
            right = '<span class="side-red-dot"></span>'
        return (
            f'<div class="side-nav-row{" active" if active else ""}">'
            f'<span class="side-icon">{GLYPH.get(page, "□")}</span>'
            f'<span>{_ui_escape(label)}</span>{right}</div>'
        )

    def nav_button(page: str, label: str, count: Optional[int] = None, dot: bool = False, target: Optional[str] = None) -> None:
        slug = _nav_slug(page)
        with st.container(key=f"nav_item_{slug}"):
            st.markdown(nav_row(page, label, count=count, dot=dot), unsafe_allow_html=True)
            if st.button(label, key=f"nav_btn_{slug}"):
                st.session_state["ez_nav_page"] = page
                if target:
                    st.session_state["ez_target_system"] = target
                    st.session_state["ez_fmt_choice"] = target
                    st.query_params["target"] = target
                st.query_params["page"] = page
                st.rerun()

    with st.sidebar:
        st.markdown(
            '<div class="ez-brand"><div class="ez-brand-mark">EZ</div>'
            '<div><div class="ez-brand-name">EZ-Invoice</div><div class="ez-brand-ver">v1.0 platform</div></div></div>',
            unsafe_allow_html=True,
        )
        with st.container(key="ez_nav_shell"):
            st.markdown('<div class="nav-section">Workflow</div>', unsafe_allow_html=True)
            nav_button("Invoices", "Invoices", invoice_count)
            nav_button("History", "History", history_count)
            nav_button("Analytics", "Analytics")
            nav_button("Exceptions", "Exceptions", dot=exception_count > 0)
            st.markdown('<div class="nav-section">Send to ERP</div>', unsafe_allow_html=True)
            nav_button("QuickBooks", "QuickBooks", target="QuickBooks")
            nav_button("Tally", "Tally", target="Tally")
            nav_button("Zoho Books", "Zoho Books", target="Zoho Books")
            nav_button("Coupa", "Coupa", target="Coupa")
            nav_button("NetSuite", "NetSuite", target="NetSuite")
            nav_button("SAP", "SAP", target="SAP")
            st.markdown('<div class="nav-section">Settings</div>', unsafe_allow_html=True)
            nav_button("GL Mapping", "GL Mapping")
            nav_button("Rules", "Rules")
            nav_button("Vendors", "Vendors")
            nav_button("Roadmap", "Roadmap")

        with st.container(key="ez_sidebar_footer"):
            if "ez_dark_mode_toggle" not in st.session_state:
                st.session_state["ez_dark_mode_toggle"] = st.session_state.get("ez_theme_mode", "Light") == "Dark"
            dark_toggle = st.toggle("Dark mode", key="ez_dark_mode_toggle")
            wanted = "Dark" if dark_toggle else "Light"
            if wanted != st.session_state.get("ez_theme_mode", "Light"):
                st.session_state["ez_theme_mode"] = wanted
                st.rerun()

            qb_dot = theme["good"] if is_connected() else theme["danger"]
            qb_s = "QuickBooks connected" if is_connected() else "QuickBooks not connected"
            tally_dot = theme["warn"] if tally_is_connected() else theme["danger"]
            tally_s = "Tally configured" if tally_is_connected() else "Tally setup"
            zoho_dot = theme["good"] if zoho_is_connected() else theme["danger"]
            zoho_s = "Zoho Books connected" if zoho_is_connected() else "Zoho Books not connected"
            st.markdown(
                f'<div class="sidebar-status">'
                f'<div class="sidebar-status-line"><span class="status-dot" style="background:{qb_dot};"></span>{qb_s}</div>'
                f'<div class="sidebar-status-line"><span class="status-dot" style="background:{tally_dot};"></span>{tally_s}</div>'
                f'<div class="sidebar-status-line"><span class="status-dot" style="background:{zoho_dot};"></span>{zoho_s}</div>'
                f'</div>',
                unsafe_allow_html=True,
            )

    return st.session_state.get("ez_nav_page", "Invoices")


def render_app_topbar(title: str) -> Tuple[List[Any], bool, bool, bool, bool, str, str, bool]:
    def heading_html(page_title: str) -> str:
        crumbs = {
            "Invoices": "Queue",
            "History": "Audit trail",
            "Analytics": "Spend overview",
            "Exceptions": "Review",
            "GL Mapping": "Ledger mapping",
            "Rules": "Automation rules",
            "Vendors": "Vendor master",
            "Roadmap": "Build plan",
        }
        crumb = crumbs.get(page_title, "Integration" if page_title in ERP_SYSTEMS else "")
        crumb_html = f'<span class="slash">/</span><span class="crumb">{_ui_escape(crumb)}</span>' if crumb else ""
        return f'<div class="ez-page-heading"><span class="title">{_ui_escape(page_title)}</span>{crumb_html}</div>'

    uploaded_value: List[Any] = []
    run_value = False
    clear_queue_value = False

    if title == "History":
        c_title, c_search, _gap, c_export = st.columns([0.24, 0.34, 0.22, 0.17])
        with c_title:
            st.markdown(heading_html(title), unsafe_allow_html=True)
        with c_search:
            st.text_input("Search history", placeholder="Search...", key="history_search_topbar", label_visibility="collapsed")
        with c_export:
            history_export = combined_activity_df()
            st.download_button(
                "Export CSV",
                history_export.to_csv(index=False).encode("utf-8") if not history_export.empty else b"",
                "ez_invoice_history.csv",
                "text/csv",
                key="history_topbar_export_csv",
                width="stretch",
                type="primary",
            )
        return (
            uploaded_value,
            run_value,
            st.session_state.get("sidebar_show_raw_v2", False),
            st.session_state.get("sidebar_show_evidence_v2", True),
            st.session_state.get("sidebar_show_analytics_v2", True),
            st.session_state.get("ez_parser_mode", "Auto"),
            st.session_state.get("ez_fmt_choice", "QuickBooks"),
            clear_queue_value,
        )

    if title == "Analytics":
        c_title, c_range, _gap, c_export = st.columns([0.36, 0.34, 0.10, 0.17])
        with c_title:
            st.markdown(heading_html(title), unsafe_allow_html=True)
        with c_range:
            st.radio(
                "Date window",
                ["7d", "30d", "90d", "All"],
                index=1,
                horizontal=True,
                key="analytics_range",
                label_visibility="collapsed",
            )
        with c_export:
            analytics_export = combined_activity_df()
            st.download_button(
                "Export CSV",
                analytics_export.to_csv(index=False).encode("utf-8") if not analytics_export.empty else b"",
                "ez_invoice_analytics.csv",
                "text/csv",
                key="analytics_topbar_export_csv",
                width="stretch",
                type="primary",
            )
        return (
            uploaded_value,
            run_value,
            st.session_state.get("sidebar_show_raw_v2", False),
            st.session_state.get("sidebar_show_evidence_v2", True),
            st.session_state.get("sidebar_show_analytics_v2", True),
            st.session_state.get("ez_parser_mode", "Auto"),
            st.session_state.get("ez_fmt_choice", "QuickBooks"),
            clear_queue_value,
        )

    if title != "Invoices":
        c_title, _c2 = st.columns([0.42, 0.58])
        with c_title:
            st.markdown(heading_html(title), unsafe_allow_html=True)
        return (
            uploaded_value,
            run_value,
            st.session_state.get("sidebar_show_raw_v2", False),
            st.session_state.get("sidebar_show_evidence_v2", True),
            st.session_state.get("sidebar_show_analytics_v2", True),
            st.session_state.get("ez_parser_mode", "Auto"),
            st.session_state.get("ez_fmt_choice", "QuickBooks"),
            clear_queue_value,
        )

    c_title, c_search, c_filter, c_upload = st.columns([0.24, 0.34, 0.13, 0.17])
    with c_title:
        st.markdown(heading_html(title), unsafe_allow_html=True)
    with c_search:
        st.text_input("Search", placeholder="Search...", key="topbar_search", label_visibility="collapsed")
    with c_filter:
        st.markdown('<span class="ez-filter-slot"></span>', unsafe_allow_html=True)
        with st.popover("Filter", width="stretch"):
            st.markdown(
                '<div class="ez-filter-panel">'
                '<div class="ez-filter-panel-head">'
                '<span>Filter workspace</span>'
                '<strong>Review controls</strong>'
                '</div>'
                '</div>'
                '<div class="ez-filter-menu-title">Display options</div>',
                unsafe_allow_html=True,
            )
            show_raw_value = st.checkbox("Show raw text", value=False, key="sidebar_show_raw_v2")
            show_evidence_value = st.checkbox("Show extraction evidence", value=True, key="sidebar_show_evidence_v2")
            show_analytics_value = st.checkbox("Show spend analytics", value=True, key="sidebar_show_analytics_v2")
            st.markdown(
                '<div class="ez-filter-section-break"></div>'
                '<div class="ez-filter-menu-title">Invoice parser</div>',
                unsafe_allow_html=True,
            )
            parser_mode_value = st.radio(
                "Invoice parser",
                ["Auto", "GST/e-Invoice adapter", "Structured adapter", "Universal extraction"],
                key="ez_parser_mode",
                label_visibility="collapsed",
            )
            target_options = ["QuickBooks", "Tally", "Zoho Books", "Coupa", "NetSuite", "SAP"]
            selected_target = st.session_state.get("ez_fmt_choice", st.session_state.get("ez_target_system", "QuickBooks"))
            if selected_target not in target_options:
                selected_target = "QuickBooks"
            st.markdown(
                '<div class="ez-filter-section-break"></div>'
                '<div class="ez-filter-menu-title">Primary target system</div>',
                unsafe_allow_html=True,
            )
            fmt_choice_value = st.radio(
                "Primary target system",
                target_options,
                index=target_options.index(selected_target),
                key="ez_fmt_choice",
                label_visibility="collapsed",
            )
    with c_upload:
        st.markdown('<span class="ez-upload-slot"></span>', unsafe_allow_html=True)
        upload_batch_id = int(st.session_state.get("upload_batch_id", 0))
        with st.popover(
            "Upload PDFs",
            width="stretch",
            key="topbar_upload_popover_" + str(upload_batch_id),
        ):
            upload_key = "topbar_pdf_uploader_" + str(upload_batch_id)
            uploaded_value = st.file_uploader(
                "Upload invoice PDFs",
                type=["pdf"],
                accept_multiple_files=True,
                key=upload_key,
                label_visibility="collapsed",
            )
            run_value = st.button("Process PDFs", type="primary", width="stretch", key="topbar_process_all")
            clear_queue_value = st.button("Clear queue", width="stretch", key="topbar_clear_queue")

    return (
        uploaded_value or [],
        run_value,
        st.session_state.get("sidebar_show_raw_v2", False),
        st.session_state.get("sidebar_show_evidence_v2", True),
        st.session_state.get("sidebar_show_analytics_v2", True),
        st.session_state.get("ez_parser_mode", "Auto"),
        st.session_state.get("ez_fmt_choice", "QuickBooks"),
        clear_queue_value,
    )


if "payloads" not in st.session_state: st.session_state.payloads = {}
if "summary" not in st.session_state: st.session_state.summary = []
if "ptimes" not in st.session_state: st.session_state.ptimes = {}
if "upload_batch_id" not in st.session_state: st.session_state.upload_batch_id = 0

_initialize_api_bridge()
if _api_bridge_active():
    try:
        _sync_api_queue()
    except EzInvoiceApiError as exc:
        st.session_state["ez_api_bridge_active"] = False
        st.session_state["ez_api_bridge_error"] = str(exc)
        if st.session_state.get("ez_backend_mode_requested") == "api":
            st.error("FastAPI queue synchronization failed: " + str(exc))

nav_page = render_shell_sidebar()
current_page = nav_page
uploaded, run, show_raw, show_evidence, show_analytics, parser_mode, fmt_choice, clear_queue = render_app_topbar(current_page)

if clear_queue:
    if _api_bridge_active():
        client = _api_bridge_client()
        organization = st.session_state.get("ez_api_organization") or {}
        try:
            if client and organization.get("id"):
                client.clear_invoices(str(organization["id"]))
        except EzInvoiceApiError as exc:
            st.error("The FastAPI invoice queue could not be cleared: " + str(exc))
            st.stop()
    st.session_state.payloads = {}
    st.session_state.summary = []
    st.session_state.ptimes = {}
    st.session_state.pop("tally_post_status", None)
    st.session_state.upload_batch_id += 1
    st.toast("Current invoice queue cleared.")
    st.rerun()

# ── Processing ────────────────────────────────────────
if run and uploaded:
    st.session_state.pop("selected_invoice_file", None)
    prog = st.progress(0, "Processing invoices...")
    if _api_bridge_active():
        client = _api_bridge_client()
        organization = st.session_state.get("ez_api_organization") or {}
        uploaded_invoice_ids: List[str] = []
        try:
            if not client or not organization.get("id"):
                raise EzInvoiceApiError("FastAPI organization is not initialized.")
            for i, f in enumerate(uploaded):
                t0 = datetime.now()
                pdf_bytes = f.read()
                invoice = client.upload_invoice(
                    organization_id=str(organization["id"]),
                    filename=f.name,
                    pdf_bytes=pdf_bytes,
                    parser_mode=parser_mode_for_api(parser_mode),
                )
                validation = client.validate_invoice(str(invoice["id"]))
                invoice["status"] = validation.get("status", invoice.get("status"))
                invoice["validation_issues"] = validation.get("issues", [])
                uploaded_invoice_ids.append(str(invoice["id"]))
                ms = (datetime.now() - t0).total_seconds() * 1000
                st.session_state.ptimes[f.name] = ms
                prog.progress(
                    (i + 1) / len(uploaded),
                    f"Processing {f.name} - {ms:.0f}ms",
                )
            id_to_key = _sync_api_queue()
            if uploaded_invoice_ids:
                selected_key = id_to_key.get(uploaded_invoice_ids[0])
                if selected_key:
                    st.session_state["selected_invoice_file"] = selected_key
        except EzInvoiceApiError as exc:
            prog.empty()
            st.error("FastAPI invoice processing failed: " + str(exc))
            st.stop()
    else:
        st.session_state.payloads = {}
        st.session_state.summary = []
        st.session_state.ptimes = {}
        for i, f in enumerate(uploaded):
            t0 = datetime.now()

            pdf_bytes = f.read()
            payload, text, engine, pages = parse_ez_invoice(f.name, pdf_bytes, parser_mode=parser_mode)
            ok, issues, df_missing = capture_check(payload)
            ms = (datetime.now() - t0).total_seconds() * 1000

            inv = payload["INVOICE"]; h = inv["INVOICE HEADER"]

            processed_data = {
                "payload": payload, "text": text, "bytes": pdf_bytes,
                "engine": engine, "pages": pages, "ok": ok,
                "issues": issues, "missing_df": df_missing
            }
            st.session_state.payloads[f.name] = processed_data
            _history_append(_history_record(f.name, processed_data, ms))
            st.session_state.ptimes[f.name] = ms
            st.session_state.summary.append({
                "File": f.name,
                "Invoice #": h.get("INVOICE NO.", ""),
                "Direction": inv.get("ROUTING", {}).get("DIRECTION", "unknown"),
                "Date": h.get("INVOICE DATE", ""),
                "Items": len(inv.get("LINE ITEMS", {}).get("ROWS", []) or []),
                "Total": f"{inv.get('PAYMENT',{}).get('ELECTRONIC',{}).get('CURRENCY','USD')} {money(h.get('INVOICE AMOUNT',0))}",
                "Status": "✅" if ok else "⚠️",
                "Mode": inv.get("PARSER_DIAGNOSTICS", {}).get("mode", ""),
                "Pages": pages,
                "Speed": f"{ms:.0f}ms"
            })
            prog.progress((i + 1) / len(uploaded), f"Processing {f.name} - {ms:.0f}ms")
            if i == 0:
                st.session_state["selected_invoice_file"] = f.name

    prog.empty()
    st.session_state.upload_batch_id += 1
    st.session_state["ez_processed_toast"] = f"{len(uploaded)} invoice(s) processed."
    st.rerun()

processed_toast = st.session_state.pop("ez_processed_toast", None)
if processed_toast:
    st.toast(processed_toast, icon="✅")

if nav_page == "History":
    render_history_page()
    st.stop()
elif nav_page == "Analytics":
    render_analytics_page()
    st.stop()
elif nav_page == "Exceptions":
    render_exceptions_page()
    st.stop()
elif nav_page in ("QuickBooks", "Tally", "Zoho Books", "Coupa", "NetSuite", "SAP"):
    render_integrations_page()
    st.stop()
elif nav_page == "GL Mapping":
    _page_heading("GL Mapping", "Client category and ledger mapping")
    classifier_settings_ui(classifier)
    st.stop()
elif nav_page in ("Rules", "Vendors"):
    _page_heading(nav_page, "Configuration workspace")
    st.info("This workspace is ready for the next build step.")
    st.stop()
elif nav_page == "Roadmap":
    render_roadmap_page()
    st.stop()

# ── Results ───────────────────────────────────────────
if st.session_state.summary:
    df = pd.DataFrame(st.session_state.summary)
    P = st.session_state.payloads
    _console_rendered = render_invoice_console(fmt_choice)
    st.stop()

    n = len(df)
    ok_n = int(df["Status"].str.contains("✅").sum())
    items_n = int(df["Items"].sum())
    amt = sum(safe_float0(d["payload"]["INVOICE"]["INVOICE HEADER"].get("INVOICE AMOUNT",0)) for d in P.values())
    tax = sum(safe_float0(d["payload"]["INVOICE"].get("USD USD USD",{}).get("TOTAL TAX",0)) for d in P.values())
    avg_ms = sum(st.session_state.ptimes.values()) / max(len(st.session_state.ptimes), 1)

    st.markdown(f"""<div class="kpi-row">
    <div class="kpi"><div class="n">{n}</div><div class="l">Invoices</div></div>
    <div class="kpi"><div class="n g">{ok_n}/{n}</div><div class="l">Passed</div></div>
    <div class="kpi"><div class="n">{items_n}</div><div class="l">Line Items</div></div>
    <div class="kpi"><div class="n p">${amt:,.2f}</div><div class="l">Total Amount</div></div>
    <div class="kpi"><div class="n a">{avg_ms:.0f}ms</div><div class="l">Avg Speed</div></div>
    </div>""", unsafe_allow_html=True)

    # Duplicate check
    inv_nums = [d["payload"]["INVOICE"]["INVOICE HEADER"].get("INVOICE NO.","") for d in P.values()]
    dupes = [x for x in set(inv_nums) if inv_nums.count(x) > 1 and x]
    if dupes:
        st.warning(f"⚠️ **Duplicate invoice(s):** {', '.join(dupes)}")



    # ── Analytics ────────────────────────────────────
    if show_analytics and n > 0:
        st.markdown("---")
        st.markdown("#### 📊 Spend Analytics")
        ac1, ac2 = st.columns(2)
        with ac1:
            cd = [{"Invoice": d["payload"]["INVOICE"]["INVOICE HEADER"].get("INVOICE NO.", fn)[:16],
                   "Amount": safe_float0(d["payload"]["INVOICE"]["INVOICE HEADER"].get("INVOICE AMOUNT",0))}
                  for fn, d in P.items()]
            if cd: st.bar_chart(pd.DataFrame(cd).set_index("Invoice"), color=theme["accent"], height=260)
        with ac2:
            ai = []
            for d in P.values():
                for r in (d["payload"]["INVOICE"].get("LINE ITEMS",{}).get("ROWS",[]) or []):
                    ai.append({"Item": r.get("DESCRIPTION","")[:25], "Amount": safe_float0(r.get("AMOUNT",0))})
            if ai:
                top = pd.DataFrame(ai).groupby("Item")["Amount"].sum().nlargest(8).reset_index()
                st.bar_chart(top.set_index("Item"), color=theme["warn"], height=260)

    # ── Table ────────────────────────────────────────
    st.markdown("---")
    st.markdown("#### 📋 Processed Invoices")
    st.dataframe(df, width="stretch", hide_index=True)

    # ── Per-Invoice ──────────────────────────────────
    st.markdown("---")
    st.markdown("#### 🔍 Invoice Details")

    for fname, data in P.items():
        payload = data["payload"]; inv = payload["INVOICE"]; h = inv["INVOICE HEADER"]
        ok = data["ok"]; issues = data["issues"]; missing_df = data["missing_df"]
        ms = st.session_state.ptimes.get(fname, 0)

        rows = inv.get("LINE ITEMS",{}).get("ROWS",[]) or []
        inv_total = safe_float0(h.get("INVOICE AMOUNT",0))
        line_sum = sum(safe_float0(r.get("AMOUNT",0)) for r in rows)
        conf = max(0, 100 - (0 if h.get("INVOICE NO.") else 25) - (0 if rows else 30) - (0 if abs(line_sum-inv_total)<0.5 else 20) - (0 if h.get("DUE DATE") else 5))
        cc = "h" if conf >= 80 else ("m" if conf >= 50 else "lo")

        title = f"{'✅' if ok else '⚠️'} **{fname}** — `{h.get('INVOICE NO.','')}` — ${money(h.get('INVOICE AMOUNT',0))} — {ms:.0f}ms"

        with st.expander(title, expanded=False):
            # Status bar
            st.markdown(f"""<div style="display:flex;align-items:center;gap:8px;margin-bottom:14px;flex-wrap:wrap;">
            <span class="{'bok' if ok else 'bwn'}">{'Passed' if ok else 'Review'}</span>
            <span class="binf">{data['engine']}</span>
            <span class="binf">{data['pages']}p</span>
            <span class="binf">{len(rows)} items</span>
            <span class="binf">{ms:.0f}ms</span>
            <div style="flex:1;min-width:100px;">
                <div style="font-size:.65rem;opacity:0.6;">Confidence {conf}%</div>
                <div class="cb"><div class="cf {cc}" style="width:{conf}%"></div></div>
            </div>
            </div>""", unsafe_allow_html=True)

            if issues:
                for it in issues: st.warning(f"⚠️ {it}")

            st.markdown("##### 📋 Key Fields")
            for lbl, val, ic in [
                ("Invoice #", h.get("INVOICE NO.", ""), "🔢"),
                ("Date", h.get("INVOICE DATE", ""), "📅"),
                ("Due Date", h.get("DUE DATE", ""), "⏰"),
                ("Sales Order", h.get("SALES ORDER NO.", ""), "📄"),
                ("Customer", h.get("CUSTOMER NO.", ""), "👤"),
                ("Total", f"${money(h.get('INVOICE AMOUNT', 0))}", "💰"),
            ]:
                if val:
                    st.markdown(f"{ic} **{lbl}:** `{val}`")
                    if show_evidence:
                        ev = snippet_evidence(data["text"], str(val).replace("$", "").replace(",", ""))
                        if ev: st.caption(f"📍 _{ev}_")

            st.markdown("##### Context / Metadata")
            fl = inv.get("CONTEXT", {})
            for k, v in fl.items():
                if v: st.markdown(f"**{k}:** `{v}`")

            st.markdown("##### 💬 Comments")
            com = inv.get("COMMENTS", {})
            for k, v in com.items():
                if v: st.markdown(f"**{k}:** {v}")

            if not missing_df.empty:
                st.markdown("##### ⚠️ Capture Issues")
                st.dataframe(missing_df, width="stretch", hide_index=True)

            # Tabs
            tabs = st.tabs(["📋 Parsed", f"💳 {fmt_choice}", "🧩 All Formats", "✅ Validation", "📊 Line Items", "📝 Raw"])

            with tabs[0]:
                st.json(payload)
                st.download_button("💾 Parsed JSON", json.dumps(payload, indent=2),
                    f"{fname.rsplit('.',1)[0]}_parsed.json", "application/json", width="stretch", key=f"dp_{fname}")
                if is_connected():
                    if st.button("Send to QuickBooks", key=f"qb_{fname}"):
                        r = send_to_quickbooks(payload, classifier)
                        if r["success"]:
                            st.success(r["message"])
                        else:
                            st.warning(r["message"])
                if zoho_is_connected():
                    if st.button("Send to Zoho Books", key=f"zoho_{fname}"):
                        r = send_to_zoho(payload, classifier)
                        if r["success"]:
                            st.success(r["message"])
                        else:
                            st.warning(r["message"])
                if tally_is_connected():
                    blocker = tally_post_blocker(payload)
                    if blocker:
                        st.warning(blocker)
                    else:
                        b1, b2 = st.columns(2)
                        with b1:
                            if st.button("Dry Run Tally", key=f"tally_dry_{fname}"):
                                r = send_to_tally(payload, classifier, dry_run=True)
                                _store_tally_status(fname, r, dry_run=True)
                                if r["success"]:
                                    st.success(r["message"])
                                else:
                                    st.warning(r["message"])
                        with b2:
                            if st.button("Send to Tally", key=f"tally_{fname}"):
                                r = send_to_tally(payload, classifier, dry_run=False)
                                _store_tally_status(fname, r, dry_run=False)
                                if r["success"]:
                                    st.success(r["message"])
                                else:
                                    st.warning(r["message"])

            with tabs[1]:
                if fmt_choice == "Tally":
                    tally_xml = build_tally_xml(payload, classifier=classifier)
                    st.code(tally_xml, language="xml")
                    st.download_button("💾 Tally XML", tally_xml,
                        f"{fname.rsplit('.',1)[0]}_tally.xml", "application/xml", width="stretch", key=f"df_{fname}")
                else:
                    conv = CONVERTERS[fmt_choice](payload)
                    st.json(conv)
                    st.download_button(f"💾 {fmt_choice} JSON", json.dumps(conv, indent=2),
                        f"{fname.rsplit('.',1)[0]}_{fmt_choice.lower()}.json", "application/json", width="stretch", key=f"df_{fname}")

            with tabs[2]:
                fc = st.selectbox("Format", ["Parsed", "QuickBooks", "Tally", "Zoho Books", "Coupa", "NetSuite", "SAP"], key=f"af_{fname}")
                if fc == "Tally":
                    tally_xml = build_tally_xml(payload, classifier=classifier)
                    st.code(tally_xml, language="xml")
                    st.download_button("💾 Download XML", tally_xml,
                        f"{fname.rsplit('.',1)[0]}_tally.xml", "application/xml", width="stretch", key=f"ds_{fc}_{fname}")
                else:
                    d2show = payload if fc == "Parsed" else CONVERTERS[fc](payload)
                    st.json(d2show)
                    st.download_button("💾 Download", json.dumps(d2show, indent=2),
                        f"{fname.rsplit('.',1)[0]}_{fc.lower()}.json", "application/json", width="stretch", key=f"ds_{fc}_{fname}")

            with tabs[3]:
                ok2, iss2, dm2 = capture_check(payload)
                if ok2: st.success("✅ All fields captured. Totals consistent.")
                else:
                    for it in iss2: st.warning(f"⚠️ {it}")
                if not dm2.empty: st.dataframe(dm2, width="stretch", hide_index=True)
                st.markdown("##### Totals Reconciliation")
                rc1, rc2, rc3 = st.columns(3)
                rc1.metric("Invoice Total", f"${inv_total:,.2f}")
                rc2.metric("Line Sum", f"${line_sum:,.2f}")
                d = line_sum - inv_total
                rc3.metric("Delta", f"${d:,.2f}", delta="Match ✓" if abs(d) < 0.5 else "Mismatch ✗")

            with tabs[4]:
                if rows:
                    enriched = classifier.classify_invoice_rows(rows)
                    has_worksheet = bool(classifier and classifier.client_map)
                    display_rows = []
                    for r in enriched:
                        display_row = dict(r)
                        display_row.pop("SUBCATEGORY", None)
                        if not has_worksheet:
                            display_row["GL_CODE"] = ""
                            display_row["CLIENT_CATEGORY"] = ""
                        display_rows.append(display_row)
                    df_items = pd.DataFrame(display_rows)
                    preferred_cols = [
                        "ITEM", "DESCRIPTION", "QUANTITY", "UOM", "UNIT PRICE", "PRICE BASIS",
                        "EXTENDED AMOUNT", "TAX AMOUNT", "AMOUNT",
                        "CATEGORY", "GL_CODE", "CLIENT_CATEGORY", "MATCH_TIER",
                    ]
                    cols = [c for c in preferred_cols if c in df_items.columns]
                    cols += [c for c in df_items.columns if c not in cols]
                    st.dataframe(df_items[cols], width="stretch", hide_index=True)
                    if len(rows) > 1:
                        cdf = pd.DataFrame([{"Item": r["DESCRIPTION"][:25], "Amount": safe_float0(r.get("AMOUNT",0))} for r in rows])
                        st.bar_chart(cdf.nlargest(10, "Amount").set_index("Item"), color=theme["accent"], height=280)
                else: st.warning("No line items")

            with tabs[5]:
                if show_raw: st.text_area("PDF text", data["text"], height=320, key=f"rt_{fname}")
                else: st.info("Enable in sidebar")

else:
    _console_rendered = render_invoice_console(fmt_choice)
